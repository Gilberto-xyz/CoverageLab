"""Source-backed 5W1H exploration and reports for the local CoverageLab studio.

The report engine is vendored, and its original calculations are shared by web
previews and isolated exports. External workbooks are read-only. A coverage file
does not contain a geographic/channel breakdown: these are additional inputs.
"""
from __future__ import annotations

from delivery_files import auxiliary, system_dir

import dataclasses
from datetime import date, datetime
import hashlib
import json
import math
from pathlib import Path
import re
import shutil
import subprocess
import sys
import threading
import uuid
import research_comparison
from project_paths import external_research_directory

from openpyxl import load_workbook
from studio_inputs import _automatic_number_locale, _number_value, _date_value

from studio_5w1h_vendor import report_engine as engine
from studio_5w1h_vendor import template_engine as templates
from studio_5w1h_vendor.confidence_design import confidence_view

ENGINE_DIR = Path(__file__).resolve().parent / "studio_5w1h_vendor"
ANALYSIS_LOCK = threading.RLock()
SEGMENTS = [
    ("1", "Cuándo · evolución MAT", "1_MarcaEjemplo"),
    ("2", "Por qué · árbol de medidas", "2_MarcaEjemplo_K"),
    ("3-1", "Qué · tamaños", "3_MarcaEjemplo_1"),
    ("3-2", "Qué · marcas", "3_MarcaEjemplo_2"),
    ("3-3", "Qué · sabores", "3_MarcaEjemplo_3"),
    ("3-4", "Qué · subsegmentos", "3_MarcaEjemplo_4"),
    ("3-5", "Qué · subcategorías", "3_MarcaEjemplo_5"),
    ("4", "Quiénes · nivel socioeconómico", "4_MarcaEjemplo"),
    ("5-1", "Dónde · regiones", "5_MarcaEjemplo_1"),
    ("5-2", "Dónde · canales", "5_MarcaEjemplo_2"),
    ("5-3", "Dónde · canal moderno", "5_MarcaEjemplo_3"),
    ("5-4", "Dónde · canal tradicional", "5_MarcaEjemplo_4"),
    ("6", "Players", "6_Categoria_XX"),
    ("6-1", "Players · precio indexado", "6_Categoria_1"),
    ("7", "Distribución", "7_Categoria_Canal"),
    ("8", "Intervalos de confianza", "8_Categoria"),
]
MODES = [
    {"id": "1", "name": "Compras y ventas juntas", "description": "Comparación en el mismo gráfico."},
    {"id": "2", "name": "Compras y ventas separadas", "description": "Un gráfico por origen."},
    {"id": "3", "name": "Solo bloque base", "description": "Muestra únicamente el primer bloque de compras."},
]


def _json_value(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, dict):
        return {str(k): _json_value(v) for k, v in value.items()}
    if isinstance(value, (tuple, list)):
        return [_json_value(v) for v in value]
    if hasattr(value, "item"):
        return _json_value(value.item())
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def _month(value):
    if value in (None, ""):
        return None
    if not isinstance(value, str) or not re.fullmatch(r"\d{4}-\d{2}(?:-\d{2})?", value):
        raise ValueError("El período debe usar AAAA-MM.")
    return datetime.strptime(value[:7], "%Y-%m")


def _next_month(start, offset):
    number = start.year * 12 + start.month - 1 + offset
    return datetime(number // 12, number % 12 + 1, 1)


def _metadata(path, file_id, read_only):
    parts = path.stem.split("_", 2)
    if len(parts) != 3 or not parts[0].isdigit():
        raise ValueError("El nombre debe ser país_categoría_fabricante.xlsx.")
    country_code = int(parts[0])
    basket, category = engine.lookup_category_info(parts[1])
    manufacturer = engine.CLIENT_NAME_SUFFIX_PATTERN.sub("", parts[2]).strip()
    metadata = {"country_code": str(country_code), "country": engine.lookup_country_name(country_code),
                "category_code": parts[1].upper(), "category": category,
                "manufacturer": manufacturer or parts[2], "basket": basket}
    return {"id": file_id, "file": file_id, "name": path.name, "metadata": metadata,
            **metadata, "read_only": read_only, "size": path.stat().st_size}


def _sheet_identity(name):
    match = re.match(r"^(6[-_]1|[1-8])_(.*)$", name.strip())
    if not match:
        return None
    segment, rest = match.groups()
    dimension = ""
    if segment in ("3", "5"):
        parsed = engine.parse_standard_sheet_name(name)
        if parsed and parsed["variant"] in ("1", "2", "3", "4", "5"):
            rest = parsed["label"]
            segment = segment + "-" + parsed["variant"]
    elif segment == "6" and rest.endswith("_1"):
        rest, segment = rest[:-2], "6-1"
    elif segment == "2":
        rest = re.sub(r"_[ULKTRMH](?:_pp)?$", "", rest, flags=re.I)
    elif segment == "7":
        rest, _, dimension = rest.rpartition("_")
    elif segment == "6":
        rest = rest.rsplit("_", 1)[0]
    segment = segment.replace("_", "-")
    title = next((label for key, label, _ in SEGMENTS if key == segment), segment)
    return {"name": name, "segment": segment, "kind": segment, "title": title,
            "brand": rest, "dimension": dimension}


def _has_numeric_data(ws):
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            try:
                _number_value(cell.value, auto_locale='decimal-dot')
                return True
            except ValueError:
                pass
    return False


def _analysis_frame(frame):
    """Interpreta una copia en memoria al analizar, sin modificar el Excel de entrada."""
    if isinstance(frame, dict):
        return {key:_analysis_frame(value) for key,value in frame.items()}
    locale = _automatic_number_locale(frame.values.tolist(), 'decimal-dot')
    def value(raw):
        if not isinstance(raw, str):
            return raw
        try:
            return _date_value(raw)
        except ValueError:
            pass
        try:
            number = _number_value(raw, auto_locale=locale)
            return number*100 if raw.strip().endswith('%') else number
        except ValueError:
            return raw
    return frame.apply(lambda column: column.map(value))


def _analysis_excel_type(base):
    class ResearchExcel(base):
        def parse(self, *args, **kwargs):
            return _analysis_frame(super().parse(*args, **kwargs))
    return ResearchExcel


def _table(frame, table_id, title):
    return {"id": table_id, "title": title, "columns": [str(c) for c in frame.columns],
            "rows": _json_value(frame.values.tolist())}


def _chart(chart_id, title, dates, series, chart_type="line"):
    return {"id": chart_id, "title": title, "type": chart_type, "dates": dates,
            "series": series}


def _base_label(label):
    return re.sub(r"\.[CV]$", "", str(label), flags=re.I)


def _display_origin_pairs(items, field="id", separator="-", purchases_only=False):
    """Empareja cada métrica con el mismo orden de orígenes que la presentación."""
    groups = {}
    for item in items:
        value = item.get(field, "")
        base, _, origin = value.rpartition(separator)
        key = base if origin in {"Ventas", "Compras"} else value
        groups.setdefault(key, []).append(item)
    priority = {"monthly": 0, "mat": 1, "share": 2, "variation": 3, "contribution": 4}
    ordered = sorted(groups, key=lambda key: priority.get(key, 5)) if field == "id" else groups
    return [item for key in ordered for item in sorted(groups[key],
        key=lambda entry: engine.origin_display_order(entry.get(field, "").rpartition(separator)[-1], purchases_only))]


def _label_origin_outputs(result):
    labels = {origin['role']: origin['label'] for origin in result.get('origins', [])}
    def display(text):
        return re.sub(r'\b(Compras|Ventas|compras|ventas)\b', lambda match: labels.get(match[0].capitalize(), match[0]), str(text))
    if all(labels.get(role, role) == role for role in ('Compras', 'Ventas')):
        return
    for chart in result['charts']:
        chart['title'] = display(chart['title'])
        for series in chart['series']:
            series['name'] = display(series['name'])
            if 'origin' in series:
                series['origin_role'] = series['origin']
                series['origin'] = labels.get(series['origin'], series['origin'])
    for table in result['tables']:
        table['title'] = display(table['title'])
    for stat in result['stats']:
        stat['label'] = display(stat['label'])
    result['warnings'] = [display(warning) for warning in result['warnings']]


def create_research_template(spec: dict, destination: Path) -> Path:
    """Create the original dynamic structure, with every demo observation blank."""
    destination = Path(destination).resolve()
    if destination.suffix.lower() != ".xlsx" or destination.exists():
        raise ValueError("La plantilla necesita un archivo .xlsx nuevo.")
    country_code = str(spec.get("country_code", ""))
    category_code = str(spec.get("category_code", "")).upper()
    if not country_code.isdigit() or int(country_code) not in engine.COUNTRY_LOOKUP:
        raise ValueError("Selecciona un país válido.")
    if category_code not in engine.CATEGORY_LOOKUP:
        raise ValueError("Selecciona una categoría válida.")
    brands = spec.get("brands", [])
    if not isinstance(brands, list) or not 1 <= len(brands) <= 40:
        raise ValueError("Indica entre una y cuarenta marcas.")
    brands = [str(b).strip() for b in brands]
    if any(not b or len(b) > 100 for b in brands):
        raise ValueError("Nombre de marca inválido.")
    selected = spec.get("segments", ["1", "2", "3-1", "3-2", "4", "5-1", "5-2", "6", "7", "8"])
    if not isinstance(selected, list) or not selected:
        raise ValueError("Selecciona al menos un segmento 5W1H.")
    mapping = {key: sheet for key, _, sheet in SEGMENTS}
    if any(str(key) not in mapping and str(key) not in mapping.values() for key in selected):
        raise ValueError("Segmento 5W1H desconocido.")
    sheet_names = list(dict.fromkeys(mapping.get(str(key), str(key)) for key in selected))
    start = _month(spec.get("start_date") or f"{date.today().year - 2}-01")
    months = spec.get("months", 36)
    if isinstance(months, bool) or not isinstance(months, int) or not 1 <= months <= 120:
        raise ValueError("La historia debe tener entre 1 y 120 meses.")
    unit = str(spec.get("unit", "K")).strip().upper()
    if mapping['2'] in sheet_names and unit not in templates.TREE_UNIT_MAP:
        raise ValueError("Unidad del árbol inválida. Usa K, L, U, T, R, M o H para la hoja Árbol de medidas.")
    category_label = str(spec.get("category_label") or engine.lookup_category_info(category_code)[1])
    destination.parent.mkdir(parents=True, exist_ok=True)
    # Include two extra template types absent from the old CLI's selectable list.
    original_segments = templates.TEMPLATE_SEGMENTS
    with ANALYSIS_LOCK:
        templates.TEMPLATE_SEGMENTS = [(label, sheet) for _, label, sheet in SEGMENTS]
        try:
            created = templates.crear_excel_desde_plantilla(
                str(destination), brands, category_label, engine.lookup_country_name(int(country_code)),
                sheet_names, str(spec.get("players_label") or "Players"),
                str(spec.get("distribution_label") or "Canal"), unit,
                "_pp" if spec.get("penetration_pp") else "",
                summary_lines=["Entrada vacía: completar observaciones reales antes de analizar.",
                               "Canales, regiones y NSE requieren datos adicionales; no se calculan desde cobertura."])
        finally:
            templates.TEMPLATE_SEGMENTS = original_segments
    workbook = load_workbook(created)
    for ws in workbook:
        if ws.title == "README":
            continue
        # All numeric observations/formulas from the demonstration workbook are removed.
        for row in ws:
            for cell in row:
                if isinstance(cell.value, (int, float)) or cell.data_type == "f":
                    cell.value = None
                elif isinstance(cell.value, str):
                    if "table -" in cell.value.lower():
                        cell.value = "table"
                    elif any(sample in cell.value for sample in ("NOMENCOL", "Cp - Palmolive", "_PERIODS =")):
                        cell.value = "Datos del estudio · completar con la fuente"
        identity = _sheet_identity(ws.title)
        if identity and identity["segment"] not in ("2", "7", "8"):
            # Monthly formats always have a row-one header and explicit date columns.
            date_cols = [cell.column for cell in ws[1]
                         if str(cell.value or "").lower().startswith(("fecha", "compras", "ventas"))]
            for row_idx in range(2, max(ws.max_row, months + 1) + 1):
                for col in date_cols:
                    ws.cell(row_idx, col, _next_month(start, row_idx - 2) if row_idx <= months + 1 else None)
                    if row_idx > months + 1:
                        ws.cell(row_idx, col).value = None
        elif identity and identity["segment"] == "2":
            end = _next_month(start, months - 1)
            ws.cell(1, 1).value = f"Marca = {identity['brand']}"
            ws.cell(2, 1).value = "table"
            ws.cell(2, 2).value = "MAT " + _next_month(end, -12).strftime("%b-%y")
            ws.cell(2, 3).value = "MAT " + end.strftime("%b-%y")
        elif identity and identity["segment"] == "7":
            # Replace brands copied from the old template's cheese example.
            for idx in range(2, ws.max_column + 1):
                ws.cell(3, idx).value = category_label if idx == 2 else (brands[idx-3] if idx-3 < len(brands) else f"Elemento {idx-2} · completar")
        elif identity and identity["segment"] == "8":
            block_index = 0
            aggregate_index = 0
            period_index = 0
            for row in ws:
                if row[0].value and len(row) > 1 and row[1].value == "Weighted PENET":
                    label = brands[block_index] if block_index < len(brands) else f"Elemento {block_index+1} · completar"
                    row[0].value = label
                    if len(row) > 5:
                        row[5].value = label
                    block_index += 1
                    period_index = 0
                elif isinstance(row[0].value, str) and re.fullmatch(r"\s*[A-Za-z]{3}-\d{2,4}\s*", row[0].value):
                    row[0].value = _next_month(start, period_index) if period_index < months else None
                    period_index += 1
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.strip().startswith("MAT "):
                        offset = -12 if cell.column == ws.max_column - 1 else 0
                        cell.value = "MAT " + _next_month(start, months - 1 + offset).strftime("%b-%y")
                if len(row) > 5 and isinstance(row[5].value, str):
                    aggregate_label = row[5].value.strip()
                    if aggregate_label and not aggregate_label.lower().startswith(("weighted", "table", "hholds", "buyers", "mat ")):
                        row[5].value = brands[aggregate_index] if aggregate_index < len(brands) else f"Elemento {aggregate_index+1} · completar"
                        aggregate_index += 1
    workbook.save(destination)
    workbook.close()
    return destination


class ResearchStudio:
    def __init__(self, base_dir: Path, external_dir: Path | None = None):
        self.base_dir = Path(base_dir).resolve()
        self.local_dir = self.base_dir / "inputs" / "5w1h"
        self.external_dir = Path(external_dir or external_research_directory(self.base_dir)).resolve()

    def source(self, file_id: str) -> Path:
        if not isinstance(file_id, str):
            raise ValueError("Archivo 5W1H inválido.")
        match = re.fullmatch(r"(local|external)[:/]([^/\\:]+\.xlsx)", file_id, flags=re.I)
        if not match or match.group(2).startswith(("~$", ".")):
            raise ValueError("Archivo 5W1H inválido.")
        origin, name = match.groups()
        directory = self.local_dir if origin.lower() == "local" else self.external_dir
        if origin.lower() == "local" and not directory.resolve().is_relative_to(self.base_dir):
            raise ValueError("La carpeta de investigación está fuera del proyecto.")
        candidate = (directory / name).resolve()
        if candidate.parent != directory.resolve() or not candidate.is_file():
            raise ValueError("No se encontró el archivo 5W1H.")
        return candidate

    def catalog(self):
        files = []
        for origin, directory in (("local", self.local_dir), ("external", self.external_dir)):
            if not directory.is_dir():
                continue
            for path in sorted(directory.glob("*.xlsx"), key=lambda p: p.name.casefold()):
                if path.name.startswith("~$") or path.stem.lower().startswith("plantilla"):
                    continue
                try:
                    files.append(_metadata(path, f"{origin}/{path.name}", origin == "external"))
                except ValueError:
                    continue
        return {"files": files, "modes": MODES,
                "segments": [{"id": key, "name": label, "label": label} for key, label, _ in SEGMENTS],
                "units": templates.TREE_UNIT_MAP,
                "note": "Canales, regiones y NSE necesitan sus propios datos 5W1H."}

    def inspect(self, file_id):
        path = self.source(file_id)
        metadata = _metadata(path, file_id, file_id.startswith("external"))
        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            sheets = []
            for ws in workbook:
                identity = _sheet_identity(ws.title)
                if identity:
                    sheets.append({**identity, "rows": ws.max_row, "columns": ws.max_column,
                                   "has_data": _has_numeric_data(ws)})
            present = {s["segment"] for s in sheets}
            return {"file": metadata, "metadata": metadata["metadata"], "sheets": sheets,
                    "unavailable_segments": [{"id": key, "title": title} for key, title, _ in SEGMENTS if key not in present],
                    "requires_input": not any(s["has_data"] for s in sheets)}
        finally:
            workbook.close()

    def analyze(self, file_id, sheet, mode="1", start=None, end=None, elements=None, pipeline=None):
        if str(mode) not in {m["id"] for m in MODES}:
            raise ValueError("Modo de investigación inválido.")
        if pipeline is not None and (isinstance(pipeline, bool) or not isinstance(pipeline, int) or not 0 <= pipeline <= 6):
            raise ValueError("Pipeline debe estar entre P0 y P6.")
        first, last = _month(start), _month(end)
        if first and last and first > last:
            raise ValueError("El período inicial debe preceder al final.")
        if elements is not None and (not isinstance(elements, list) or any(not isinstance(x, str) for x in elements)):
            raise ValueError("Selección de elementos inválida.")
        info = self.inspect(file_id)
        identity = next((s for s in info["sheets"] if s["name"] == sheet), None)
        if identity is None:
            raise ValueError("La hoja seleccionada no es un segmento 5W1H.")
        path = self.source(file_id)
        provenance = {"file": file_id, "sheet": sheet, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                      "engine": "5W1H original · motor integrado", "mode": str(mode), "pipeline": pipeline,
                      "filters": {"start": start, "end": end, "elements": elements},
                      "export_scope": "Informe completo de las hojas seleccionadas; los filtros web son de exploración."}
        result = {"file": info["file"], "sheet": sheet, "title": identity["title"] + " · " + identity["brand"],
                  "charts": [], "tables": [], "stats": [], "filters": {"dates": [], "elements": []},
                  "warnings": [], "provenance": provenance, "ready": False}
        if not identity["has_data"]:
            result["warnings"].append("La hoja está vacía. Pega observaciones reales en el editor para analizarla.")
            return result
        with ANALYSIS_LOCK:
            engine.ensure_heavy_modules_loaded()
            excel = _analysis_excel_type(engine.pd.ExcelFile)(path)
            try:
                self._analyze(excel, identity, result, str(mode), first, last, elements, pipeline)
            finally:
                excel.close()
        result["ready"] = bool(result["charts"] or result["tables"])
        purchases_only = len(result.get('origins', [])) == 2 and all(origin['kind'] == 'Compras' for origin in result['origins'])
        result["charts"] = _display_origin_pairs(result["charts"], purchases_only=purchases_only)
        result["tables"] = _display_origin_pairs(result["tables"], purchases_only=purchases_only)
        result["stats"] = _display_origin_pairs(result["stats"], field="label", separator=" · ", purchases_only=purchases_only)
        if purchases_only:
            for chart in result['charts']:
                metric = chart['id'].rsplit('-', 1)[0]
                paired = [c for c in result['charts'] if c['id'].rsplit('-', 1)[0] == metric]
                if len(paired) != 2:
                    continue
                values = [v for c in paired for s in c['series'] for v in s['values'] if isinstance(v, (int, float)) and math.isfinite(v)]
                if values:
                    chart['axis_range'] = [min(0, min(values)), max(0, max(values))]
        _label_origin_outputs(result)
        if elements == []:
            result["charts"] = []
            result["ready"] = False
            result["warnings"].append("No hay elementos seleccionados. Marca al menos uno para explorar y generar; los totales visibles conservan la base completa de la hoja.")
        return _json_value(result)

    def _analyze(self, excel, identity, result, mode, first, last, elements, pipeline):
        sheet, segment = identity["name"], identity["segment"]
        lang = "P" if result["file"]["country"] == "Brasil" else "E"
        if segment == "7":
            frame, source_label = engine.parse_distribution_sheet(excel, sheet, result["file"]["manufacturer"])
            names = [str(c) for c in frame.columns[1:]]
            result["filters"]["elements"] = names
            if elements is not None and any(name not in names for name in elements):
                raise ValueError("El elemento seleccionado no existe en esta hoja.")
            chosen = set(names if elements is None else elements)
            result["tables"].append(_table(frame[[frame.columns[0]] + [name for name in names if name in chosen]], "distribution", "Distribución · datos de la fuente"))
            series = [{"name": name, "values": frame[name].tolist(), "unit": "%",
                       "color": engine.TREND_COLOR_SEQUENCE[i % len(engine.TREND_COLOR_SEQUENCE)], "origin": source_label}
                      for i, name in enumerate(names) if name in chosen]
            result["charts"].append(_chart("distribution", identity["title"], frame.iloc[:, 0].astype(str).tolist(), series, "bar"))
            return
        if segment == "2":
            frame, periods = engine._limpiar_tabla_excel(excel.parse(sheet, header=None))
            if frame is None or not periods or frame.iloc[:, 1:].notna().sum().sum() == 0:
                result["warnings"].append("Completa los valores MAT del árbol de medidas.")
                return
            unit = engine._unidad_desde_nombre_hoja(sheet)
            metrics = engine.calcular_cambios(frame, periods[0], periods[-1], unidad=unit,
                                             penetracion_en_pp=engine._modo_penetracion_desde_nombre_hoja(sheet) == "pp")
            incomplete = [name for name, metric in metrics.items()
                          if not all(isinstance(metric[key], (int, float)) and math.isfinite(metric[key])
                                     for key in ("valor_origen", "value"))]
            if incomplete:
                result["warnings"].append("Completa ambos períodos MAT de: " + ", ".join(incomplete))
                return
            rows = []
            for index, (name, metric) in enumerate(metrics.items()):
                # The engine's labels already contain 000s; reuse its presentation
                # formatter so the web does not show raw currency/people as thousands.
                old = float(engine._format_val(name, metric["valor_origen"]).replace(",", ""))
                current = float(engine._format_val(name, metric["value"]).replace(",", ""))
                rows.append([name, old, current, metric["change"], metric["change_suffix"]])
                result["charts"].append(_chart(f"tree-{index}", name, [periods[0], periods[-1]],
                    [{"name": name, "values": [old, current], "color": "#00968f", "unit": ""}], "bar"))
            result["tables"].append({"id": "tree", "title": "Árbol · niveles y cambio", "columns": ["Métrica", periods[0], periods[-1], "Cambio", "Unidad cambio"], "rows": rows})
            result["provenance"]["unit"] = unit
            result["provenance"]["tree_display"] = "Niveles con las unidades, escala y precisión del informe original; cambio calculado sobre datos sin redondear."
            return
        if segment == "8":
            raw = excel.parse(sheet, header=None)
            monthly = engine.ppt8_parse_monthly_blocks(raw)
            aggregate = engine.ppt8_parse_agg_blocks(raw)
            aggregate, _ = engine.ppt8_filter_agg_blocks(monthly, aggregate)
            rows, errors = engine.ppt8_compute_rows(monthly, aggregate)
            result["warnings"].extend(errors)
            if rows:
                records = [dataclasses.asdict(row) for row in rows]
                columns = list(records[0])
                result["tables"].append({"id": "confidence", "title": "Intervalos de confianza y error muestral", "columns": columns,
                                         "display_columns": ["Marca", "Variación fabricante", "Variación MAT", "Límite inferior", "Límite superior", "Compradores · prom. 12m", "Penetración · prom. 12m", "Nivel de penetración", "Error muestral"],
                                         "column_units": ["", "%", "%", "%", "%", "", "%", "", "percent"],
                                         "confidence_view": confidence_view(rows),
                                         "rows": [[record[c] for c in columns] for record in records]})
                result["charts"].append(_chart("confidence", "Variación MAT e intervalo de confianza", [r.marca for r in rows],
                    [{"name": label, "values": [getattr(r, field) for r in rows], "color": color, "unit": "%"}
                     for field, label, color in (("lim_inf", "Límite inferior", "#94a3b8"), ("mat_sep25", "Variación MAT", "#00968f"), ("lim_sup", "Límite superior", "#2c3e50"))], "bar"))
                result["provenance"]["interval_z"] = engine.PPT8_Z_SCORE_INTERVAL
            return
        if segment == "1":
            frame = engine.parse_sheet_with_date_fallback(excel, sheet)
            if frame.iloc[:, 1:].apply(engine.pd.to_numeric, errors="coerce").notna().sum().sum() == 0:
                return
            raw_pipeline = engine._detect_ventas_pipeline(frame.columns)
            p = raw_pipeline if pipeline is None else pipeline
            if len(frame) < 24 + p:
                result["warnings"].append(f"El comparativo MAT e interanual necesita {24 + p} meses; esta hoja tiene {len(frame)}.")
                return
            mat = engine.df_mat(frame.copy(), p)
            valid = mat.dropna(how="all", subset=list(mat.columns[1:]))
            dates = [v.strftime("%Y-%m") for v in valid.iloc[:, 0]]
            mask = [not (first and v < first or last and v > last) for v in valid.iloc[:, 0]]
            selected_dates = [d for d, keep in zip(dates, mask) if keep]
            for chart_id, cols in (("mat", list(valid.columns[1:3])), ("yoy", list(valid.columns[3:]))):
                result["charts"].append(_chart(chart_id, "Volumen MAT" if chart_id == "mat" else "Variación interanual MAT", selected_dates,
                    [{"name": str(c), "values": [v for v, keep in zip(valid[c], mask) if keep], "unit": "percent" if chart_id == "yoy" else "",
                      "color": ("#d4ac0d", "#2c3e50")[i % 2]} for i, c in enumerate(cols)]))
            result["tables"].append(_table(valid, "mat", "Evaluación MAT y variación"))
            result["filters"]["dates"] = dates
            return
        frame = engine.parse_sheet_with_compras_header(excel, sheet)
        blocks, detected_pipeline = engine.split_research_origins(frame, sheet_name=sheet)
        configs = engine.prepare_series_configs(blocks, lang, detected_pipeline if pipeline is None else pipeline, sheet_name=sheet)
        # Use the original semantic ordering to compare the same channel across origins.
        configs = engine.reorder_series_configs_by_compras(configs)
        comparison_colors = research_comparison.palette(configs, engine.TREND_COLOR_SEQUENCE)
        result['origins'] = [{'role': c.display_tipo, 'label': engine.series_display_label(c),
                              'kind': c.data.attrs.get('source_kind', c.display_tipo), 'pipeline': c.pipeline} for c in configs]
        result['provenance']['origins'] = result['origins']
        for config in configs:
            if config.data.attrs.get('missing_comparison'):
                result['warnings'].append(f"El bloque {config.data.attrs['missing_comparison']} está vacío y todavía no se puede comparar.")
        if len(configs) == 2 and all(c.data.attrs.get('source_kind') == 'Compras' for c in configs):
            result['comparison_note'] = 'Comparación entre versiones de compras: ambos bloques conservan sus fechas y se calculan sin desfase.'
            ranges = [(c.data.iloc[0, 0], c.data.iloc[-1, 0]) for c in configs]
            if ranges[0] != ranges[1]:
                result['warnings'].append('Los bloques cubren períodos distintos. Las tendencias se alinean por mes; revisa los cortes indicados en cada tabla MAT.')
        if mode == "3":
            configs = [c for c in configs if c.display_tipo == "Compras"]
        else:
            result['purchase_comparison'] = research_comparison.mat_model(configs, engine.pd, str(last.date()) if last is not None else None)
            if result['purchase_comparison'] and not result['purchase_comparison']['ready']:
                result['warnings'].append(result['purchase_comparison']['note'])
        all_elements = list(dict.fromkeys(_base_label(col) for c in configs for col in c.data.columns[1:]
                                         if engine.pd.to_numeric(c.data[col], errors="coerce").notna().any()))
        if elements is not None and any(name not in all_elements for name in elements):
            raise ValueError("El elemento seleccionado no existe en esta hoja.")
        chosen = set(all_elements if elements is None else elements)
        comparison = result.get('purchase_comparison')
        if comparison and comparison.get('ready'):
            comparison['rows'] = [row for row in comparison['rows'] if any(_base_label(label) in chosen for label in (row['left_label'], row['right_label']) if label)]
        result["filters"]["elements"] = all_elements
        date_union = sorted(set(v.strftime("%Y-%m") for c in configs for v in c.data.iloc[:, 0]))
        dates = [d for d in date_union if not (first and d < first.strftime("%Y-%m") or last and d > last.strftime("%Y-%m"))]
        result["filters"]["dates"] = date_union
        monthly_combined, mat_combined = [], []
        palette = {name: engine.TREND_COLOR_SEQUENCE[i % len(engine.TREND_COLOR_SEQUENCE)] for i, name in enumerate(all_elements)}
        for config in configs:
            for col in config.data.columns[1:]:
                if col in comparison_colors: palette[_base_label(col)] = comparison_colors[col]
        for config in configs:
            data = config.data.copy()
            # A cut changes the trailing MAT windows; a visual start does not remove their history.
            if last:
                data = data[data.iloc[:, 0] <= last].reset_index(drop=True)
            period_numbers = [dt.year * 12 + dt.month for dt in data.iloc[:, 0]]
            contiguous = all(b == a + 1 for a, b in zip(period_numbers, period_numbers[1:]))
            if not contiguous:
                result["warnings"].append(f"{config.display_tipo}: hay meses repetidos o faltantes. Completa la secuencia para calcular MAT.")
            if segment == "6-1":
                ratio, columns, original = engine.prepare_price_index_dataframe(config._replace(data=data), [])
                if ratio.empty:
                    result["warnings"].append("Completa precios y una base de categoría positiva para calcular el índice.")
                    continue
                ratio = ratio[ratio.iloc[:, 0].map(lambda dt: not (first and dt < first or last and dt > last))]
                result["charts"].append(_chart(f"price-{config.display_tipo}", f"Precio indexado · {config.display_tipo} · categoría = 100", [dt.strftime("%Y-%m") for dt in ratio.iloc[:, 0]],
                    [{"name": _base_label(col), "values": ratio[col].tolist(), "unit": "index", "color": palette.get(_base_label(col), "#64748b")}
                     for col in columns if col == "Total" or _base_label(col) in chosen]))
                result["tables"].append(_table(ratio, f"price-{config.display_tipo}", "Índice de precio · categoría = 100"))
                continue
            monthly_series, mat_series = [], []
            for col in data.columns[1:]:
                name = _base_label(col)
                if name not in chosen:
                    continue
                numbers = engine.pd.to_numeric(data[col], errors="coerce")
                if not numbers.notna().any():
                    continue
                shifted = numbers.shift(config.pipeline)
                rolling = shifted.rolling(12, min_periods=12).sum() if contiguous else shifted * float("nan")
                monthly = {dt.strftime("%Y-%m"): val for dt, val in zip(data.iloc[:, 0], shifted)}
                mats = {dt.strftime("%Y-%m"): val for dt, val in zip(data.iloc[:, 0], rolling)}
                label = name + " · " + config.display_tipo
                common = {"name": label, "color": palette[name], "unit": "", "origin": config.display_tipo,
                          "pipeline": config.pipeline, "dash": config.display_tipo == "Ventas"}
                monthly_series.append({**common, "values": [monthly.get(dt) for dt in dates]})
                mat_series.append({**common, "values": [mats.get(dt) for dt in dates]})
            relevant = data.iloc[max(0, len(data) - 24 - config.pipeline):len(data) - config.pipeline if config.pipeline else len(data), 1:]
            complete = not relevant.apply(engine.pd.to_numeric, errors="coerce").isna().any().any()
            if not complete:
                result["warnings"].append(f"{config.display_tipo}: faltan observaciones numéricas en la base del comparativo MAT.")
            apo = engine.aporte(data.copy(), config.pipeline, lang, engine.series_display_label(config)) if contiguous and complete else engine.pd.DataFrame()
            if apo.attrs.get("skip_table"):
                result["warnings"].append(f"{config.display_tipo}: el comparativo MAT necesita {24 + config.pipeline} meses completos.")
            elif not apo.empty:
                keep_cols = [apo.columns[0]] + [c for c in apo.columns[1:] if _base_label(c) in chosen or c == "Total"]
                result["tables"].append(_table(apo[keep_cols], f"contribution-{config.display_tipo}", f"{config.display_tipo} · participación, variación y aporte MAT"))
                shares = apo.attrs.get("share_period_values", [])
                if shares:
                    series = [{"name": _base_label(col), "color": palette.get(_base_label(col), "#64748b"),
                               "values": [record.get(col) for _, record in shares], "unit": "percent"}
                              for col in shares[0][1] if _base_label(col) in chosen]
                    result["charts"].append(_chart(f"share-{config.display_tipo}", f"Participación MAT · {config.display_tipo}", [period for period, _ in shares], series, "bar"))
                for row_index, label in ((2, "Variación total MAT"), (3, "Aporte total MAT")):
                    columns = [col for col in apo.columns[1:] if col != "Total" and _base_label(col) in chosen]
                    values = [engine._parse_percent_cell(apo.loc[row_index, col]) for col in columns]
                    result["charts"].append(_chart(
                        f"{'variation' if row_index == 2 else 'contribution'}-{config.display_tipo}",
                        f"{'Variación MAT' if row_index == 2 else 'Aporte al crecimiento MAT'} · {config.display_tipo}",
                        [_base_label(col) for col in columns],
                        [{"name": "Variación" if row_index == 2 else "Aporte", "values": values,
                          "unit": "%" if row_index == 2 else "pp", "color": "#078c86"}], "bar"))
                    if "Total" in apo.columns:
                        value = engine._parse_percent_cell(apo.loc[row_index, "Total"])
                        result["stats"].append({"label": label + " · " + config.display_tipo, "value": value / 100 if value is not None else None, "unit": "percent"})
            if mode == "2":
                result["charts"].append(_chart(f"monthly-{config.display_tipo}", f"Tendencia mensual · {config.display_tipo}", dates, monthly_series))
                result["charts"].append(_chart(f"mat-{config.display_tipo}", f"Tendencia MAT · {config.display_tipo}", dates, mat_series))
            else:
                monthly_combined.extend(monthly_series)
                mat_combined.extend(mat_series)
        if mode != "2" and segment != "6-1":
            result["charts"][:0] = [_chart("monthly", "Tendencia mensual · compras y ventas" if mode == "1" else "Tendencia mensual · compras", dates, monthly_combined),
                                    _chart("mat", "Tendencia de volumen MAT", dates, mat_combined)]
            align = research_comparison.alignment(configs)
            if align:
                for chart in result['charts'][:2]:
                    visible = {s['name'].split(' · ')[0] for s in chart['series']}
                    chart['comparison_legend'] = [{'name':_base_label(a or b), 'color':comparison_colors[a or b]} for a,b in align['pairs']
                                                   if any(_base_label(label) in visible for label in (a,b) if label)]
                    chart['comparison_legend'] += [{'name':align['labels'][0], 'color':'#172c3c'},
                                                   {'name':align['labels'][1], 'color':'#172c3c', 'dash':True}]
        result["charts"] = [c for c in result["charts"] if any(any(v is not None and not (isinstance(v, float) and not math.isfinite(v)) for v in s["values"]) for s in c["series"])]
        if chosen != set(all_elements):
            result["warnings"].append("Participaciones y aportes conservan como base el total de la hoja, incluso al ocultar elementos.")
        if not result["charts"]:
            result["warnings"].append("No hay datos numéricos disponibles en el período seleccionado.")

    def export(self, file_id, options=None, destination=None, *, progress=None):
        """Run the original full report against an immutable workbook copy."""
        options = dict(options or {})
        mode = str(options.get("mode", "1"))
        if mode not in {m["id"] for m in MODES}:
            raise ValueError("Modo de investigación inválido.")
        info = self.inspect(file_id)
        source = self.source(file_id)
        sheets = options.get("sheets") or [s["name"] for s in info["sheets"] if s["has_data"]]
        valid_names = {s["name"] for s in info["sheets"] if s["has_data"]}
        if not isinstance(sheets, list) or not sheets or any(s not in valid_names for s in sheets):
            raise ValueError("El informe necesita hojas con datos reales.")
        # A partially filled tree must never be exported with NaN levels or zero changes.
        for sheet in sheets:
            if _sheet_identity(sheet)["segment"] == "2":
                preview = self.analyze(file_id, sheet, mode=mode)
                if not preview["ready"]:
                    raise ValueError(f"{sheet}: " + " ".join(preview["warnings"]))
        if any(options.get(key) for key in ("start", "end", "elements")):
            raise ValueError("Los filtros son de exploración; exporta el informe completo de las hojas seleccionadas.")
        fingerprint = hashlib.sha256(source.read_bytes()).hexdigest()
        expected = options.get("fingerprint")
        if expected and expected != fingerprint:
            raise ValueError("El archivo cambió; vuelve a revisar la investigación.")
        pipeline = options.get("pipeline")
        if pipeline is not None and (isinstance(pipeline, bool) or not isinstance(pipeline, int) or not 0 <= pipeline <= 6):
            raise ValueError("Pipeline debe estar entre P0 y P6.")
        created = datetime.now()
        output_root = (self.base_dir / "Output").resolve()
        automatic_destination = destination is None or Path(destination).resolve() == output_root
        run_suffix = f"{created:%Y%m%d_%H%M%S}_{uuid.uuid4().hex[:6]}"
        if automatic_destination:
            if not output_root.is_relative_to(self.base_dir):
                raise ValueError("Output debe estar dentro del proyecto.")
            meta = info["metadata"]
            stem = "-".join(templates.sanitizar_segmento_archivo(str(meta[key]), 45) for key in ("country", "category", "manufacturer"))
            destination = output_root / f"{stem}_5W1H_{run_suffix}"
        destination = Path(destination).resolve()
        if destination.exists() and any(destination.iterdir()):
            raise ValueError("La carpeta de entrega debe estar vacía.")
        destination.mkdir(parents=True, exist_ok=True)
        # Only registered source paths are read. The worker never writes into the external program.
        work_dir = system_dir(destination) / "_input"
        work_dir.mkdir(exist_ok=True)
        copied = work_dir / source.name
        shutil.copy2(source, copied)
        if hashlib.sha256(copied.read_bytes()).hexdigest() != fingerprint:
            raise ValueError("El archivo cambió mientras se preparaba la entrega; vuelve a revisarlo.")
        overrides = options.get("pipeline_overrides", {})
        if not isinstance(overrides, dict) or any(not isinstance(k, str) or isinstance(v, bool) or not isinstance(v, int) or not 0 <= v <= 6 for k, v in overrides.items()):
            raise ValueError("La selección de pipelines por marca es inválida.")
        normalized = {engine.normalize_brand_key(k): v for k, v in overrides.items()}
        applied = {}
        if pipeline is not None or overrides or len(sheets) != len(info["sheets"]):
            wb = load_workbook(copied)
            for title in list(wb.sheetnames):
                if title not in sheets and title.strip().lower() != "readme":
                    wb.remove(wb[title])
            for ws in wb:
                if ws.title not in sheets:
                    continue
                identity = _sheet_identity(ws.title)
                selected_pipeline = normalized.get(engine.normalize_brand_key(identity["brand"]), pipeline)
                if selected_pipeline is None:
                    continue
                for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12)):
                    for cell in row:
                        if isinstance(cell.value, str) and re.fullmatch(r"Ventas[_ -]\d+", cell.value.strip(), flags=re.I):
                            cell.value = f"Ventas_{selected_pipeline}"
                            applied[ws.title] = selected_pipeline
            wb.save(copied)
            wb.close()
        request = {"source": str(copied), "destination": str(destination), "mode": mode, "sheets": sheets,
                   "size_threshold": options.get("size_threshold"), "result": str(auxiliary(destination, "_result.json", write=True))}
        if progress:
            request['progress'] = progress
        threshold = request["size_threshold"]
        if threshold is not None and (isinstance(threshold, bool) or not isinstance(threshold, (int, float)) or not 0 < threshold < 100):
            raise ValueError("El umbral de agrupación debe estar entre 0 y 100%.")
        request_path = work_dir / "request.json"
        request_path.write_text(json.dumps(request, ensure_ascii=False), encoding="utf-8")
        log_path = auxiliary(destination, "generacion_5w1h.log", write=True)
        with log_path.open("w", encoding="utf-8") as log:
            process = subprocess.run([sys.executable, "-X", "utf8", str(Path(__file__).resolve()), str(request_path)],
                                     cwd=destination, stdout=log, stderr=subprocess.STDOUT, timeout=600,
                                     creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        if process.returncode != 0:
            excerpt = log_path.read_text(encoding="utf-8", errors="replace")[-2500:]
            raise RuntimeError("El motor 5W1H no pudo completar el informe. " + excerpt)
        returned = json.loads(auxiliary(destination, "_result.json").read_text(encoding="utf-8"))
        outputs = {label: returned[label] for label in ("Excel", "PPT")}
        for path in outputs.values():
            candidate = Path(path).resolve()
            if candidate.parent != destination or not candidate.is_file():
                raise RuntimeError("El motor produjo una entrega inválida.")
        if automatic_destination:
            # The original engine determines the report cut from the selected
            # data. Reuse that canonical title, including MM-YY, after it finishes.
            title = returned["title"]
            if not isinstance(title, str) or not title or re.search(r'[<>:"/\\|?*\x00-\x1f]', title) or title in {".", ".."}:
                raise RuntimeError("El motor produjo un nombre de entrega inválido.")
            final_directory = (output_root / f"{title}__{run_suffix}").resolve()
            original_directory = destination.resolve()
            # Validate both absolute targets before the directory move, including
            # symlink resolution, so it can never leave this project's Output.
            if (original_directory.parent != output_root or final_directory.parent != output_root
                    or original_directory == output_root or final_directory == output_root
                    or not original_directory.is_dir() or final_directory.exists()):
                raise RuntimeError("No se pudo asignar una carpeta segura a la entrega.")
            original_directory.rename(final_directory)
            destination = final_directory
            outputs = {label: str(destination / Path(path).name) for label, path in outputs.items()}
            log_path = auxiliary(destination, log_path.name, write=True)
            returned.update(outputs)
            auxiliary(destination, "_result.json", write=True).write_text(json.dumps(returned, ensure_ascii=False), encoding="utf-8")
            request.update(source=str(destination / "_sistema" / "_input" / source.name), destination=str(destination),
                           result=str(auxiliary(destination, "_result.json")))
            (destination / "_sistema" / "_input" / "request.json").write_text(json.dumps(request, ensure_ascii=False), encoding="utf-8")
            with log_path.open("a", encoding="utf-8") as log:
                log.write(f"\nCarpeta final de esta entrega: {destination}\n")
        provenance = {"source": file_id, "source_sha256": fingerprint, "created_at": created.isoformat(),
                      "options": options, "sheets": sheets, "metadata": info["metadata"],
                      "pipelines_applied": applied,
                      "scope": "Informe completo de las hojas seleccionadas", "engine": json.loads((ENGINE_DIR / "provenance.json").read_text(encoding="utf-8"))}
        provenance_path = auxiliary(destination, "investigacion_5w1h.json", write=True)
        provenance_path.write_text(json.dumps(provenance, ensure_ascii=False, indent=2), encoding="utf-8")
        outputs.update(Provenance=str(provenance_path), Registro=str(log_path))
        return {"directory": str(destination), "outputs": outputs, "title": returned["title"],
                "created_at": created.isoformat(), "source": file_id, "fingerprint": fingerprint,
                "pipelines_applied": applied}


if __name__ == '__main__':
    request = json.loads(Path(sys.argv[1]).read_text(encoding='utf-8'))
    engine.ensure_heavy_modules_loaded()
    engine.pd.ExcelFile = _analysis_excel_type(engine.pd.ExcelFile)
    reporter = None
    if request.get('progress'):
        from export_progress import ExportProgressReporter, RESEARCH_STAGES
        spec = request['progress']
        reporter = ExportProgressReporter(spec['path'], spec['started_at'], spec.get('estimate_seconds'), stages=RESEARCH_STAGES)
        reporter.update('starting', 1, 1)
    result = engine.run_report(request['source'], request['destination'], request.get('mode','1'),
                               request.get('sheets'), request.get('size_threshold'), progress=reporter.update if reporter else None)
    Path(request['result']).write_text(json.dumps(result, ensure_ascii=False), encoding='utf-8')
