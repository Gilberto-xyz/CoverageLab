"""Excel de entrada de los tres momentos de CoverageLab, sin automatizar Excel.

Las referencias públicas son nombres de archivo, nunca rutas. Las entradas de
cobertura permanecen en la raíz del proyecto para que el motor existente pueda
usarlas. Los libros 5W1H externos se leen y, al editarlos, se copian a inputs/5w1h.
"""
from __future__ import annotations

import base64
import binascii
from datetime import date, datetime
import hashlib
import io
import math
import os
from pathlib import Path
import re
import tempfile
import threading
import unicodedata
import zipfile

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

import archivos_studio as nomenclature
from project_paths import external_research_directory


MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ARCHIVE_BYTES = 100 * 1024 * 1024
MAX_ROWS = 20000
MAX_COLUMNS = 256
MAX_SHEETS = 100
MAX_EDIT_CELLS = 50000
MAX_READ_CELLS = 20000
MAX_CELL_TEXT = 32767
_FILE_LOCKS: dict[str, threading.RLock] = {}
_LOCKS_LOCK = threading.Lock()
_MONTHS = {name: index for index, names in enumerate([
    ("ene", "jan", "enero", "january"), ("feb", "febrero", "february"),
    ("mar", "marzo", "march"), ("abr", "apr", "abril", "april"),
    ("may", "mayo"), ("jun", "junio", "june"), ("jul", "julio", "july"),
    ("ago", "aug", "agosto", "august"), ("sep", "sept", "septiembre", "september"),
    ("oct", "octubre", "october"), ("nov", "noviembre", "november"),
    ("dic", "dec", "diciembre", "december"),
], 1) for name in names}
_DISPLAY_MONTHS = ("ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic")


class InputConflictError(ValueError):
    """El usuario debe volver a leer el archivo antes de guardar."""

    status = 409


class InputNumberFormatError(ValueError):
    """Los separadores del número no corresponden al formato de la fuente."""


def file_fingerprint(path: Path) -> str:
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def _normalized(value: str) -> str:
    return "".join(char for char in unicodedata.normalize("NFD", value.strip().lower())
                   if unicodedata.category(char) != "Mn")


def _integer(value, label: str, minimum: int, maximum: int) -> int:
    if isinstance(value, bool):
        raise ValueError(f"{label} debe ser un número entero.")
    try:
        result = int(value)
    except (ValueError, TypeError, OverflowError) as exc:
        raise ValueError(f"{label} debe ser un número entero.") from exc
    if str(result) != str(value) or not minimum <= result <= maximum:
        raise ValueError(f"{label} debe estar entre {minimum} y {maximum}.")
    return result


def _filename(value: str) -> str:
    if not isinstance(value, str) or not value or len(value) > 200:
        raise ValueError("Nombre de archivo inválido.")
    if re.search(r'[<>:"/\\|?*\x00-\x1f]', value) or value in {".", ".."}:
        raise ValueError("Usa solamente el nombre del archivo, sin carpetas.")
    if value != value.strip() or value.endswith((" ", ".")) or not value.lower().endswith(".xlsx"):
        raise ValueError("El archivo debe tener extensión .xlsx.")
    if value.startswith("~$"):
        raise ValueError("No se puede usar un archivo temporal de Excel.")
    if re.fullmatch(r"(?:CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(?:\..*)?", value, re.I):
        raise ValueError("Ese nombre está reservado por Windows.")
    return value


def _manufacturer(value) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError("Escribe el nombre del fabricante.")
    clean = value.strip()
    if len(clean) > 100 or re.search(r'[<>:"/\\|?*\x00-\x1f]', clean) or clean.endswith("."):
        raise ValueError("El fabricante contiene caracteres inválidos para un archivo.")
    return clean


def _sheet_title(value: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError("Escribe un nombre para la hoja.")
    name = value.strip()
    if len(name) > 31 or re.search(r"[\\/*\[\]:?\x00-\x1f]", name) or name.startswith("'") or name.endswith("'"):
        raise ValueError("La hoja admite hasta 31 caracteres y no admite \\ / * [ ] : ?.")
    return name


def _date_value(value, *, epoch=None) -> datetime:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if not math.isfinite(value) or not 1 <= value <= 109574:
            raise ValueError("Fecha serial de Excel inválida.")
        parsed = from_excel(value, epoch=epoch) if epoch is not None else from_excel(value)
        if not isinstance(parsed, datetime):
            raise ValueError("Fecha serial de Excel inválida.")
        return parsed
    if not isinstance(value, str):
        raise ValueError("Usa una fecha como 2026-01-01, ene 26 o 01/01/2026.")
    clean = _normalized(value)
    try:
        return datetime.fromisoformat(clean.replace("z", "+00:00")).replace(tzinfo=None)
    except ValueError:
        pass
    match = re.fullmatch(r"([a-z]+)[ /-]+(\d{2}|\d{4})", clean)
    if match and match[1] in _MONTHS:
        year = int(match[2])
        return datetime(year + 2000 if year < 100 else year, _MONTHS[match[1]], 1)
    match = re.fullmatch(r"(\d{4})[-/](\d{1,2})", clean)
    if match:
        try:
            return datetime(int(match[1]), int(match[2]), 1)
        except ValueError:
            pass
    match = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", clean)
    if match:
        try:
            return datetime(int(match[3]), int(match[2]), int(match[1]))
        except ValueError:
            pass
    raise ValueError(f"Fecha inválida: {value}. Usa 2026-01-01, ene 26 o 01/01/2026.")


def _automatic_number_locale(values, fallback):
    """Usa separadores inequívocos del bloque antes del formato regional."""
    evidence = set()
    for row in values:
        for raw in row:
            raw = raw.get('value') if isinstance(raw, dict) else raw
            if not isinstance(raw, str):
                continue
            clean = raw.strip().replace('\u00a0', '').replace(' ', '').rstrip('%')
            if not re.fullmatch(r'[+-]?[\d.,]+', clean):
                continue
            if ',' in clean and '.' in clean:
                locale = 'decimal-dot' if clean.rfind('.') > clean.rfind(',') else 'decimal-comma'
                try:
                    _number_value(raw, number_locale=locale)
                    evidence.add(locale)
                except ValueError:
                    pass
            else:
                match = re.fullmatch(r'[+-]?(\d+)([.,])(\d+)', clean)
                if match and (len(match[3]) != 3 or not match[1].strip('0') or len(match[1]) > 3):
                    evidence.add('decimal-dot' if match[2] == '.' else 'decimal-comma')
                elif re.fullmatch(r'[+-]?\d{1,3}(?:,\d{3}){2,}', clean):
                    evidence.add('decimal-dot')
                elif re.fullmatch(r'[+-]?\d{1,3}(?:\.\d{3}){2,}', clean):
                    evidence.add('decimal-comma')
    return next(iter(evidence)) if len(evidence) == 1 else fallback


def _number_value(value, *, number_locale: str = "auto", auto_locale: str | None = None) -> int | float:
    if isinstance(value, bool):
        raise ValueError("Usa un número, no verdadero/falso.")
    if isinstance(value, (int, float)):
        try:
            finite = math.isfinite(value)
        except OverflowError:
            finite = False
        if not finite:
            raise ValueError("El número debe ser finito.")
        return value
    if not isinstance(value, str):
        raise ValueError("Usa un número válido.")
    clean = value.strip().replace("\u00a0", "").replace(" ", "")
    percent = clean.endswith("%")
    if percent:
        clean = clean[:-1]
    if not re.fullmatch(r"[+\-\d.,eE]+", clean) or not re.search(r"\d", clean):
        raise ValueError(f"Número inválido: {value}.")
    if number_locale in {"decimal-comma", "decimal-dot"}:
        decimal, grouping = (",", ".") if number_locale == "decimal-comma" else (".", ",")
        parts = clean.split(decimal)
        if len(parts) > 2:
            raise InputNumberFormatError(f"Número inválido para el formato seleccionado: {value}.")
        integer = parts[0]
        if grouping in integer:
            if not re.fullmatch(r"[+-]?\d{1,3}(?:" + re.escape(grouping) + r"\d{3})+", integer):
                raise InputNumberFormatError(f"Número inválido para el formato seleccionado: {value}.")
            integer = integer.replace(grouping, "")
        if len(parts) == 2 and grouping in parts[1]:
            raise InputNumberFormatError(f"Número inválido para el formato seleccionado: {value}.")
        clean = integer + ("." + parts[1] if len(parts) == 2 else "")
    elif "." in clean and "," in clean:
        decimal, grouping = (".", ",") if clean.rfind(".") > clean.rfind(",") else (",", ".")
        parts = clean.split(decimal)
        if len(parts) != 2 or not re.fullmatch(r"[+-]?\d{1,3}(?:" + re.escape(grouping) + r"\d{3})+", parts[0]):
            raise InputNumberFormatError(f"Número inválido: {value}.")
        clean = parts[0].replace(grouping, "") + "." + parts[1]
    else:
        match = re.fullmatch(r"([+-]?\d+)[.,](\d{3})", clean)
        if match and abs(int(match[1])) > 0:
            if auto_locale in {'decimal-dot', 'decimal-comma'}:
                return _number_value(value, number_locale=auto_locale)
            raise InputNumberFormatError(
                f"Número ambiguo: {value}. Selecciona punto o coma decimal antes de guardar.")
        grouping = "," if "," in clean else "."
        if clean.count(grouping) > 1 and re.fullmatch(r"[+-]?\d{1,3}(?:" + re.escape(grouping) + r"\d{3})+", clean):
            clean = clean.replace(grouping, "")
        else:
            clean = clean.replace(",", ".")
    if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", clean):
        raise InputNumberFormatError(f"Número inválido: {value}.")
    result = float(clean)
    if not math.isfinite(result):
        raise ValueError("El número debe ser finito.")
    if percent:
        result /= 100
    return int(result) if result.is_integer() else result


def _check_archive(data: bytes) -> None:
    if not isinstance(data, bytes) or len(data) > MAX_FILE_BYTES:
        raise ValueError("El libro admite como máximo 10 MiB.")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 10000 or sum(item.file_size for item in entries) > MAX_ARCHIVE_BYTES:
                raise ValueError("El contenido del libro excede el límite de tamaño.")
            names = {item.filename for item in entries}
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                raise ValueError("El archivo no es un libro Excel .xlsx válido.")
            if any(item.flag_bits & 1 for item in entries):
                raise ValueError("Los libros cifrados no se pueden editar en la web.")
            if any("vbaproject" in item.filename.lower() for item in entries):
                raise ValueError("Solo se admiten libros .xlsx sin macros.")
    except (zipfile.BadZipFile, OSError) as exc:
        raise ValueError("El archivo no es un libro Excel .xlsx válido.") from exc


def _bounds(workbook) -> None:
    if len(workbook.worksheets) > MAX_SHEETS:
        raise ValueError(f"El libro admite como máximo {MAX_SHEETS} hojas.")
    if any(sheet.max_row > MAX_ROWS or sheet.max_column > MAX_COLUMNS for sheet in workbook.worksheets):
        raise ValueError(f"Cada hoja admite como máximo {MAX_ROWS} filas y {MAX_COLUMNS} columnas.")


def _load(path: Path):
    data = path.read_bytes()
    _check_archive(data)
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
        _bounds(workbook)
        return workbook
    except Exception as exc:
        if isinstance(exc, ValueError) and str(exc).startswith("El libro"):
            raise
        raise ValueError(f"No se pudo leer el libro: {exc}") from exc


def _lock(path: Path):
    with _LOCKS_LOCK:
        return _FILE_LOCKS.setdefault(str(path).casefold(), threading.RLock())


def _coverage_header(sheet) -> int:
    weighted = None
    for row in range(1, min(sheet.max_row, 50) + 1):
        first, second = sheet.cell(row, 1).value, sheet.cell(row, 2).value
        if isinstance(first, str) and "table" in first.lower():
            return row
        if isinstance(second, str) and "weighted" in second.lower():
            weighted = row
    return weighted or 2


def _date_columns(sheet, kind: str, *, epoch) -> dict[int, int]:
    """Columnas de períodos del contrato, sin inferir fechas de las métricas.

    Un número de ventas nunca se trata como un serial por su magnitud. En los
    bloques 5W1H el título y su separador identifican la columna del calendario.
    Los árboles y la distribución contienen encabezados MAT, no fechas de filas.
    """
    if kind == "coverage":
        return {1: _coverage_header(sheet) + 1}
    segment = re.match(r"^([1-8])[-_]", sheet.title)
    if not segment or segment[1] in {"2", "7"}:
        return {}
    columns = {}
    if segment[1] in {"3", "4", "5", "6"}:
        for row in sheet.iter_rows(min_row=1, max_row=min(50, sheet.max_row)):
            for cell in row:
                if not isinstance(cell.value, str) or not re.fullmatch(
                        r"(?:compras|ventas(?:[_ -]?\d+)?|fecha|periodo)", _normalized(cell.value)):
                    continue
                # Un bloque de ventas empieza después de una columna vacía;
                # evita confundir una métrica titulada Ventas con su calendario.
                if cell.column != 1 and sheet.cell(cell.row, cell.column - 1).value is not None:
                    continue
                columns.setdefault(cell.column, cell.row + 1)
        return columns
    # Evolución (1) y confianza (8) tienen su calendario en A. La confianza
    # incluye una marca encima; busca el primer mes sin convertir su encabezado.
    for row in range(1, min(sheet.max_row, 50) + 1):
        value = sheet.cell(row, 1).value
        if segment[1] == "1" and isinstance(value, str) and _normalized(value) in {"fecha", "periodo", "base"}:
            return {1: row + 1}
        if segment[1] == "8" and isinstance(value, str):
            headers = [sheet.cell(row, col).value for col in (2, 3)]
            if all(isinstance(header, str) for header in headers) and (
                    "penet" in _normalized(headers[0]) and "buyers" in _normalized(headers[1])):
                return {1: row + 1}
        if not isinstance(value, (date, datetime, str)):
            continue
        try:
            _date_value(value, epoch=epoch)
        except ValueError:
            continue
        return {1: row}
    return {}


def _display_date(value, *, epoch) -> tuple[str | None, str]:
    if value is None or value == "":
        return None, "blank"
    try:
        parsed = _date_value(value, epoch=epoch)
    except (ValueError, OverflowError, TypeError):
        return str(value), "invalid-date"
    return f"{_DISPLAY_MONTHS[parsed.month - 1]} {parsed.year}", "date"


def _coverage_sheet(workbook, name: str, manufacturer: str, unit: str | None = None) -> None:
    sheet = workbook.create_sheet(name)
    sheet.append(["Linea PowerView", "Weighted " + unit if unit else None])
    sheet.append(["Base", "Weighted R_VOL1", "Weighted PENET", "Weighted VO1_BUY",
                  "Weighted VO1_DAY", "Weighted FREQ", "BUYERS", "SELL-IN: " + manufacturer])
    # Conserva exactamente los valores y posiciones de archivos_studio.py.
    # Los formatos ayudan al llenado y no cambian el contrato posicional.
    sheet.freeze_panes = "B3"
    for cell in sheet[2]:
        cell.font = Font(bold=True, color="173442")
        cell.fill = PatternFill("solid", fgColor="EAF5F2")
        cell.alignment = Alignment(wrap_text=True)
    sheet.row_dimensions[2].height = 30
    sheet.column_dimensions["A"].width = 17
    for column in range(2, 9):
        sheet.column_dimensions[get_column_letter(column)].width = 22


class InputStudio:
    def __init__(self, root: Path | str, research_root: Path | str | None = None):
        self.root = Path(root).resolve()
        self.research_local = self.root / "inputs" / "5w1h"
        self.research_external = (Path(research_root).resolve() if research_root is not None
                                  else external_research_directory(self.root))

    def _path(self, file: str, kind: str = "coverage") -> tuple[Path, str]:
        if kind == "coverage":
            base, name, origin = self.root, _filename(file), "local"
        elif kind == "research":
            if not isinstance(file, str):
                raise ValueError("Referencia de archivo inválida.")
            match = re.fullmatch(r"(local|external)[:/](.+)", file)
            if not match:
                raise ValueError("Usa una referencia local:archivo.xlsx o external:archivo.xlsx.")
            origin, name = match[1], _filename(match[2])
            base = self.research_local if origin == "local" else self.research_external
        else:
            raise ValueError("Tipo de entrada desconocido.")
        path = (base / name).resolve()
        if origin == "local" and not base.resolve().is_relative_to(self.root):
            raise ValueError("La carpeta de entradas está fuera del proyecto.")
        if not path.is_relative_to(base.resolve()):
            raise ValueError("El archivo está fuera de la carpeta de entradas.")
        if not path.is_file():
            raise ValueError("El archivo de entrada no existe.")
        return path, origin

    def catalog(self) -> dict:
        seen = set()
        countries = []
        for alias, code in nomenclature.countries.items():
            if code in seen:
                continue
            seen.add(code)
            name = nomenclature.obtener_nombre_pais(alias)
            countries.append({"code": code, "name": name,
                              "population_coverage": nomenclature.pop_coverage.get(name)})
        return {"countries": countries,
                "categories": [{"id": str(index), "code": item["cod"],
                                "name": item["descripcion"], "basket": item["categoria"]}
                               for index, item in enumerate(nomenclature.categorias)],
                "limits": {"file_bytes": MAX_FILE_BYTES, "rows": MAX_ROWS,
                           "columns": MAX_COLUMNS, "sheets": MAX_SHEETS,
                           "edit_cells": MAX_EDIT_CELLS, "read_cells": MAX_READ_CELLS}}

    def _metadata(self, name: str, kind: str, catalog: dict | None = None) -> dict:
        """La identidad se lee del nombre que utilizan los scripts existentes."""
        match = re.fullmatch(r"(\d+)_([^_]+)_(.+)\.xlsx", name, re.I)
        if not match:
            return {"recognized": False, "country_code": None, "country": None,
                    "category_code": None, "category": None, "category_id": None,
                    "basket": None, "manufacturer": None}
        catalog = catalog or self.catalog()
        country_code, category_code, manufacturer = match.groups()
        country = next((item for item in catalog["countries"] if item["code"] == country_code), None)
        categories = [item for item in catalog["categories"] if item["code"].casefold() == category_code.casefold()]
        category = categories[0] if categories else None
        if kind == "research":
            manufacturer = re.sub(r"[_\s]+5w1h(?= ?\(\d+\)$|$)", "", manufacturer, flags=re.I)
        return {"recognized": bool(country and category), "country_code": country_code,
                "country": country["name"] if country else None,
                "category_code": category_code.upper(), "category": category["name"] if category else None,
                "category_id": category["id"] if len(categories) == 1 else None,
                "basket": category["basket"] if category else None, "manufacturer": manufacturer}

    def files(self) -> dict:
        result = {"coverage": [], "research": []}
        catalog = self.catalog()
        for kind, origin, base in (("coverage", "local", self.root),
                                   ("research", "local", self.research_local),
                                   ("research", "external", self.research_external)):
            if not base.is_dir():
                continue
            if origin == "local" and not base.resolve().is_relative_to(self.root):
                continue
            for path in sorted(base.glob("*.xlsx"), key=lambda item: item.name.casefold()):
                if path.name.startswith("~$") or path.name == "file_temp_coverage.xlsx":
                    continue
                try:
                    path = path.resolve()
                    if not path.is_relative_to(base.resolve()) or not path.is_file():
                        continue
                    stat = path.stat()
                    ref = path.name if kind == "coverage" else f"{origin}/{path.name}"
                    result[kind].append({"file": ref, "kind": kind, "name": path.name,
                                         "origin": origin, "editable": origin == "local",
                                         "size": stat.st_size, "modified": stat.st_mtime,
                                         "metadata": self._metadata(path.name, kind, catalog)})
                except OSError:
                    continue
        return result

    def _unique(self, base: Path, name: str) -> Path:
        _filename(name)
        base.mkdir(parents=True, exist_ok=True)
        if not base.resolve().is_relative_to(self.root):
            raise ValueError("La carpeta de entradas está fuera del proyecto.")
        candidate = base / name
        index = 1
        while candidate.exists():
            candidate = base / f"{Path(name).stem} ({index}).xlsx"
            index += 1
            if index > 10000:
                raise ValueError("Hay demasiados archivos con el mismo nombre.")
        return candidate

    def _write(self, workbook, path: Path, expected: str | None = None) -> None:
        _bounds(workbook)
        temp_path = None
        try:
            with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".xlsx", delete=False) as temp:
                temp_path = Path(temp.name)
            workbook.save(temp_path)
            if temp_path.stat().st_size > MAX_FILE_BYTES:
                raise ValueError("El libro guardado excede el límite de 10 MiB.")
            if expected is not None:
                if not path.is_file() or file_fingerprint(path) != expected:
                    raise InputConflictError("El archivo cambió. Vuelve a cargarlo antes de guardar.")
                os.replace(temp_path, path)
            else:
                # O_EXCL impide que una segunda creación sobrescriba la primera.
                with path.open("xb") as destination, temp_path.open("rb") as source:
                    while chunk := source.read(1024 * 1024):
                        destination.write(chunk)
        except PermissionError as exc:
            raise ValueError("No se pudo guardar. Cierra el archivo en Excel y vuelve a intentar.") from exc
        finally:
            if temp_path is not None:
                temp_path.unlink(missing_ok=True)

    def _identity(self, spec: dict) -> tuple[dict, dict, str]:
        if not isinstance(spec, dict):
            raise ValueError("Datos de creación inválidos.")
        catalog = self.catalog()
        country_code = str(spec.get("country_code", ""))
        country = next((item for item in catalog["countries"] if item["code"] == country_code), None)
        if country is None:
            raise ValueError("Selecciona un país del catálogo.")
        category_id = spec.get("category_id")
        category_code = spec.get("category_code")
        category = next((item for item in catalog["categories"]
                         if str(item["id"]) == str(category_id)), None) if category_id is not None else None
        if category is None and category_id is None:
            category = next((item for item in catalog["categories"] if item["code"] == category_code), None)
        if category is None:
            raise ValueError("Selecciona una categoría del catálogo.")
        return country, category, _manufacturer(spec.get("manufacturer"))

    def create(self, spec: dict) -> dict:
        country, category, manufacturer = self._identity(spec)
        kind = spec.get("kind", "coverage")
        suffix = "_5w1h" if kind == "research" else ""
        name = f"{country['code']}_{category['code']}_{manufacturer}{suffix}.xlsx"
        base = self.research_local if kind == "research" else self.root
        if not isinstance(kind, str) or kind not in {"coverage", "research"}:
            raise ValueError("Tipo de entrada desconocido.")
        with _lock(base):
            path = self._unique(base, name)
            if kind == "research":
                from studio_5w1h import create_research_template
                complete = {**spec, "country_code": country["code"], "category_code": category["code"],
                            "manufacturer": manufacturer, "category_label": spec.get("category_label") or category["name"]}
                create_research_template(complete, path)
            else:
                brands = spec.get("brands") or [manufacturer]
                unit = spec.get("unit")
                if unit is not None:
                    if not isinstance(unit, str) or not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_ .-]{0,49}", unit.strip()):
                        raise ValueError("Usa una unidad como R_VOL1, VOL1_P, VOL1_KG o VOL1_LTS.")
                    unit = re.sub(r"^Weighted\s+", "", unit.strip(), flags=re.I)
                if not isinstance(brands, list) or not 1 <= len(brands) <= MAX_SHEETS:
                    raise ValueError(f"Agrega entre 1 y {MAX_SHEETS} marcas.")
                workbook = openpyxl.Workbook()
                workbook.remove(workbook.active)
                titles = set()
                for brand in brands:
                    brand = _manufacturer(brand)
                    title = _sheet_title("P0_" + brand)
                    if title.casefold() in titles:
                        raise ValueError("Los nombres de las marcas deben ser distintos.")
                    titles.add(title.casefold())
                    _coverage_sheet(workbook, title, brand, unit)
                self._write(workbook, path)
        ref = path.name if kind == "coverage" else f"local/{path.name}"
        return self.workbook(ref, kind=kind)

    def import_workbook(self, payload: dict) -> dict:
        if not isinstance(payload, dict):
            raise ValueError("Datos de importación inválidos.")
        kind = payload.get("kind", "coverage")
        if not isinstance(kind, str) or kind not in {"coverage", "research"}:
            raise ValueError("Tipo de entrada desconocido.")
        name = _filename(payload.get("filename"))
        data = payload.get("content", payload.get("data"))
        if isinstance(data, str):
            if len(data) > ((MAX_FILE_BYTES + 2) // 3) * 4:
                raise ValueError("El libro admite como máximo 10 MiB.")
            try:
                data = base64.b64decode(data, validate=True)
            except (binascii.Error, ValueError) as exc:
                raise ValueError("El contenido del archivo no es base64 válido.") from exc
        _check_archive(data)
        try:
            workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
            _bounds(workbook)
        except Exception as exc:
            raise ValueError(f"No se pudo leer el libro importado: {exc}") from exc
        reuse = kind == "research" and payload.get("reuse_existing") is True
        try:
            if reuse:
                from studio_5w1h import _sheet_identity
                if not re.fullmatch(r"\d+_[^_]+_.+\.xlsx", name, re.I):
                    raise ValueError("Usa el nombre país_categoría_fabricante.xlsx para identificar esta investigación.")
                if not any(_sheet_identity(sheet) for sheet in workbook.sheetnames):
                    raise ValueError("El archivo no contiene hojas con la estructura 5W1H. Puedes revisarlo en Archivos de entrada.")
        finally:
            workbook.close()
        base = self.root if kind == "coverage" else self.research_local
        with _lock(base):
            existing = (base / name).resolve()
            if reuse and existing.parent == base.resolve() and existing.is_file() and existing.read_bytes() == data:
                result = self.workbook(f"local/{existing.name}", kind=kind)
                result["reused_existing"] = True
                return result
            path = self._unique(base, name)
            # Copia los bytes originales para conservar íntegramente gráficos,
            # formatos y cachés de fórmulas en la importación.
            with path.open("xb") as destination:
                destination.write(data)
        ref = path.name if kind == "coverage" else f"local/{path.name}"
        return self.workbook(ref, kind=kind)

    def workbook(self, file: str, *, kind: str = "coverage", sheet: str | None = None,
                 row: int = 1, limit: int = 200, column: int = 1, columns: int = 32) -> dict:
        row = _integer(row, "La primera fila", 1, MAX_ROWS)
        limit = _integer(limit, "Las filas visibles", 1, 500)
        column = _integer(column, "La primera columna", 1, MAX_COLUMNS)
        columns = _integer(columns, "Las columnas visibles", 1, MAX_COLUMNS)
        if limit * columns > MAX_READ_CELLS or row + limit - 1 > MAX_ROWS or column + columns - 1 > MAX_COLUMNS:
            raise ValueError("La ventana de lectura excede el límite de filas o columnas.")
        path, origin = self._path(file, kind)
        file = path.name if kind == "coverage" else f"{origin}/{path.name}"
        with _lock(path):
            before = file_fingerprint(path)
            workbook = _load(path)
            if file_fingerprint(path) != before:
                raise InputConflictError("El archivo cambió mientras se leía. Vuelve a cargarlo.")
            selected = sheet or workbook.sheetnames[0]
            if selected not in workbook.sheetnames:
                raise ValueError("La hoja seleccionada no existe.")
            ws = workbook[selected]
            end_row = min(row + limit - 1, max(row, ws.max_row))
            end_column = min(column + columns - 1, max(column, ws.max_column))
            date_columns = _date_columns(ws, kind, epoch=workbook.epoch)
            values, types, readonly, formats, fills = [], [], [], [], []
            display_values, display_types = [], []
            for cells in ws.iter_rows(min_row=row, max_row=end_row, min_col=column, max_col=end_column):
                vals, typs, locked, fmts, colors = [], [], [], [], []
                shown, shown_types = [], []
                for cell in cells:
                    raw = value = cell.value
                    is_formula = cell.data_type == "f"
                    if isinstance(value, (datetime, date)):
                        value, typ = value.isoformat()[:10], "date"
                    elif is_formula:
                        typ = "formula"
                    elif isinstance(value, bool):
                        typ = "boolean"
                    elif isinstance(value, (int, float)):
                        value, typ = (value if math.isfinite(value) else None), "number"
                    else:
                        typ = "text" if value is not None else "blank"
                    if not is_formula and cell.column in date_columns and cell.row >= date_columns[cell.column]:
                        display, display_type = _display_date(raw, epoch=workbook.epoch)
                        if (display_type == "invalid-date" and kind == "research"
                                and ws.title.startswith(("8_", "8-")) and cell.column == 1
                                and all(ws.cell(cell.row, col).value is None for col in (2, 3))):
                            display_type = typ  # Separador/título de otra marca.
                    else:
                        display, display_type = (str(value) if value is not None else None), typ
                    vals.append(value)
                    typs.append(typ)
                    shown.append(display)
                    shown_types.append(display_type)
                    locked.append(is_formula or isinstance(cell, MergedCell) or bool(ws.protection.sheet))
                    fmts.append(cell.number_format)
                    color = cell.fill.fgColor
                    colors.append(color.rgb[-6:] if cell.fill.fill_type == "solid" and color.type == "rgb" else None)
                values.append(vals); types.append(typs); readonly.append(locked); formats.append(fmts); fills.append(colors)
                display_values.append(shown); display_types.append(shown_types)
            result = {"file": file, "kind": kind, "name": path.name, "fingerprint": before,
                      "metadata": self._metadata(path.name, kind),
                      "origin": origin, "editable": origin == "local", "copy_on_save": origin == "external",
                      "sheets": [{"name": ws.title, "rows": ws.max_row, "columns": ws.max_column}
                                 for ws in workbook.worksheets], "sheet": selected,
                      "row": row, "column": column, "values": values, "types": types,
                      "display_values": display_values, "display_types": display_types,
                      "date_columns": [{"column": col, "data_row": first} for col, first in sorted(date_columns.items())],
                      "readonly": readonly, "formats": formats, "fills": fills,
                      "validation": self._validation(workbook, kind)}
            if kind == "coverage":
                result["data_row"] = _coverage_header(ws) + 1
            return result

    def _validation(self, workbook, kind: str) -> dict:
        if kind != "coverage":
            return {"ready": None, "errors": [], "warnings": [],
                    "message": "Revisa los segmentos y períodos en el módulo 5W1H antes de generar."}
        details, all_errors, all_warnings = [], [], []
        for sheet in workbook.worksheets:
            errors, warnings, dates, records = [], [], [], 0
            if sheet.max_column < 8:
                errors.append("Faltan columnas: la entrada requiere fecha, seis KPI y Sell-in (8 columnas).")
            header = _coverage_header(sheet)
            for row in range(header + 1, sheet.max_row + 1):
                cells = [sheet.cell(row, column) for column in range(1, 9)]
                if all(cell.value is None for cell in cells):
                    continue
                records += 1
                try:
                    dates.append(_date_value(cells[0].value, epoch=workbook.epoch).date())
                except ValueError:
                    errors.append(f"A{row}: fecha inválida o vacía.")
                for column, cell in enumerate(cells[1:], 2):
                    if cell.data_type == "f":
                        errors.append(f"{cell.coordinate}: la entrada tiene una fórmula. Pega el valor calculado en una copia para analizar.")
                    elif cell.value is None:
                        if column in {2, 8}:
                            errors.append(f"{cell.coordinate}: falta {'Worldpanel' if column == 2 else 'Sell-in'}.")
                        else:
                            warnings.append(f"{cell.coordinate}: KPI sin dato; no se guardará como cero.")
                    else:
                        try:
                            _number_value(cell.value)
                        except ValueError:
                            errors.append(f"{cell.coordinate}: número inválido.")
            if not records:
                errors.append("Agrega las fechas y los datos mensuales debajo de los encabezados.")
            elif records < 12:
                warnings.append(f"Hay {records} meses. La cobertura MAT necesita 12 meses completos.")
            periods = [item.year * 12 + item.month - 1 for item in dates]
            if len(set(periods)) != len(periods):
                errors.append("Hay meses repetidos. Cada mes debe aparecer una vez por hoja.")
            if periods != sorted(periods):
                errors.append("Las fechas no están en orden cronológico.")
            unique_periods = sorted(set(periods))
            if any(current - previous != 1 for previous, current in zip(unique_periods, unique_periods[1:])):
                errors.append("Faltan meses entre las fechas. La cobertura MAT requiere meses consecutivos.")
            details.append({"sheet": sheet.title, "records": records,
                            "ready": not errors and records >= 12,
                            "errors": errors[:100], "warnings": warnings[:100]})
            all_errors.extend(f"{sheet.title}: {message}" for message in errors)
            all_warnings.extend(f"{sheet.title}: {message}" for message in warnings)
        return {"ready": all(item["ready"] for item in details), "sheets": details,
                "errors": all_errors[:100], "warnings": all_warnings[:100],
                "error_count": len(all_errors), "warning_count": len(all_warnings)}

    def _convert(self, raw, cell, *, numeric: bool, dated: bool, epoch, number_locale: str = "auto", auto_locale=None):
        typ = None
        if isinstance(raw, dict):
            if set(raw) - {"type", "value"} or raw.get("type") not in {"text", "number", "date", "blank"}:
                raise ValueError("Tipo de celda desconocido.")
            typ, raw = raw.get("type"), raw.get("value")
        if raw is None or raw == "" or typ == "blank":
            return None, None
        if isinstance(raw, str):
            if len(raw) > MAX_CELL_TEXT or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", raw):
                raise ValueError("El texto de la celda es demasiado largo o contiene caracteres inválidos.")
            if raw.lstrip().startswith("="):
                raise ValueError("Pega valores, no fórmulas. Las fórmulas existentes se conservan protegidas.")
        if dated or typ == "date":
            return _date_value(raw, epoch=epoch), "yyyy-mm-dd"
        if numeric or typ == "number":
            value = _number_value(raw, number_locale=number_locale, auto_locale=auto_locale)
            if isinstance(raw, str) and raw.strip().endswith("%"):
                return value, "0.0%"
            return value, None
        if typ == "text":
            if not isinstance(raw, str):
                raise ValueError("Una celda de texto debe contener texto.")
            return raw, None
        if isinstance(raw, bool):
            return raw, None
        if isinstance(raw, (int, float)):
            return _number_value(raw), None
        if not isinstance(raw, str):
            raise ValueError("Valor de celda inválido.")
        # En 5W1H las cabeceras incluyen períodos: las fechas reconocidas se
        # tipan, mientras que nombres, unidades y códigos conservan su texto.
        try:
            value = _date_value(raw)
            return value, "yyyy-mm-dd"
        except ValueError:
            pass
        try:
            value = _number_value(raw, number_locale=number_locale, auto_locale=auto_locale)
            if raw.strip().endswith("%"):
                return value, "0.0%"
            return value, None
        except ValueError as exc:
            if isinstance(exc, InputNumberFormatError):
                raise
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                raise ValueError("Esta celda requiere un número válido.")
            if isinstance(cell.value, (datetime, date)):
                raise ValueError("Esta celda requiere una fecha válida.")
            return raw, None

    @staticmethod
    def _copy_value(raw, cell):
        """Copia el contenido sin exigir el tipo que antes tenía la celda."""
        if isinstance(raw, dict):
            if set(raw) - {'type', 'value'} or raw.get('type') not in {'text','number','date','blank'}:
                raise ValueError('Tipo de celda desconocido.')
            raw = None if raw.get('type') == 'blank' else raw.get('value')
        if raw is None or raw == '':
            return None, None
        if isinstance(raw, str):
            if len(raw)>MAX_CELL_TEXT or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', raw):
                raise ValueError('El texto excede los límites de una celda de Excel.')
        elif not isinstance(raw, (bool, int, float, date, datetime)):
            raise ValueError('Valor de celda inválido.')
        elif isinstance(raw, (int, float)) and not math.isfinite(raw):
            raise ValueError('Excel no admite este valor numérico.')
        clear_date_format = openpyxl.styles.numbers.is_date_format(cell.number_format) and not isinstance(raw, (date,datetime))
        return raw, 'General' if clear_date_format else None

    def save(self, file: str, *, fingerprint: str, edits: list, kind: str = "coverage",
             number_locale: str = "auto", number_locale_hint: str = "decimal-dot", preserve_text: bool = False) -> dict:
        path, origin = self._path(file, kind)
        if not isinstance(preserve_text, bool):
            raise ValueError('Modo de guardado desconocido.')
        literal = kind == 'research' or preserve_text
        if not isinstance(number_locale, str) or number_locale not in {"auto", "decimal-comma", "decimal-dot"}:
            raise ValueError("Formato numérico desconocido. Usa automático, punto decimal o coma decimal.")
        if not isinstance(number_locale_hint, str) or number_locale_hint not in {'decimal-dot', 'decimal-comma'}:
            raise ValueError("Formato numérico regional desconocido.")
        if not isinstance(fingerprint, str) or not re.fullmatch(r"[a-f0-9]{64}", fingerprint):
            raise ValueError("Falta la versión del archivo. Vuelve a cargarlo antes de guardar.")
        if not isinstance(edits, list) or not edits or len(edits) > 1000:
            raise ValueError("Envía al menos un bloque de celdas para guardar (máximo 1000).")
        with _lock(path):
            if file_fingerprint(path) != fingerprint:
                raise InputConflictError("El archivo cambió. Vuelve a cargarlo antes de guardar.")
            workbook = _load(path)
            prepared, touched = [], set()
            for edit in edits:
                if not isinstance(edit, dict) or set(edit) - {"sheet", "row", "column", "values"}:
                    raise ValueError("Bloque de edición inválido.")
                name = edit.get("sheet")
                if name not in workbook.sheetnames:
                    raise ValueError("La hoja seleccionada no existe.")
                sheet = workbook[name]
                if sheet.protection.sheet:
                    raise ValueError(f"La hoja {name} está protegida. Se conserva su protección.")
                row = _integer(edit.get("row"), "La primera fila", 1, MAX_ROWS)
                column = _integer(edit.get("column"), "La primera columna", 1, MAX_COLUMNS)
                values = edit.get("values")
                if not isinstance(values, list) or not values or not all(isinstance(v, list) and v for v in values):
                    raise ValueError("El pegado debe ser una tabla rectangular con al menos una celda.")
                width = len(values[0])
                if any(len(v) != width for v in values):
                    raise ValueError("Todas las filas del pegado deben tener la misma cantidad de columnas.")
                if row + len(values) - 1 > MAX_ROWS or column + width - 1 > MAX_COLUMNS:
                    raise ValueError("El bloque excede el límite de filas o columnas.")
                if len(prepared) + len(values) * width > MAX_EDIT_CELLS:
                    raise ValueError(f"Guarda como máximo {MAX_EDIT_CELLS} celdas por operación.")
                header = _coverage_header(sheet)
                auto_locale = _automatic_number_locale(values, number_locale_hint) if number_locale == 'auto' and not literal else None
                for offset_row, line in enumerate(values):
                    for offset_column, raw in enumerate(line):
                        cell = sheet.cell(row + offset_row, column + offset_column)
                        if isinstance(cell, MergedCell):
                            raise ValueError(f"{name}!{cell.coordinate}: no se puede editar parte de una celda combinada.")
                        if cell.data_type == "f":
                            raise ValueError(f"{name}!{cell.coordinate}: la fórmula está protegida.")
                        identity = (name, cell.coordinate)
                        if identity in touched:
                            raise ValueError("Los bloques de pegado se superponen. Guarda cada celda una vez.")
                        touched.add(identity)
                        is_data = kind == "coverage" and cell.row > header
                        try:
                            if literal:
                                value, number_format = self._copy_value(raw, cell)
                            else:
                                value, number_format = self._convert(raw, cell,
                                    numeric=is_data and 2 <= cell.column <= 8,
                                    dated=is_data and cell.column == 1, epoch=workbook.epoch,
                                    number_locale=number_locale, auto_locale=auto_locale)
                            # La penetración de cobertura y los porcentajes
                            # del motor 5W1H usan puntos porcentuales: 20% es
                            # 20, no 0.20. Un número sin % no se reescala.
                            source_value = raw.get("value") if isinstance(raw, dict) else raw
                            if not literal and is_data and cell.column == 3 and isinstance(source_value, str) and source_value.strip().endswith("%"):
                                value *= 100
                                number_format = '0.0"%"'
                        except ValueError as exc:
                            raise ValueError(f"{name}!{cell.coordinate}: {exc}") from exc
                        prepared.append((cell, value, number_format))
            # Valida todo antes de alterar el libro o crear la copia local.
            for cell, value, number_format in prepared:
                cell.value = value
                if literal and isinstance(value, str):
                    cell.data_type = 's'
                if number_format:
                    cell.number_format = number_format
            target, ref = path, file
            copied_from = None
            if origin == "external":
                with _lock(self.research_local):
                    target = self._unique(self.research_local, path.name)
                    if file_fingerprint(path) != fingerprint:
                        raise InputConflictError("El archivo externo cambió. Vuelve a cargarlo antes de guardar.")
                    self._write(workbook, target)
                    ref, copied_from = f"local/{target.name}", file
            else:
                self._write(workbook, target, fingerprint)
        result = self.workbook(ref, kind=kind, sheet=edits[0]["sheet"])
        result.update(saved_cells=len(prepared), copied_from=copied_from)
        return result

    def add_sheet(self, file: str, *, fingerprint: str, name: str, kind: str = "coverage",
                  source_sheet: str | None = None) -> dict:
        path, origin = self._path(file, kind)
        title = _sheet_title(name)
        with _lock(path):
            if file_fingerprint(path) != fingerprint:
                raise InputConflictError("El archivo cambió. Vuelve a cargarlo antes de agregar una hoja.")
            workbook = _load(path)
            if workbook.security.lockStructure:
                raise ValueError("La estructura del libro está protegida. No se puede agregar una hoja.")
            if len(workbook.worksheets) >= MAX_SHEETS:
                raise ValueError(f"El libro admite como máximo {MAX_SHEETS} hojas.")
            if any(existing.casefold() == title.casefold() for existing in workbook.sheetnames):
                raise ValueError("Ya existe una hoja con ese nombre.")
            if source_sheet is not None:
                if source_sheet not in workbook.sheetnames:
                    raise ValueError("La hoja de origen no existe.")
                copy_sheet = workbook.copy_worksheet(workbook[source_sheet])
                copy_sheet.title = title
            elif kind == "coverage":
                if not re.match(r"^P[0-6]_", title):
                    title = _sheet_title("P0_" + title)
                if any(existing.casefold() == title.casefold() for existing in workbook.sheetnames):
                    raise ValueError("Ya existe una hoja con ese nombre.")
                original_unit = workbook.worksheets[0].cell(1, 2).value
                unit = re.sub(r"^Weighted\s+", "", original_unit, flags=re.I) if isinstance(original_unit, str) else None
                _coverage_sheet(workbook, title, re.sub(r"^P[0-6]_", "", title), unit)
            else:
                raise ValueError("Para 5W1H selecciona una hoja de origen y duplica su estructura.")
            copied_from = None
            if origin == "external":
                with _lock(self.research_local):
                    target = self._unique(self.research_local, path.name)
                    if file_fingerprint(path) != fingerprint:
                        raise InputConflictError("El archivo externo cambió. Vuelve a cargarlo antes de guardar.")
                    self._write(workbook, target)
                    file, copied_from = f"local/{target.name}", file
            else:
                self._write(workbook, path, fingerprint)
        result = self.workbook(file, kind=kind, sheet=title)
        result["copied_from"] = copied_from
        return result
