"""The approved presentation shares unrounded data and one visual error scale."""
from io import BytesIO
import math
import json
from pathlib import Path
import shutil
import subprocess
from types import SimpleNamespace
import unittest

from pptx import Presentation
from pptx.util import Inches
from studio_5w1h_vendor import confidence_design as design


def observation(**changes):
    data = dict(marca="Marca", var_fabricante=-4.583746979661118,
                mat_sep25=2.6902707894574496, lim_inf=-3.130474533704275,
                lim_sup=8.511016112619174, nivel_pen="Marca Pequeña (1% a 10%)",
                pen12m=3.733333333333333, compradores_prom=124.25, error_muestral=.14755025219591095)
    return SimpleNamespace(**{**data, **changes})


class ConfidenceDesignTests(unittest.TestCase):
    def test_visual_cutoffs_use_unrounded_ratio_and_keep_extreme_separate(self):
        for limit, ink, accent, background in design.ERROR_BANDS[:-1]:
            with self.subTest(limit=limit):
                style = design.error_style(limit)
                self.assertEqual((style['ink'], style['accent'], style['background']), (ink, accent, background))
                self.assertNotEqual(design.error_style(math.nextafter(limit, math.inf)), style)
        self.assertFalse(design.error_style(.51)['extreme'])
        self.assertTrue(design.error_style(.510001)['extreme'])
        self.assertEqual(design.error_style(.43)['background'], 'FFD3D6')
        self.assertEqual(design.error_style(.43)['accent'], 'EF1B2D')
        self.assertEqual(design.error_style(.5533323335290206)['background'], '151419')

    def test_missing_data_is_neutral_not_green(self):
        for value in (None, math.nan, math.inf, -1):
            self.assertEqual(design.error_style(value)['background'], 'F1F4F7')
        self.assertEqual(design.percent(None), '—')
        self.assertEqual(design.percent(.14755025219591095, ratio=True), '14.8%')

    def test_view_preserves_values_and_uses_literal_interval_comparison(self):
        row=observation()
        view=design.confidence_view([row])
        self.assertEqual(view['rows'][0]['manufacturer'], row.var_fabricante)
        self.assertEqual(view['rows'][0]['error'], row.error_muestral)
        self.assertEqual(view['rows'][0]['buyers'], 124.25)
        self.assertEqual(design.buyers_label(118.08333333333333), '118')
        self.assertEqual(design.buyers_label(7.25), '7')
        self.assertEqual(design.buyers_label(124.5), '125')
        self.assertEqual(design.buyers_label(None), '—')
        self.assertEqual(view['rows'][0]['penetration'], ['Pequeña', '1% a 10%'])
        self.assertEqual(view['rows'][0]['position'], 'Por debajo')
        self.assertEqual(design.position_label(observation(var_fabricante=row.lim_inf)), 'Dentro')
        self.assertEqual(design.position_label(observation(var_fabricante=10)), 'Por encima')
        self.assertEqual(design.position_label(observation(var_fabricante=math.nan)), '—')
        self.assertEqual(design.penetration_label(observation(nivel_pen='Fuera de rango',pen12m=.28)), ('Fuera de rango','Menor de 1%'))

    def test_axis_contains_each_bound_and_estimate(self):
        rows=[observation(), observation(lim_inf=-78, lim_sup=199, var_fabricante=250)]
        axis=design.interval_axis(rows)
        self.assertLessEqual(axis['min'], -78)
        self.assertGreaterEqual(axis['max'], 250)
        self.assertLessEqual(len(axis['ticks']), 8)
        self.assertEqual(design.interval_axis([observation(lim_inf=math.nan,lim_sup=math.nan,var_fabricante=0,mat_sep25=0)])['min'], 0)

    def test_table_is_native_without_hidden_sample_or_error_range_text(self):
        ppt=Presentation();ppt.slide_width=Inches(13.3333);ppt.slide_height=Inches(7.5)
        slide=ppt.slides.add_slide(ppt.slide_layouts[6])
        rows=[observation(marca='Mondelez'), observation(marca='Tobleron',error_muestral=.5533323335290206), observation(marca='Milka',error_muestral=.15314149191764648)]
        table=design.add_table(slide,rows,['#007A99','CE6611','#238368'],Inches(.33),Inches(1.1),Inches(12.66),5).table
        self.assertEqual(len(table.columns), 8)
        self.assertEqual(table.cell(0,0).text, 'Compradores\nprom. 12 meses')
        self.assertEqual(table.cell(0,1).text, 'Marca')
        self.assertEqual(table.cell(2,0).text, '124')
        self.assertEqual(table.cell(1,7).text, 'Error muestral')
        for i, expected in enumerate(('14.8%','55.3%','15.3%'),2):
            self.assertEqual(table.cell(i,7).text, expected)
        all_text=' '.join(cell.text for row in table.rows for cell in row.cells)
        self.assertNotIn('Penetración 12 meses',all_text)
        self.assertNotIn('43.0%',all_text)
        self.assertNotIn('>51',all_text)
        self.assertNotIn('Muestra de color',all_text)
        self.assertEqual(str(table.cell(3,7).fill.fore_color.rgb),'151419')
        self.assertEqual(str(table.cell(2,7).fill.fore_color.rgb),'FCE8D8')
        self.assertTrue(table.cell(2,6).text_frame.paragraphs[1].runs[0].font.bold)
        for row in rows:
            self.assertTrue(any(s.name==row.marca+': error señal vertical' for s in slide.shapes))
            self.assertTrue(any(s.name==row.marca+': error señal horizontal' for s in slide.shapes))
        self.assertFalse(any('Referencia 0%' in s.text for s in slide.shapes if s.has_text_frame))
        buffer=BytesIO();ppt.save(buffer);buffer.seek(0)
        reopened=Presentation(buffer)
        self.assertEqual(sum(s.has_table for s in reopened.slides[0].shapes),1)

    def test_reference_footer_retains_source_and_reserves_readable_space(self):
        layout=design.slide_layout(7.5,.72)
        self.assertEqual(layout['capacity'],3)
        ppt=Presentation();ppt.slide_width=Inches(13.3333);ppt.slide_height=Inches(7.5)
        slide=ppt.slides.add_slide(ppt.slide_layouts[6])
        rows=[observation() for _ in range(layout['capacity'])]
        shape=design.add_table(slide,rows,['183143']*len(rows),Inches(.33),Inches(.72),Inches(12.66),layout['table_height'])
        footer=design.add_buyer_reference(slide,Inches(.33),Inches(layout['reference_top']),Inches(12.66))
        self.assertLess(shape.top+shape.height+Inches(.5),Inches(layout['comment_top']))
        self.assertLess(Inches(layout['comment_top']+.4),footer.top)
        self.assertLess(footer.top+footer.height,Inches(7.35))
        self.assertGreater(footer.top,Inches(5.8))
        self.assertAlmostEqual(footer.left+footer.width/2,Inches(.33+12.66/2),delta=2)
        self.assertLess(footer.width,Inches(9.2))
        self.assertEqual((len(footer.table.rows),len(footer.table.columns)),(7,3))
        self.assertEqual(footer.table.cell(2,1).text,'Entre 70 y 100 casos')
        self.assertEqual(footer.table.cell(3,1).text,'Entre 100 y 200 casos')
        self.assertIn('Compra por ocasión',footer.table.cell(3,2).text)
        self.assertIn('Variaciones de volumen',footer.table.cell(5,2).text)
        self.assertEqual(footer.table.cell(6,1).text,'Más de 400 casos')
        self.assertTrue(all(run.font.size.pt>=9 for row in footer.table.rows for cell in row.cells
                            for p in cell.text_frame.paragraphs for run in p.runs))
        buf=BytesIO();ppt.save(buf);buf.seek(0)
        self.assertEqual(sum(s.has_table for s in Presentation(buf).slides[0].shapes),2)

    def test_many_brands_require_pages_instead_of_tiny_text(self):
        capacity=design.page_capacity(5)
        self.assertGreaterEqual(capacity,5)
        self.assertLessEqual(capacity,8)
        ppt=Presentation();slide=ppt.slides.add_slide(ppt.slide_layouts[6])
        with self.assertRaisesRegex(ValueError,'varias diapositivas'):
            design.add_table(slide,[observation() for _ in range(capacity+1)],['183143']*(capacity+1),0,0,Inches(12),5)

    def test_browser_component_uses_shared_styles_and_hides_internal_ranges(self):
        node=shutil.which('node')
        if not node:
            self.skipTest('Node no disponible')
        view=design.confidence_view([observation(marca='Media'),observation(marca='Rojo',error_muestral=.43),
                                     observation(marca='Extremo',error_muestral=.5533323335290206),
                                     observation(marca='Sin dato',error_muestral=None)])
        script=r'''
const fs=require('node:fs'),vm=require('node:vm'),assert=require('node:assert/strict');
class Node {
 constructor(tag){this.tag=tag;this.children=[];this.attrs={};this.className='';this.textContent='';this.style={setProperty(k,v){this[k]=v;}};}
 append(...items){this.children.push(...items);}
 setAttribute(k,v){this.attrs[k]=String(v);}
}
const document={createElement:t=>new Node(t),createElementNS:(ns,t)=>new Node(t),createTextNode:t=>({textContent:t,children:[]})};
const source=fs.readFileSync('explorer_web/research.js','utf8');
const prefix=source.slice(source.indexOf('  const $ ='),source.indexOf('  const normalize ='));
const a=source.indexOf('  function confidenceContent('),b=source.indexOf('  function renderAnalysis(',a);
assert(a>=0&&b>a);
const context={document,finite:v=>typeof v==='number'&&Number.isFinite(v)};vm.createContext(context);
vm.runInContext(prefix+source.slice(a,b)+';this.render=confidenceContent;',context);
const view=JSON.parse(process.argv[1]);
const root=context.render({title:'Intervalos',confidence_view:view});
const all=n=>[n,...n.children.flatMap(all)],text=n=>n.textContent+n.children.map(text).join('');
const signals=all(root).filter(n=>String(n.className).split(' ').includes('confidence-signal'));
assert.equal(signals.length,4);
assert.equal(signals[0].style['--error-background'],'#FCE8D8');
assert.equal(signals[1].style['--error-background'],'#FFD3D6');
assert.equal(signals[1].style['--error-accent'],'#EF1B2D');
assert.equal(signals[2].style['--error-background'],'#151419');
assert.equal(signals[3].style['--error-background'],'#F1F4F7');
assert(signals.every(n=>!text(n).includes('>')));
assert.deepEqual(signals.map(n=>text(n.children.find(c=>c.tag==='strong'))),['14.8%','43.0%','55.3%','—']);
assert(!text(root).includes('Muestra de color'));assert(!text(root).includes('Penetración 12 meses'));assert(!text(root).includes('Referencia 0%'));
assert(text(root).includes('Compradores · prom. 12 meses'));
assert.equal(all(root).filter(n=>n.className==='confidence-buyers')[0].textContent,'124');
'''
        result=subprocess.run([node,'-e',script,json.dumps(view)],cwd=Path(__file__).resolve().parents[1],
                              capture_output=True,text=True,encoding='utf-8')
        self.assertEqual(result.returncode,0,result.stdout+result.stderr)


if __name__ == '__main__':
    unittest.main()
