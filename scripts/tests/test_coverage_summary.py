"""Resumen de todas las marcas y slides elegidos de forma independiente."""
import contextlib
import io
import json
import math
from pathlib import Path
import shutil
import subprocess
import tempfile
import time
import unittest

from openpyxl import Workbook, load_workbook
import pipeline_explorer as explorer


class CoverageSummaryTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.core = explorer.engine()

    def setUp(self):
        (explorer.ROOT / 'outputs').mkdir(exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=explorer.ROOT / 'outputs')
        self.root = Path(self.temp.name)
        self.source = self.root / '52_CARB_Resumen.xlsx'
        book = Workbook(); book.remove(book.active)
        for name, buyers in [('P1_BaseBaja', 199.6), ('P2_Limite', 200), ('P3_Grande', 900)]:
            sheet = book.create_sheet(name)
            sheet.append(['Table', 'Weighted R_VOL1', 'PENET', 'VO1_BUY', 'VO1_DAY', 'FREQ', 'BUYERS', 'SELL-IN'])
            for i in range(36):
                date = self.core.pd.Timestamp(2023, 1, 1) + self.core.pd.DateOffset(months=i)
                volume = 1200 + i * 17 + 180 * math.sin(i * .7)
                sheet.append([date.to_pydatetime(), volume, 20 + i * .25, 6, 2, 3, buyers, volume * 1.7 + math.cos(i) * 75])
        book.save(self.source); book.close()
        shutil.copy2(explorer.ROOT / 'Modelo_PPT.pptx', self.root / 'Modelo_PPT.pptx')
        self.state = explorer.ExplorerState(self.root)
        with contextlib.redirect_stdout(io.StringIO()):
            self.analysis = self.state.analyze({'file': self.source.name, 'animate_trend_pipeline': False})
        self.selections = {'P1_BaseBaja': 0, 'P2_Limite': 2, 'P3_Grande': 4}

    def tearDown(self):
        self.temp.cleanup()

    def request(self, slides=None):
        payload = {'analysis_id': self.analysis['id'], 'selections': self.selections}
        if slides is not None: payload['ppt_slide_sheets'] = slides
        return payload

    def export(self, payload):
        job = self.state.export(payload)
        deadline = time.monotonic() + 90
        while self.state.job(job['id'])['status'] == 'running' and time.monotonic() < deadline:
            time.sleep(.1)
        job = self.state.job(job['id'])
        self.assertEqual(job['status'], 'complete', job.get('error') or job.get('log'))
        return job, {out['label']: out['path'] for out in job['outputs']}

    def test_mask_rejects_unselected_duplicates_and_wrong_types_before_export(self):
        for value in [None, 'P2_Limite', [False], ['No existe'], ['P2_Limite', 'P2_Limite']]:
            with self.subTest(value=value), self.assertRaises(ValueError):
                self.state.save_selection(self.request() | {'ppt_slide_sheets': value})
        self.assertEqual(self.state.jobs, {})
        self.assertEqual(self.state.reviewed_slide_sheets(self.request(), self.selections), list(self.selections))
        self.assertEqual(self.state.reviewed_slide_sheets(self.request([]), self.selections), [])

    def test_selection_recovers_empty_mask_without_changing_pipelines_or_extra_preferences(self):
        self.state.save_selection(self.request([]) | {'ppt_extras': {'P1_BaseBaja': ['frequency']}})
        restored = explorer.ExplorerState(self.root).analyze({'file': self.source.name, 'animate_trend_pipeline': False})['saved_selection']
        self.assertEqual(restored['ppt_slide_sheets'], [])
        self.assertEqual(restored['selections'], self.selections)
        self.assertEqual(restored['ppt_extras'], {'P1_BaseBaja': ['frequency']})

    def check_summary(self, outputs, expected):
        ppt = self.core.Presentation(outputs['presentation'])
        summary_rows = [row for slide in ppt.slides for shape in slide.shapes if shape.has_table for row in shape.table.rows
                        if row.cells[0].text in expected]
        self.assertEqual({row.cells[0].text for row in summary_rows}, set(expected))
        for row in summary_rows:
            name = next(name for name in self.selections if name.endswith(row.cells[0].text))
            self.assertEqual(row.cells[1].text, str(self.selections[name]))
        bank = load_workbook(outputs['bank'], data_only=True)
        rows = list(bank.active.values); bank.close()
        brand_column = rows[0].index('Fabricante/Marca')
        data = []
        for row in rows[1:]:
            if not any(value is not None for value in row): break
            data.append(row)
        self.assertEqual({row[brand_column] for row in data}, set(expected))
        self.assertEqual(len(data), len(expected))
        template = load_workbook(outputs['template'], read_only=True)
        self.assertTrue(set(self.selections).issubset(template.sheetnames)); template.close()
        return ppt

    def test_real_export_retains_all_summary_rows_but_only_selected_brand_slides_and_extras(self):
        job, outputs = self.export(self.request(['P2_Limite']) | {'ppt_extras': {'P1_BaseBaja': ['frequency'], 'P2_Limite': ['frequency']}})
        ppt = self.check_summary(outputs, ['BaseBaja', 'Limite', 'Grande'])
        brand_text = '\n'.join(shape.text for slide in list(ppt.slides)[7:-1] for shape in slide.shapes if shape.has_text_frame)
        self.assertIn('Limite', brand_text)
        self.assertNotIn('BaseBaja', brand_text)
        self.assertNotIn('Grande', brand_text)
        self.assertEqual(sum(slide.name.startswith('KPI:frequency') for slide in ppt.slides), 1)
        config = json.loads((Path(job['directory']) / '_sistema' / 'selection.json').read_text(encoding='utf-8'))
        self.assertEqual(config['ppt_slide_sheets'], ['P2_Limite'])
        self.assertEqual(config['selections'], self.selections)

    def test_real_export_can_generate_summary_only_without_any_brand_slides(self):
        _, outputs = self.export(self.request([]))
        ppt = self.check_summary(outputs, ['BaseBaja', 'Limite', 'Grande'])
        self.assertEqual(len(ppt.slides), 8, 'Se conservan introducción, resumen y cierre, sin slides de marca')

    def test_web_threshold_masks_and_current_pipeline_values_remain_consistent(self):
        node = shutil.which('node')
        if not node: self.skipTest('Node no disponible')
        script = r'''
const fs=require('node:fs'),vm=require('node:vm'),assert=require('node:assert/strict');
class Element {
 constructor(tag){this.tag=tag;this.children=[];this.dataset={};this.events={};this.value='';this.textContent='';}
 append(...nodes){this.children.push(...nodes);} replaceChildren(...nodes){this.children=nodes;}
 setAttribute(k,v){this[k]=v;} addEventListener(k,v){this.events[k]=v;}
 querySelector(selector){return all(this).find(n=>selector==='details'?n.tag==='details':selector.startsWith('.')?n.className?.split(' ').includes(selector.slice(1)):n.tag==='input'&&n.type==='search');}
 querySelectorAll(){return all(this).filter(n=>['button','input','select'].includes(n.tag));}
}
const all=n=>[n,...n.children.flatMap(all)],context={window:{},document:{createElement:tag=>new Element(tag)}};
vm.createContext(context);vm.runInContext(fs.readFileSync('explorer_web/coverage-summary.js','utf8'),context);
const module=context.window.CoverageSummary;
const sheets=[{name:'A',brand:'Alta',buyers:200},{name:'B',brand:'Baja',buyers:199.6},{name:'C',brand:'Sin dato',buyers:null}];
for(const s of sheets){s.dates=['2025-12-01'];s.pipelines=[{pipeline:0,available:true,coverage:80},{pipeline:1,available:true,coverage:90,client_variation:-5,wp_variation:4}];}
sheets.push({name:'D',error:'Fecha inválida'});
const state={analysis:{sheets},included:{A:true,B:true,C:true,D:false},selections:{A:1,B:0,C:0},pptSlides:{},pptExtras:{},busy:false};
module.preset(state,'eligible');assert(state.included.A);assert(!state.included.B);assert(!state.included.C);assert(!state.included.D);
assert.equal(module.counts(state).summary,1);assert.equal(module.counts(state).slides,1);
module.preset(state,'summary');assert.equal(module.counts(state).summary,1);assert.equal(module.counts(state).slides,0);
module.preset(state,'all');assert.equal(module.counts(state).summary,3);assert.equal(module.counts(state).slides,3);
module.preset(state,'slides-eligible');assert.equal(module.counts(state).summary,3);assert.equal(module.counts(state).slides,1);
assert(state.included.B,'La base baja se conserva en el resumen');
module.preset(state,'all');
const root=new Element('section');let changes=0,inspected;
const helpers={fmt:(v,d=1,s='')=>v==null?'—':v.toFixed(d)+s,month:v=>v,changed:()=>changes++,inspect:n=>inspected=n};
module.render(root,state,helpers);
const rows=all(root).find(n=>n.tag==='tbody').children;
assert(rows[0].children.some(n=>n.textContent==='90.0%'),'Cobertura del pipeline elegido P1');
assert(rows[0].children.some(n=>n.textContent==='-5.0%'),'Variación del mismo pipeline');
assert.equal(all(rows[0]).find(n=>n.tag==='select').value,'1');
const slideCheck=rows[0].children[1].children[0];slideCheck.checked=false;slideCheck.events.change();
assert.equal(state.pptSlides.A,false);assert.equal(state.included.A,true);assert.equal(changes,1);
all(rows[0]).find(n=>n.tag==='button').events.click();assert.equal(inspected,'A');
const search=all(root).find(n=>n.type==='search');search.value='baja';search.events.input();assert(rows[0].hidden);assert(!rows[1].hidden);
root.querySelector('.coverage-summary-scroll').scrollTop=150;root.querySelector('.coverage-summary-scroll').scrollLeft=200;
module.render(root,state,helpers);assert.equal(root.querySelector('.coverage-summary-scroll').scrollTop,150);assert.equal(root.querySelector('.coverage-summary-scroll').scrollLeft,200);
const source=fs.readFileSync('explorer_web/app.js','utf8');
const start=source.indexOf('function selectedSheets()'),end=source.indexOf('function draftKey()',start);
Object.assign(context,{state,MosaicComparison:{eligible:()=>[]}});vm.runInContext(source.slice(start,end)+';this.payload=payload;',context);
const request=context.payload();assert.equal(request.selections.A,1);assert(!request.ppt_slide_sheets.includes('A'));assert(request.ppt_slide_sheets.includes('B'));
state.busy=true;module.render(root,state,helpers);assert(all(root).find(n=>n.tag==='select').disabled);
state.busy=false;module.setBusy(root,false);assert(!all(root).find(n=>n.tag==='select').disabled,'Se habilita al terminar la carga');
assert(all(root).find(n=>n.tag==='select'&&n['aria-label']==='Pipeline de D').disabled,'La hoja inválida sigue deshabilitada');
'''
        result = subprocess.run([node, '-e', script], cwd=explorer.ROOT, capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)


if __name__ == '__main__': unittest.main()
