"""Compatibilidad de auxiliares y contrato de los gráficos adicionales del PPT."""
import json
import base64
import io
import shutil
import subprocess
from pathlib import Path
import tempfile
import types
import unittest
from unittest import mock

import pipeline_explorer as explorer
from delivery_files import auxiliary, organize_auxiliaries


class DeliveryExtrasTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.core = explorer.engine()

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=explorer.ROOT / "outputs")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.directory = self.root / "Output" / "Entrega"
        self.directory.mkdir(parents=True)

    def test_old_delivery_is_migrated_and_recognized_again(self):
        output = self.directory / "Template.xlsx"
        output.write_bytes(b"existing-deliverable")
        job = {"id": "old", "status": "complete", "directory": str(self.directory),
               "started_at": "2026-10-13T12:00:00-06:00", "file": "52_CARB_Test.xlsx",
               "outputs": [{"label": "template", "path": str(output), "name": output.name}]}
        for name, value in (("job.json", job), ("selection.json", {"selections": {"A": 2}}),
                            ("result.json", {"directory": str(self.directory)})):
            (self.directory / name).write_text(json.dumps(value), encoding="utf-8")
        (self.directory / "generation.log").write_text("Registro conservado", encoding="utf-8")
        for _ in range(2):
            state = explorer.ExplorerState(self.root)
            self.assertEqual(state.job("old")["log"], "Registro conservado")
            self.assertEqual(state.output_path(state.job("old"), "selection"), self.directory / "_sistema" / "selection.json")
        self.assertEqual(output.read_bytes(), b"existing-deliverable")
        self.assertEqual({path.name for path in self.directory.iterdir()}, {"_sistema", "Template.xlsx"})

    def test_old_read_only_auxiliary_still_reads_in_place(self):
        old = self.directory / "selection.json"
        old.write_text("{}", encoding="utf-8")
        with mock.patch.object(Path, "replace", side_effect=PermissionError):
            self.assertEqual(organize_auxiliaries(self.directory, [old.name]), {})
        self.assertEqual(auxiliary(self.directory, old.name), old)

    def test_old_research_delivery_retargets_auxiliary_links_and_work_copy(self):
        outputs = []
        for label, name in (("Excel", "Informe.xlsx"), ("PPT", "Informe.pptx"),
                            ("Provenance", "investigacion_5w1h.json"), ("Registro", "generacion_5w1h.log")):
            path = self.directory / name
            path.write_bytes(b"preserved")
            outputs.append({"label": label, "name": name, "path": str(path)})
        copied = self.directory / "_input"
        copied.mkdir()
        (copied / "Entrada.xlsx").write_bytes(b"original copy")
        (copied / "request.json").write_text(json.dumps({"source": str(copied / "Entrada.xlsx")}), encoding="utf-8")
        (self.directory / "_result.json").write_text("{}", encoding="utf-8")
        job = {"id":"research-old", "kind":"research", "file":"local/52_CARB_Test.xlsx",
               "fingerprint":"f"*64, "status":"complete", "directory":str(self.directory),
               "started_at":"2026-10-13T12:00:00-06:00", "outputs":outputs}
        (self.directory / "research-job.json").write_text(json.dumps(job), encoding="utf-8")
        for _ in range(2):
            state = explorer.ExplorerState(self.root)
            recovered = state.workflow.research_jobs[job["id"]]
            self.assertTrue(all(Path(out["path"]).is_file() for out in recovered["outputs"]))
        self.assertEqual({path.name for path in self.directory.iterdir()}, {"Informe.xlsx", "Informe.pptx", "_sistema"})
        request = json.loads((self.directory / "_sistema" / "_input" / "request.json").read_text(encoding="utf-8"))
        self.assertEqual(Path(request["source"]).read_bytes(), b"original copy")

    def test_internal_path_escape_is_rejected(self):
        with self.assertRaises(ValueError):
            auxiliary(self.directory, "../outside.json", write=True)
        outside = self.root / "outside"
        outside.mkdir()
        try:
            (self.directory / "_sistema").symlink_to(outside, target_is_directory=True)
        except OSError:
            self.skipTest("Esta sesión no permite crear enlaces simbólicos")
        with self.assertRaises(ValueError):
            auxiliary(self.directory, "job.json", write=True)

    def test_extras_reject_unknown_or_excluded_sheet_and_duplicate_keys(self):
        for value in (None, [], {"Excluded": ["frequency"]}, {"A": ["unknown"]},
                      {"A": ["frequency", "frequency"]}, {"A": "frequency"}):
            with self.subTest(value=value), self.assertRaises(ValueError):
                self.core.validate_ppt_extra_charts(value, ["A"])
        self.assertEqual(self.core.validate_ppt_extra_charts({}, ["A"]), {})

    def test_every_kpi_is_editable_and_default_adds_no_slides(self):
        core = self.core
        ppt = core.Presentation(explorer.ROOT / "Modelo_PPT.pptx")
        builder = types.SimpleNamespace(ppt=ppt, lang_index=2)
        frame = core.pd.DataFrame({core.COL_DATA:core.pd.date_range("2024-01-01", periods=14, freq="MS")})
        for column, _ in core.PPT_EXTRA_KPIS.values():
            frame[column] = list(range(1, 15))
        frame.loc[0, core.COL_FREQ] = float("nan")
        count = len(ppt.slides)
        self.assertEqual(core.add_extra_kpi_slides(builder, frame, [], "Marca", "kg"), 0)
        self.assertEqual(len(ppt.slides), count)
        self.assertEqual(core.add_extra_kpi_slides(builder, frame, list(core.PPT_EXTRA_KPIS), "Marca", "kg"), 5)
        ppt.save(self.root / "proof.pptx")
        reloaded = core.Presentation(self.root / "proof.pptx")
        for slide in list(reloaded.slides)[count:]:
            chart = next(shape.chart for shape in slide.shapes if shape.has_chart)
            self.assertEqual(chart.series[0].values[-1], 14)
            texts = " ".join(shape.text for shape in slide.shapes if shape.has_text_frame)
            self.assertIn("Promedio 12 meses: 8" if slide.name.startswith("KPI:buyers_series") else "Promedio 12 meses: 8.50", texts)
            self.assertIn("Promedio total", texts)
            self.assertTrue(all(shape.left >= 0 and shape.top >= 0 and shape.left+shape.width <= ppt.slide_width and shape.top+shape.height <= ppt.slide_height for shape in slide.shapes))

    def test_web_capture_is_validated_and_inserted_without_native_chart(self):
        from PIL import Image
        core = self.core
        buffer = io.BytesIO()
        Image.new('RGB', (2400,1350), '#FFFFFF').save(buffer,format='PNG')
        encoded = 'data:image/png;base64,'+base64.b64encode(buffer.getvalue()).decode()
        extras = {'A':['annual_mat']}
        images = explorer.validate_ppt_images({'A':{'annual_mat':encoded}},extras)
        for value in ({}, {'A':{}}, {'B':{'annual_mat':encoded}}, {'A':{'frequency':encoded}},
                      {'A':{'annual_mat':'C:/private/file.png'}}, {'A':{'annual_mat':'data:image/png;base64,invalid'}}):
            with self.subTest(value=str(value)[:70]), self.assertRaises(ValueError):
                explorer.validate_ppt_images(value,extras)
        target = self.root/'capture.png';target.write_bytes(images['A']['annual_mat'])
        ppt = core.Presentation(explorer.ROOT/'Modelo_PPT.pptx')
        builder = types.SimpleNamespace(ppt=ppt,lang_index=2)
        self.assertEqual(core.add_extra_kpi_slides(builder,None,['annual_mat'],'Marca','kg',images={'annual_mat':str(target)}),1)
        ppt.save(self.root/'captured.pptx')
        slide = core.Presentation(self.root/'captured.pptx').slides[-1]
        self.assertFalse(any(shape.has_chart or shape.has_table for shape in slide.shapes))
        picture = next(shape for shape in slide.shapes if shape.shape_type == 13)
        self.assertEqual(picture.image.blob,images['A']['annual_mat'])
        self.assertAlmostEqual(picture.width/picture.height,16/9,places=5)
        self.assertTrue(picture.left>=0 and picture.top>=0 and picture.left+picture.width<=ppt.slide_width and picture.top+picture.height<=ppt.slide_height)

    def test_annual_comparison_preserves_pipeline_yoy_coverage_and_editable_data(self):
        from io import BytesIO
        from openpyxl import load_workbook
        core = self.core
        frame = core.pd.DataFrame({core.COL_DATA:core.pd.date_range("2023-01-01", periods=36, freq="MS"),
            core.COL_SELL_OUT:list(range(100,136)), core.COL_SELL_IN:list(range(200,236))})
        coverage = core.compute_coverage_dataframe(frame, "Peru", "Relativa", True)
        p = 3
        selected = coverage[f"P{p}"]
        rows = core.annual_mat_rows(frame, p, selected)
        self.assertIsNone(rows[0]["sell_in"])
        self.assertIsNone(rows[1]["sell_in_yoy"])
        self.assertEqual(rows[-1]["sell_in"], sum(frame[core.COL_SELL_IN].iloc[-15:-3]))
        self.assertEqual(rows[-1]["coverage"], selected.iloc[-1])
        self.assertAlmostEqual(rows[-1]["sell_in_yoy"], (rows[-1]["sell_in"]/rows[-2]["sell_in"]-1)*100)
        self.assertEqual(core.validate_ppt_extra_charts({"A":["annual_mat"]}, ["A"]), {"A":["annual_mat"]})
        ppt = core.Presentation(explorer.ROOT / "Modelo_PPT.pptx")
        builder = types.SimpleNamespace(ppt=ppt, lang_index=2)
        self.assertEqual(core.add_extra_kpi_slides(builder, frame, ["annual_mat"], "Marca", "kg", p, selected), 1)
        ppt.save(self.root / "annual.pptx")
        chart = next(shape.chart for shape in core.Presentation(self.root / "annual.pptx").slides[-1].shapes if shape.has_chart)
        self.assertEqual(len(chart.plots), 2)
        self.assertTrue(chart.plots[0]._element.tag.endswith('}barChart'))
        self.assertTrue(chart.plots[1]._element.tag.endswith('}lineChart'))
        self.assertEqual(chart.series[0].values, tuple(row["sell_in"] for row in rows[1:]))
        self.assertEqual(chart.series[1].values, tuple(row["worldpanel"] for row in rows[1:]))
        self.assertEqual(chart.series[2].values, tuple(row["coverage"] for row in rows[1:]))
        axes = chart._chartSpace.chart.plotArea.xpath('c:valAx')
        self.assertEqual(len(axes), 2)
        self.assertEqual(axes[-1].xpath('c:axPos')[0].get('val'), 'r')
        self.assertEqual(axes[-1].xpath('c:numFmt')[0].get('formatCode'), '0"%"')
        table = next(shape.table for shape in core.Presentation(self.root / "annual.pptx").slides[-1].shapes if shape.has_table)
        self.assertEqual(table.cell(1,2).text, f"{rows[-1]['sell_in_yoy']:+.1f}%")
        self.assertEqual(table.cell(1,1).text, '—')
        embedded = load_workbook(BytesIO(chart.part.chart_workbook.xlsx_part.blob), data_only=True)
        self.assertEqual(embedded.active.cell(3, 4).value, rows[-1]["coverage"])
        self.assertEqual(core.add_extra_kpi_slides(builder, frame, ["annual_mat"], "Marca", "kg", p,
            core.pd.Series([float('nan')]*len(frame))), 0)


    def test_bulk_ppt_extras_use_each_included_sheets_data_and_pipeline(self):
        node = shutil.which("node")
        if not node:
            self.skipTest("Node no está disponible para comprobar la selección web")
        script = r'''
const fs=require('node:fs'),vm=require('node:vm'),assert=require('node:assert/strict');
const context={window:{}};vm.createContext(context);
vm.runInContext(fs.readFileSync('explorer_web/dashboard.js','utf8'),context);
const dashboard=context.window.CoverageDashboard;
function sheet(name,pipeline,count=12,annual=true) {
  const values=[1,2];
  return {name,dates:['2025-01-01','2026-01-01'],frequency:values,
    pipelines:{[pipeline]:{available:true,sell_in:values}},
    template:{correlations:[{label:'Actual',pipelines:Array.from({length:7},(_,p)=>p===pipeline ? .9 : null)}]},
    dashboard:{kpis:[{key:'frequency',label:'Frecuencia',total_count:count}],
      mat_periods:annual ? [{date:'2026-01-01',sell_in:{[pipeline]:2},sell_in_yoy:{[pipeline]:10},
        worldpanel:2,worldpanel_yoy:10,coverage:{[pipeline]:100}}] : []}};
}
const sheets=[sheet('A',0),sheet('B',4),sheet('Sin frecuencia',2,0),
  sheet('Sin anual',1,12,false),sheet('Excluida',3),{name:'Inválida',error:'Sin datos'}];
const state={analysis:{sheets,chart_colors:{}},selections:{A:0,B:4,'Sin frecuencia':2,'Sin anual':1,Excluida:3},
  included:Object.fromEntries(sheets.map(s=>[s.name,s.name!=='Excluida'])),
  pptExtras:{A:['buyers_series'],B:['annual_mat'],Excluida:['purchase_average']},busy:false};
assert.equal(dashboard.pptAvailable(sheets[1],4,'annual_mat',{}),true);
assert.equal(dashboard.pptAvailable(sheets[1],0,'annual_mat',{}),false,'Usar el pipeline propio');
assert.equal(dashboard.pptAvailable(sheets[2],2,'frequency',{}),false);
let saved=0,synced=0,messages=[];
Object.assign(context,{state,CoverageDashboard:dashboard,pptExtraTitles:{frequency:'Frecuencia',annual_mat:'Anual'},
  saveDraft(){saved++;},status(message){messages.push(message);},
  document:{querySelectorAll(){return [];}}});
const source=fs.readFileSync('explorer_web/app.js','utf8');
vm.runInContext(source.slice(source.indexOf('function pptExtraTargets('),source.indexOf('function renderDashboard()')),context);
context.syncPptExtraControls=()=>synced++;
context.includePptExtraInAll('frequency');
assert.deepEqual([...state.pptExtras.A],['buyers_series','frequency']);
assert.deepEqual([...state.pptExtras.B],['annual_mat','frequency']);
assert.deepEqual([...state.pptExtras['Sin anual']],['frequency']);
assert.equal(state.pptExtras['Sin frecuencia'],undefined);
assert.deepEqual(state.pptExtras.Excluida,['purchase_average']);
assert.equal(state.pptExtras['Inválida'],undefined);
context.includePptExtraInAll('frequency');
assert.equal(state.pptExtras.A.filter(key=>key==='frequency').length,1,'Sin duplicados');
context.includePptExtraInAll('annual_mat');
assert.ok(state.pptExtras.A.includes('annual_mat'));
assert.ok(state.pptExtras['Sin frecuencia'].includes('annual_mat'));
assert.ok(!state.pptExtras['Sin anual'].includes('annual_mat'));
const snapshot=JSON.stringify(state.pptExtras),before=saved;
state.busy=true;context.includePptExtraInAll('frequency');state.busy=false;
state.job={status:'running'};context.includePptExtraInAll('frequency');state.job=null;
context.includePptExtraInAll('coverage');
assert.equal(saved,before);assert.equal(JSON.stringify(state.pptExtras),snapshot);
assert.equal(saved,3);assert.equal(synced,3);
assert.ok(messages[0].includes('3 hoja(s)'));assert.ok(messages[0].includes('2 hoja(s) sin datos'));
assert.equal(state.selections.B,4,'No modificar pipelines ni incluir hojas excluidas');
context.pptExtraTitles.trend_correlations='Correlaciones MAT';
context.includePptExtraInAll('trend_correlations');
assert.ok(state.pptExtras.A.includes('trend_correlations'));
assert.ok(state.pptExtras.B.includes('trend_correlations'));
assert.ok(!state.pptExtras.Excluida.includes('trend_correlations'));
assert.equal(saved,4);
state.active='A';state.pptExtraScope='sheet';context.setTrendCorrelations(false);
assert.ok(!state.pptExtras.A.includes('trend_correlations'));
assert.ok(state.pptExtras.B.includes('trend_correlations'),'El control local solo cambia la hoja actual');
context.setTrendCorrelations(true);
assert.ok(state.pptExtras.A.includes('trend_correlations'));
state.pptExtras.Excluida.push('trend_correlations');
context.removePptExtraFromAll('trend_correlations');
assert.ok(!state.pptExtras.A.includes('trend_correlations'));
assert.ok(!state.pptExtras.B.includes('trend_correlations'));
assert.ok(state.pptExtras.Excluida.includes('trend_correlations'),'No cambiar hojas excluidas');
assert.ok(state.pptExtras.A.includes('buyers_series'),'Conservar los demás extras');
assert.ok(state.pptExtras.B.includes('annual_mat'));
const afterOff=JSON.stringify(state.pptExtras),savedAfterOff=saved;
state.busy=true;context.setTrendCorrelations(true);context.removePptExtraFromAll('trend_correlations');state.busy=false;
assert.equal(saved,savedAfterOff);assert.equal(JSON.stringify(state.pptExtras),afterOff);
state.pptExtraScope='all';context.setPptExtra('annual_mat',false);
assert.ok(!state.pptExtras.A.includes('annual_mat'));
assert.ok(!state.pptExtras.B.includes('annual_mat'));
assert.ok(!state.pptExtras['Sin frecuencia'].includes('annual_mat'));
assert.ok(state.pptExtras.A.includes('buyers_series'));
assert.ok(state.pptExtras.B.includes('frequency'));
context.setPptExtra('annual_mat',true);
assert.ok(state.pptExtras.A.includes('annual_mat'));
assert.ok(state.pptExtras.B.includes('annual_mat'));
assert.deepEqual(state.pptExtras.Excluida,['purchase_average','trend_correlations']);
'''
        subprocess.run([node, "-e", script], cwd=explorer.ROOT, check=True, capture_output=True, text=True)

    def test_trend_correlations_image_preserves_existing_slide_elements(self):
        core = self.core
        ppt = core.Presentation(explorer.ROOT / "Modelo_PPT.pptx")
        slide = ppt.slides.add_slide(ppt.slide_layouts[core.PPT_LAYOUT_INDEX])
        frame = core.pd.DataFrame({core.COL_DATA:core.pd.date_range("2023-01-01", periods=36, freq="MS").strftime("%m-%y"),
                                  core.COL_SELL_IN:list(range(10,46)),core.COL_SELL_OUT:list(range(20,56))})
        core.generar_grafico_tendencia(slide,"Prueba",2,frame,2,core.build_labels(2,"Prueba","12-25"),animate_pipeline=False)
        before = [(shape.shape_id,shape.left,shape.top,shape.width,shape.height,shape._element.xml.encode('utf-8'))
                  for shape in slide.shapes]
        count = len(slide.shapes)
        self.assertIsNone(core.add_trend_correlations_table(slide,ppt,2,None))
        self.assertEqual(len(slide.shapes),count)
        matrix=[[.123456,None,-.8,.5,.6,.7,.8],[None]*7,[.2]*7]
        picture = core.add_trend_correlations_table(slide,ppt,2,matrix)
        self.assertEqual(picture.image.content_type,"image/png")
        description=picture._element.nvPicPr.cNvPr.get('descr')
        self.assertIn('12.3% | — | -80.0%',description)
        self.assertNotIn('Anterior',description)
        self.assertEqual(len(description.splitlines()),3)
        self.assertEqual(before,[(shape.shape_id,shape.left,shape.top,shape.width,shape.height,shape._element.xml.encode('utf-8'))
                                 for shape in list(slide.shapes)[:count]])
        shape = slide.shapes[-1]
        self.assertLessEqual(shape.top+shape.height,ppt.slide_height-core.Inches(.45))
        self.assertGreaterEqual(shape.left,core.Inches(2))
        self.assertLessEqual(shape.left+shape.width,ppt.slide_width)
        with self.assertRaises(ValueError):
            explorer.validate_ppt_images({}, {"A":["trend_correlations"]})
        captured=io.BytesIO();core.Image.new('RGB',(2400,240),'white').save(captured,format='PNG')
        encoded='data:image/png;base64,'+base64.b64encode(captured.getvalue()).decode()
        validated=explorer.validate_ppt_images({'A':{'trend_correlations':encoded}},{'A':['trend_correlations']})
        self.assertTrue(validated['A']['trend_correlations'].startswith(b'\x89PNG'))
        with self.assertRaises(ValueError):
            explorer.validate_ppt_images({'A':{'annual_mat':encoded}},{'A':['annual_mat']})
        with mock.patch.object(core.plt,'subplots',side_effect=AssertionError('No debe redibujar la captura web')):
            same=core.add_trend_correlations_table(slide,ppt,2,matrix,image=io.BytesIO(validated['A']['trend_correlations']))
        self.assertEqual(same.image.size,(2400,240))
        builder=types.SimpleNamespace(ppt=ppt,lang_index=2)
        slides=len(ppt.slides)
        self.assertEqual(core.add_extra_kpi_slides(builder,frame,["trend_correlations"],"Prueba","kg"),0)
        self.assertEqual(len(ppt.slides),slides)
        self.assertEqual(core.validate_ppt_extra_charts({"A":["trend_correlations"]},["A"]),
                         {"A":["trend_correlations"]})

    def test_sparse_correlations_show_percent_and_keep_real_zero(self):
        core=self.core
        ppt=core.Presentation(explorer.ROOT/'Modelo_PPT.pptx')
        slide=ppt.slides.add_slide(ppt.slide_layouts[core.PPT_LAYOUT_INDEX])
        matrix=[[0,None,-.8,float('nan'),.1,float('inf'),None],[None]*7,[None]*7]
        picture=core.add_trend_correlations_table(slide,ppt,2,matrix)
        rows=picture._element.nvPicPr.cNvPr.get('descr').splitlines()
        self.assertEqual(rows,['r MAT (%) · P2 | P0 | P2 | P4','Actual · 12m | 0.0% | -80.0% | 10.0%'])
        self.assertLess(picture.image.size[1],300)
        count=len(slide.shapes)
        self.assertIsNone(core.add_trend_correlations_table(slide,ppt,2,[[None]*7]*3))
        self.assertEqual(len(slide.shapes),count)

    def test_web_correlation_view_filters_missing_rows_and_pipelines(self):
        node=shutil.which('node')
        if not node:self.skipTest('Node no está disponible')
        script=r'''
const fs=require('node:fs'),vm=require('node:vm'),assert=require('node:assert/strict');
const context={window:{}};vm.createContext(context);
vm.runInContext(fs.readFileSync('explorer_web/dashboard.js','utf8'),context);
const api=context.window.CoverageDashboard;
const sheet={template:{correlations:[{label:'Actual',pipelines:[0,null,-.8,NaN,.1,Infinity,null]},
  {label:'Anterior',pipelines:Array(7).fill(null)},{label:'24m',pipelines:Array(7).fill(null)}]}};
const before=JSON.stringify(sheet),view=api.correlationView(sheet);
assert.deepEqual(Array.from(view.pipelines),[0,2,4]);assert.equal(view.rows.length,1);
assert.equal(view.rows[0].label,'Actual');assert.equal(JSON.stringify(sheet),before);
assert.equal(api.correlationPercent(0),'0.0%');assert.equal(api.correlationPercent(.12869),'12.9%');
assert.equal(api.correlationPercent(-.8),'-80.0%');assert.equal(api.correlationPercent(null),'—');
assert.equal(api.correlationPercent(Infinity),'—');
sheet.template.correlations[2].pipelines[1]=.5;
assert.equal(api.correlationView(sheet).rows[1].label,'24m');
assert.deepEqual(Array.from(api.correlationView(sheet).pipelines),[0,1,2,4]);
assert.equal(api.correlationView({template:{correlations:[]}}).rows.length,0);
'''
        subprocess.run([node,'-e',script],cwd=explorer.ROOT,check=True,capture_output=True,text=True)


if __name__ == "__main__":
    unittest.main()
