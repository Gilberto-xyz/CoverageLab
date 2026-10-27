"""Cobertura y penetración conservan la comparación de barras de la PPT."""
from pathlib import Path
import tempfile
import unittest
from unittest import mock
import zipfile
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
import pipeline_explorer as explorer


class ExcelCoverageChartTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.core = explorer.engine()

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parents[1] / 'outputs')
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / 'Template.xlsx'

    def make_template(self, months=36):
        core = self.core
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'P0_Marca'
        sheet.append([core.COL_DATA, core.COL_SELL_OUT, core.COL_SELL_IN_SIM, core.COL_PENET,
                      *[f'P{p}' for p in range(7)], core.COL_COMPRA_MEDIA, core.COL_COMPRA_OCA,
                      core.COL_FREQ, core.COL_BUYERS])
        for month in range(months):
            date = core.pd.Timestamp(2023, 1, 1) + core.pd.DateOffset(months=month)
            # Incluye cero real y períodos sin cobertura al inicio y al cierre.
            coverage = (0 if month == 11 else 70 + month / 10) if 11 <= month < months - 2 else None
            sheet.append([date.to_pydatetime(), 1000 + month, 1200 + month, 15 + month / 10,
                          *([coverage] * 7), 2.5, 1.2, 2.1, 450])
        self.original_data = [[cell.value for cell in row] for row in sheet.iter_rows(max_col=15)]
        workbook.save(self.path)
        workbook.close()

    def add_charts(self, country='Peru', english=False, rounded=False):
        with mock.patch.object(self.core, 'ROUND_COVERAGE', rounded):
            self.core.add_native_excel_charts(str(self.path), coverage_label='Cobertura P3',
                trend_axis='single', evolution_slide_variant='simple', include_english=english,
                pais_nombre=country, pipeline_overrides={'P0_Marca': 3})

    def test_saved_chart_has_two_labelled_bar_series_on_one_percentage_axis(self):
        self.make_template()
        self.add_charts()
        workbook = load_workbook(self.path)
        self.addCleanup(workbook.close)
        sheet = workbook['P0_Marca']
        chart = sheet._charts[0]
        self.assertEqual(chart.type, 'col')
        self.assertEqual(chart.grouping, 'clustered')
        self.assertEqual(len(chart._charts), 1)
        self.assertEqual(len(chart.series), 2)
        self.assertEqual(chart.y_axis.scaling.min, 0)
        self.assertEqual(chart.y_axis.numFmt.formatCode, '0"%"')
        self.assertIn('$D$13:$D$35', chart.series[0].val.numRef.f)
        self.assertIn('$H$13:$H$35', chart.series[1].val.numRef.f)
        for series, color in zip(chart.series, ('FFC000', 'D9D9D9')):
            self.assertEqual(series.graphicalProperties.solidFill.srgbClr, color)
            self.assertTrue(series.dLbls.showVal)
            self.assertEqual(series.dLbls.dLblPos, 'outEnd')
            self.assertEqual(series.dLbls.numFmt, '0.0"%"')
            self.assertIn('$A$13:$A$35', series.cat.numRef.f)
        self.assertEqual([[cell.value for cell in row] for row in sheet.iter_rows(max_row=37, max_col=15)], self.original_data)
        with zipfile.ZipFile(self.path) as archive:
            xml = ET.fromstring(archive.read('xl/charts/chart1.xml'))
        ns = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
        self.assertEqual(len(xml.findall('.//c:barChart', ns)), 1)
        self.assertEqual(len(xml.findall('.//c:lineChart', ns)), 0)
        self.assertEqual(len(xml.findall('.//c:valAx', ns)), 1)

    def test_localized_labels_rounding_and_regeneration_keep_both_series(self):
        for country, english, expected in [('Peru', False, 'Penetración Mensual'),
                                            ('Brasil', False, 'Penetração Mensal'),
                                            ('Peru', True, 'PENETRATION BY PERIOD')]:
            with self.subTest(country=country, english=english):
                self.make_template(months=24)
                self.add_charts(country, english, rounded=True)
                self.add_charts(country, english, rounded=True)
                workbook = load_workbook(self.path)
                try:
                    coverage = [chart for chart in workbook.active._charts
                                if getattr(chart.anchor, '_from', None).row == 4
                                and chart.anchor._from.col == 26]
                    self.assertEqual(len(coverage), 1)
                    chart = coverage[0]
                    self.assertEqual(chart.series[0].title.v, expected)
                    self.assertEqual(chart.series[0].dLbls.numFmt, '0.0"%"')
                    self.assertEqual(chart.series[1].dLbls.numFmt, '0"%"')
                finally:
                    workbook.close()


if __name__ == '__main__':
    unittest.main()
