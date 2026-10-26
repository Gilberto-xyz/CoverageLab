from datetime import datetime, timezone
from pathlib import Path
import hashlib
import json
import shutil
import subprocess
import tempfile
import time
import unittest
from unittest import mock

from openpyxl import Workbook
from pptx import Presentation
import pipeline_explorer as explorer
from export_progress import ExportProgressReporter, RESEARCH_STAGES, snapshot


class ResearchProgressTests(unittest.TestCase):
    def test_stage_weights_real_counts_and_eta(self):
        with tempfile.TemporaryDirectory(dir=explorer.ROOT/'outputs') as directory:
            clock=[1000.0]; path=Path(directory)/'progress.json'
            started=datetime.fromtimestamp(clock[0],timezone.utc).isoformat()
            reporter=ExportProgressReporter(path,started,clock=lambda:clock[0],stages=RESEARCH_STAGES)
            reporter.update('research',0,10,'Hoja 1')
            self.assertIsNone(snapshot(json.loads(path.read_text()),'running',started,1000)['remaining_seconds'])
            clock[0]=1020;reporter.update('research',2,10,'Hoja 2')
            running=snapshot(json.loads(path.read_text()),'running',started,1020)
            self.assertEqual(running['percent'],21)
            self.assertGreater(running['remaining_seconds'],80)
            reporter.update('saving_ppt',1,1);reporter.update('excel',4,10)
            self.assertEqual(json.loads(path.read_text())['percent'],93.4)
            reporter.update('packaging',1,1)
            self.assertLess(json.loads(path.read_text())['percent'],100)

    def test_real_worker_live_progress_completion_and_restore(self):
        with tempfile.TemporaryDirectory(dir=explorer.ROOT/'outputs') as directory:
            root=Path(directory); folder=root/'inputs'/'5w1h';folder.mkdir(parents=True)
            path=folder/'51_DIAP_Prueba.xlsx';wb=Workbook();wb.remove(wb.active)
            for name in ['5_Uno_2','5_Dos_2','5_Tres_2']:
                ws=wb.create_sheet(name);ws.append(['Compras','Moderno','Tradicional',None,'Compras_Mosaico','Tradicional','Moderno'])
                for index in range(36):
                    date=datetime(2023+index//12,index%12+1,1)
                    ws.append([date,100,200,None,date,240,110])
            wb.save(path);wb.close();before=hashlib.sha256(path.read_bytes()).hexdigest()
            state=explorer.ExplorerState(root);workflow=state.workflow
            analyzed=workflow.research_analyze({'file':'local/'+path.name,'sheet':'5_Uno_2','mode':'1'})
            job=workflow.export_research({'analysis_id':analyzed['id']})
            seen=[];deadline=time.monotonic()+100
            while time.monotonic()<deadline:
                current=workflow.research_job(job['id']);seen.append(current['progress'])
                if current['status']!='running':break
                time.sleep(.1)
            self.assertEqual(current['status'],'complete',current.get('error'))
            self.assertTrue(any(p['stage']=='research' for p in seen))
            self.assertTrue(any(p['remaining_seconds'] is not None and p['percent']<100 for p in seen))
            self.assertEqual([p['percent'] for p in seen],sorted(p['percent'] for p in seen))
            self.assertEqual(current['progress']['percent'],100)
            self.assertGreater(current['elapsed_seconds'],0)
            self.assertGreater(current['work_units'],0)
            self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(),before)
            folder=Path(current['directory'])
            self.assertTrue((folder/'_sistema'/'export-progress.json').is_file())
            self.assertFalse(Path(job['progress_path']).exists())
            self.assertEqual({p.suffix for p in folder.iterdir() if p.is_file()},{'.xlsx','.pptx'})
            ppt=Presentation(next(out['path'] for out in current['outputs'] if out['label']=='PPT'))
            trends=[slide for slide in ppt.slides if any(s.name=='5W1H · tendencia conjunta · compras' for s in slide.shapes)]
            summaries=[slide for slide in ppt.slides if any(s.name=='5W1H · comparação MAT entre compras' for s in slide.shapes)]
            self.assertEqual(len(trends),3);self.assertEqual(len(summaries),3)
            for slide in trends:
                trend=next(s for s in slide.shapes if s.name=='5W1H · tendencia conjunta · compras')
                self.assertGreaterEqual(trend.height/914400,5.9)
                self.assertFalse(any(s.name.startswith('5W1H · aportes') for s in slide.shapes))
            for slide in summaries:
                tables=[s for s in slide.shapes if s.name.startswith('5W1H · aportes')]
                self.assertEqual(len(tables),2)
            restored=explorer.ExplorerState(root).workflow.research_job(job['id'])
            self.assertEqual(restored['progress']['percent'],100)
            self.assertEqual(restored['elapsed_seconds'],current['elapsed_seconds'])
            self.assertEqual(restored['available_outputs'],['Excel','PPT','Provenance','Registro'])

    def test_browser_progress_view_and_result_renderer(self):
        node=shutil.which('node')
        if not node:self.skipTest('Node no disponible')
        script=r'''
const fs=require('node:fs'),vm=require('node:vm'),assert=require('node:assert/strict');
class Element{
 constructor(tag,attrs={},text=''){this.tag=tag;this.attrs={...attrs};this.children=[];this.textContent=text;this.style={setProperty:(k,v)=>this.attrs[k]=v,removeProperty:k=>delete this.attrs[k]};this.classList={toggle(){}};}
 append(...children){this.children.push(...children);} replaceChildren(...children){this.children=children;}
 setAttribute(k,v){this.attrs[k]=v;}removeAttribute(k){delete this.attrs[k];}
 querySelector(selector){return this.children.find(child=>child.attrs.class===selector.slice(1));}
}
const $=(...args)=>new Element(...args),nodes={generate:$('button')},state={file:'local/A',job:null};let ticks=0;
const context={window:{},state,nodes,$,Date,setInterval(){ticks++;return 1;},clearInterval(){ticks--;}};
vm.createContext(context);vm.runInContext(fs.readFileSync('explorer_web/generation.js','utf8'),context);
let source=fs.readFileSync('explorer_web/research.js','utf8'),start=source.indexOf('  let generationTimer;'),end=source.indexOf('  const pendingInputs',start);
vm.runInContext(source.slice(start,end)+';this.render=renderResearchProgress;',context);
state.job={id:'r',file:'local/A',status:'running',progressReceivedAt:Date.now(),progress:{stage:'research',completed:2,total:22,percent:12,remaining_seconds:90,elapsed_seconds:10}};
context.render();assert(nodes.generate.querySelector('.generation-main').textContent.includes('5W1H 2/22'));
assert(nodes.generate.querySelector('.generation-time').textContent.includes('restantes'));assert.equal(ticks,1);
state.job.status='complete';state.job.completedObservedAt=Date.now();context.render();assert.equal(nodes.generate.attrs['--generation-progress'],'100%');
state.job.completedObservedAt=Date.now()-5000;context.render();assert.equal(nodes.generate.textContent,'Generar informe 5W1H');assert.equal(ticks,0);
state.job.status='failed';context.render();assert.equal(nodes.generate.querySelector('.generation-main').textContent,'Reintentar generación');
state.file='local/B';context.render();assert.equal(nodes.generate.textContent,'Generar informe 5W1H');
'''
        result=subprocess.run([node,'-e',script],cwd=explorer.ROOT,capture_output=True,text=True)
        self.assertEqual(result.returncode,0,result.stdout+result.stderr)


if __name__=='__main__':unittest.main()
