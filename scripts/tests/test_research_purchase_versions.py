"""Comparar dos compras sin disfrazar la segunda de ventas ni desplazar sus fechas."""
from datetime import datetime
import hashlib
from pathlib import Path
import tempfile
import unittest
import zipfile

from openpyxl import Workbook, load_workbook
from pptx import Presentation
from studio_5w1h import ResearchStudio, _analysis_excel_type, engine


class PurchaseVersionsTests(unittest.TestCase):
    def setUp(self):
        output = Path(__file__).resolve().parents[1] / 'outputs'
        self.temp = tempfile.TemporaryDirectory(dir=output)
        self.root = Path(self.temp.name)
        self.path = self.root / 'inputs' / '5w1h' / '51_DIAP_Prueba.xlsx'
        self.path.parent.mkdir(parents=True)
        self.studio = ResearchStudio(self.root)
        self.file = 'local/' + self.path.name
        self.make()

    def tearDown(self):
        self.temp.cleanup()

    def make(self, separator=True, comparison='Compras_ajuste', offset=0, same=False, base='Compras_mosaico'):
        wb = Workbook()
        ws = wb.active
        ws.title = '5_Babysec_2'
        ws.append([base, 'Moderno', 'Tradicional'] + ([None] if separator else []) + [comparison, 'Tradicional', 'Moderno'])
        for index in range(36):
            period = datetime(2023 + index // 12, index % 12 + 1, 1)
            shifted = index + offset
            other_period = datetime(2023 + shifted // 12, shifted % 12 + 1, 1)
            base = 120 if index >= 24 else 100
            adjusted = base if same else 150 if index >= 24 else 100
            ws.append([period, base, base * 2] + ([None] if separator else []) + [other_period, adjusted * 2, adjusted])
        wb.save(self.path)
        wb.close()

    def test_names_same_months_and_no_sales_pipeline_even_with_numeric_version_name(self):
        self.make(separator=False, comparison='Compras_ajuste_2')
        result = self.studio.analyze(self.file, '5_Babysec_2', pipeline=5)
        self.assertTrue(result['ready'])
        self.assertEqual([(o['label'], o['kind'], o['pipeline']) for o in result['origins']],
                         [('Compras_mosaico', 'Compras', 0), ('Compras_ajuste_2', 'Compras', 0)])
        monthly = next(c for c in result['charts'] if c['id'] == 'monthly')
        base = next(s for s in monthly['series'] if s['name'] == 'Moderno · Compras_mosaico')
        adjusted = next(s for s in monthly['series'] if s['name'] == 'Moderno · Compras_ajuste_2')
        self.assertEqual(base['values'][:2], [100, 100])
        self.assertEqual(adjusted['values'][:2], [100, 100])
        self.assertEqual(base['values'][-1], 120)
        self.assertEqual(adjusted['values'][-1], 150)
        self.assertEqual(base['color'], adjusted['color'])
        separate = self.studio.analyze(self.file, '5_Babysec_2', mode='2', pipeline=5)
        self.assertEqual([chart['axis_range'] for chart in separate['charts'][:2]], [[0, 300], [0, 300]])
        self.assertFalse(base['dash'])
        self.assertTrue(adjusted['dash'])
        self.assertEqual([t['title'] for t in result['tables']], [
            'Compras_mosaico · participación, variación y aporte MAT', 'Compras_ajuste_2 · participación, variación y aporte MAT'])
        self.assertFalse(any('Ventas' in c['title'] for c in result['charts']))
        variation = next(c for c in result['charts'] if c['id'] == 'variation-Ventas')
        self.assertEqual(variation['series'][0]['values'], [50, 50])
        self.assertEqual(variation['axis_range'], [0, 50])
        self.assertIn(r'Compras\_ajuste', engine.format_title_with_bold_suffix('Comparación', 'Compras_ajuste'))

    def test_base_mode_filters_and_identical_versions(self):
        self.make(same=True)
        result = self.studio.analyze(self.file, '5_Babysec_2', mode='2', elements=['Moderno'])
        self.assertEqual(result['charts'][0]['title'], 'Tendencia mensual · Compras_mosaico')
        self.assertEqual(result['charts'][1]['title'], 'Tendencia mensual · Compras_ajuste')
        self.assertEqual(result['charts'][0]['series'][0]['values'], result['charts'][1]['series'][0]['values'])
        self.assertEqual(result['charts'][0]['axis_range'], result['charts'][1]['axis_range'])
        solo = self.studio.analyze(self.file, '5_Babysec_2', mode='3')
        self.assertEqual(len(solo['tables']), 1)
        self.assertTrue(all(s['origin'] == 'Compras_mosaico' for s in solo['charts'][0]['series']))

    def test_sales_headers_keep_real_pipeline_and_empty_comparison_is_explained(self):
        self.make(comparison='Ventas_2')
        result = self.studio.analyze(self.file, '5_Babysec_2')
        self.assertEqual([(o['kind'], o['pipeline']) for o in result['origins']], [('Compras', 0), ('Ventas', 2)])
        monthly = next(c for c in result['charts'] if c['id'] == 'monthly')
        sales = next(s for s in monthly['series'] if s['origin'] == 'Ventas')
        self.assertEqual(sales['values'][:3], [None, None, 100])
        wb = load_workbook(self.path)
        for row in wb.active.iter_rows(min_row=2, min_col=6):
            for cell in row:
                cell.value = None
        wb.active['E1'] = 'Compras_ajuste'
        wb.save(self.path)
        wb.close()
        empty = self.studio.analyze(self.file, '5_Babysec_2')
        self.assertTrue(empty['ready'])
        self.assertEqual(len(empty['origins']), 1)
        self.assertTrue(any('Compras_ajuste' in warning and 'vacío' in warning for warning in empty['warnings']))

    def test_different_start_months_align_by_date_in_combined_export_and_preview(self):
        self.make(offset=1)
        result = self.studio.analyze(self.file, '5_Babysec_2')
        monthly = next(c for c in result['charts'] if c['id'] == 'monthly')
        adjusted = next(s for s in monthly['series'] if s['name'] == 'Moderno · Compras_ajuste')
        self.assertIsNone(adjusted['values'][0])
        self.assertEqual(adjusted['values'][1], 100)
        self.assertTrue(any('períodos distintos' in warning for warning in result['warnings']))
        with _analysis_excel_type(engine.pd.ExcelFile)(self.path) as excel:
            frame = engine.parse_sheet_with_compras_header(excel, '5_Babysec_2')
        blocks, p = engine.split_research_origins(frame)
        configs = engine.prepare_series_configs(blocks, 'E', p)
        self.assertEqual(engine.purchase_comparison_limits(configs), (0, 324))
        for config in configs:
            grouped, changed = engine.group_size_elements_by_threshold(config.data, config.pipeline, 40)
            self.assertTrue(changed)
            self.assertEqual(grouped.attrs['source_label'], engine.series_display_label(config))
            self.assertEqual(grouped.attrs['source_kind'], 'Compras')
        combined = engine.combine_research_series(configs)
        self.assertEqual(len(combined), 37)
        self.assertTrue(engine.pd.isna(combined.loc[0, 'Moderno.V']))
        self.assertEqual(combined.loc[1, 'Moderno.V'], 100)

    def test_plain_purchases_left_and_custom_version_right_for_arbitrary_names(self):
        for index, variant in enumerate(['Compras_Mosaico', 'Compras_ajuste', 'Compras_ajt',
                                         'Compras_lo que se me ocurra', 'Compras_ajuste_2']):
            with self.subTest(variant=variant):
                self.make(separator=bool(index % 2), base='Compras', comparison=variant)
                result = self.studio.analyze(self.file, '5_Babysec_2', mode='2', pipeline=5)
                self.assertTrue(result['ready'])
                self.assertEqual([(o['label'], o['pipeline']) for o in result['origins']], [('Compras', 0), (variant, 0)])
                self.assertEqual([c['id'] for c in result['charts']], [
                    f'{metric}-{origin}' for metric in ['monthly', 'mat', 'share', 'variation', 'contribution']
                    for origin in ['Compras', 'Ventas']])
                self.assertEqual([c['title'] for c in result['charts'][:2]],
                                 ['Tendencia mensual · Compras', f'Tendencia mensual · {variant}'])
                self.assertEqual([t['title'] for t in result['tables']], [
                    'Compras · participación, variación y aporte MAT', f'{variant} · participación, variación y aporte MAT'])
                self.assertTrue(result['stats'][0]['label'].endswith(' · Compras'))
                self.assertTrue(result['stats'][1]['label'].endswith(' · ' + variant))
                self.assertEqual(result['charts'][0]['axis_range'], result['charts'][1]['axis_range'])
                self.assertEqual(result['charts'][0]['series'][0]['values'][-1], 120)
                self.assertEqual(result['charts'][1]['series'][0]['values'][-1], 150)
                solo = self.studio.analyze(self.file, '5_Babysec_2', mode='3')
                self.assertTrue(all(s['origin'] == 'Compras' for s in solo['charts'][0]['series']))

    def test_real_export_preserves_labels_positions_source_and_second_excel_block(self):
        variant = 'Compras_lo que se me ocurra'
        self.make(separator=False, base='Compras', comparison=variant)
        before = self.path.read_bytes()
        result = self.studio.export(self.file, {'mode': '2', 'sheets': ['5_Babysec_2'], 'pipeline': 5,
                                               'pipeline_overrides': {'Babysec': 3}}, self.root / 'Output')
        self.assertEqual(self.path.read_bytes(), before)
        self.assertEqual(result['fingerprint'], hashlib.sha256(before).hexdigest())
        self.assertEqual(result['pipelines_applied'], {})
        ppt = Presentation(result['outputs']['PPT'])
        found = set()
        for slide in ppt.slides:
            names = {shape.name: shape for shape in slide.shapes}
            for metric in ['tendencia', 'aportes', 'participación']:
                left = names.get(f'5W1H · {metric} · Compras')
                right = names.get(f'5W1H · {metric} · {variant}')
                if left is not None:
                    self.assertIsNotNone(right)
                    self.assertLess(left.left, right.left)
                    self.assertLessEqual(left.left + left.width, right.left)
                    self.assertEqual(left.top, right.top)
                    found.add(metric)
        self.assertEqual(found, {'tendencia', 'aportes', 'participación'})
        wb = load_workbook(result['outputs']['Excel'])
        ws = wb['5_Babysec_2']
        blocks = engine.detect_excel_table_blocks(ws)
        self.assertEqual(blocks['compras']['display_tipo'], 'Compras')
        self.assertEqual(blocks['ventas']['display_tipo'], variant)
        self.assertEqual(blocks['ventas']['pipeline'], 0)
        self.assertEqual(blocks['ventas']['date_col'], 4)
        self.assertTrue(ws._charts)
        titles = [s.title.v for chart in ws._charts for s in chart.series if s.title and s.title.v]
        self.assertTrue(any(variant in label for label in titles))
        self.assertTrue(any('Compras' in label for label in titles))
        wb.close()
        with zipfile.ZipFile(result['outputs']['Excel']) as archive:
            charts = ''.join(archive.read(n).decode() for n in archive.namelist() if n.startswith('xl/charts/chart') and n.endswith('.xml'))
        self.assertIn(variant, charts)

    def test_joint_export_keeps_tables_off_tall_trend_without_a_common_mat(self):
        self.make(base='Compras',offset=30)
        analysis=self.studio.analyze(self.file,'5_Babysec_2')
        self.assertFalse(analysis['purchase_comparison']['ready'])
        exported=self.studio.export(self.file,{'mode':'1','sheets':['5_Babysec_2']},self.root/'Output')
        ppt=Presentation(exported['outputs']['PPT'])
        trend=next(slide for slide in ppt.slides if any(s.name=='5W1H · tendencia conjunta · compras' for s in slide.shapes))
        self.assertFalse(any(s.name.startswith('5W1H · aportes') for s in trend.shapes))
        table_slide=next(slide for slide in ppt.slides if any(s.name.startswith('5W1H · aportes') for s in slide.shapes))
        self.assertIn('No hay 12 meses comunes', '\n'.join(s.text for s in table_slide.shapes if s.has_text_frame))


if __name__ == '__main__':
    unittest.main()
