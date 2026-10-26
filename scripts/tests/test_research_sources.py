"""Archivos elegidos desde cualquier carpeta y descubrimiento de salidas 5W1H."""
import base64
from datetime import datetime
from pathlib import Path
import shutil
import subprocess
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
from pptx import Presentation
import pipeline_explorer as explorer


class ResearchSourcesTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=explorer.ROOT / 'outputs')
        self.root = Path(self.temp.name)
        self.state = explorer.ExplorerState(self.root)
        self.source = self.root / 'otra_carpeta' / '51_DIAP_Softys_5w1h.xlsx'
        self.source.parent.mkdir()
        wb = Workbook()
        for index, title in enumerate(['5_Babysec_2', '3_Babysec_1', '5_Babysec_1', '4_Babysec']):
            ws = wb.active if index == 0 else wb.create_sheet()
            ws.title = title
            ws.append(['Compras', 'Moderno', 'Tradicional', None, 'Ventas_0', 'Moderno', 'Tradicional'])
            for month in range(36):
                date = datetime(2023 + month // 12, month % 12 + 1, 1)
                ws.append([date, 120 if month >= 24 else 100, 80 if month >= 24 else 100,
                           None, date, 240 if month >= 24 else 200, 160 if month >= 24 else 200])
        wb.save(self.source)
        wb.close()

    def tearDown(self):
        self.temp.cleanup()

    def choose(self):
        return self.state.workflow.import_input({'kind': 'research', 'filename': self.source.name,
            'content': base64.b64encode(self.source.read_bytes()).decode(), 'reuse_existing': True})

    def test_import_preserves_original_bytes_and_reuses_identical_local_file(self):
        original = self.source.read_bytes()
        chosen = self.choose()
        local = self.state.workflow.research.source(chosen['file'])
        self.assertEqual(local.read_bytes(), original)
        self.assertEqual(self.source.read_bytes(), original)
        again = self.choose()
        self.assertEqual(again['file'], chosen['file'])
        self.assertTrue(again['reused_existing'])
        self.assertEqual(len(list(local.parent.glob('*.xlsx'))), 1)
        wb = load_workbook(self.source)
        wb.active['B2'] = 999
        wb.save(self.source)
        wb.close()
        changed = self.choose()
        self.assertNotEqual(changed['file'], chosen['file'])
        self.assertEqual(local.read_bytes(), original)

    def test_inspection_exposes_all_dimensions_and_explains_missing_ones(self):
        chosen = self.choose()
        info = self.state.workflow.research.inspect(chosen['file'])
        self.assertEqual({s['segment'] for s in info['sheets']}, {'5-2', '3-1', '5-1', '4'})
        self.assertTrue(all(s['has_data'] for s in info['sheets']))
        self.assertIn('2', {s['id'] for s in info['unavailable_segments']})
        self.assertNotIn('4', {s['id'] for s in info['unavailable_segments']})

    def test_variation_and_contribution_charts_match_engine_tables_and_filters(self):
        chosen = self.choose()
        research = self.state.workflow.research
        result = research.analyze(chosen['file'], '5_Babysec_2')
        for origin in ('Compras', 'Ventas'):
            table = next(t for t in result['tables'] if t['id'] == f'contribution-{origin}')
            for prefix, row, unit in [('variation', 2, '%'), ('contribution', 3, 'pp')]:
                chart = next(c for c in result['charts'] if c['id'] == f'{prefix}-{origin}')
                self.assertEqual(chart['dates'], ['Moderno', 'Tradicional'])
                self.assertEqual(chart['series'][0]['unit'], unit)
                self.assertEqual(chart['series'][0]['values'], [float(v.strip('%')) for v in table['rows'][row][1:3]])
        self.assertEqual(next(c for c in result['charts'] if c['id'] == 'variation-Compras')['series'][0]['values'], [20, -20])
        filtered = research.analyze(chosen['file'], '5_Babysec_2', elements=['Moderno'], mode='3')
        self.assertFalse(any(c['id'].endswith('-Ventas') for c in filtered['charts']))
        chart = next(c for c in filtered['charts'] if c['id'] == 'contribution-Compras')
        self.assertEqual(chart['dates'], ['Moderno'])
        self.assertEqual(chart['series'][0]['values'], [10])

    def test_invalid_workbook_is_not_added_to_research_catalog(self):
        wb = Workbook()
        wb.save(self.source)
        wb.close()
        with self.assertRaisesRegex(ValueError, 'estructura 5W1H'):
            self.choose()
        self.assertEqual(self.state.workflow.research.catalog()['files'], [])

    def test_sales_left_and_purchases_right_in_web_and_exported_slides(self):
        chosen = self.choose()
        research = self.state.workflow.research
        result = research.analyze(chosen['file'], '5_Babysec_2', mode='2')
        self.assertEqual([c['id'] for c in result['charts']], [
            f'{metric}-{origin}' for metric in ['monthly', 'mat', 'share', 'variation', 'contribution']
            for origin in ['Ventas', 'Compras']])
        self.assertEqual([t['id'] for t in result['tables']], ['contribution-Ventas', 'contribution-Compras'])
        for field in ['charts', 'stats']:
            self.assertTrue(result[field][0]['id'].endswith('Ventas') if field == 'charts'
                            else result[field][0]['label'].endswith('Ventas'))
        before = self.source.read_bytes()
        exported = research.export(chosen['file'], {'mode': '2', 'sheets': ['5_Babysec_2']}, self.root / 'Output')
        ppt = Presentation(exported['outputs']['PPT'])
        found = set()
        for slide in ppt.slides:
            for metric in ['tendencia', 'aportes', 'participación']:
                pictures = {shape.name: shape for shape in slide.shapes if shape.name.startswith(f'5W1H · {metric} · ')}
                if pictures:
                    sales = pictures[f'5W1H · {metric} · Ventas']
                    purchases = pictures[f'5W1H · {metric} · Compras']
                    self.assertLess(sales.left, purchases.left)
                    self.assertLessEqual(sales.left + sales.width, purchases.left)
                    self.assertEqual(sales.top, purchases.top)
                    found.add(metric)
        self.assertEqual(found, {'tendencia', 'aportes', 'participación'})
        self.assertEqual(self.source.read_bytes(), before)

    def test_web_picker_opens_selected_book_and_resets_after_failure(self):
        node = shutil.which('node')
        if not node: self.skipTest('Node no disponible')
        script = r'''
const fs=require('node:fs'),vm=require('node:vm'),assert=require('node:assert/strict');
const state={busy:false,manual:false,file:null,inspection:null},nodes={upload:{value:'selected'}};
const calls=[],messages=[];let opens=0,fail=false;
state.bridge={api:async(path,payload)=>{calls.push({path,payload});if(fail)throw Error('Libro no válido');return {file:'local/51_DIAP_Softys_5w1h.xlsx',name:'Softys',reused_existing:true};}};
const context={state,nodes,Uint8Array,String,btoa,updateActions(){},message:text=>messages.push(text),catalog:async()=>{},
 open:async file=>{assert.equal(state.busy,false);opens++;state.file=file;state.inspection={sheets:[1,2,3]};}};
vm.createContext(context);
const source=fs.readFileSync('explorer_web/research.js','utf8'),a=source.indexOf('  async function importResearch('),b=source.indexOf('  function renderDimensions(',a);
vm.runInContext(source.slice(a,b)+';this.pick=importResearch;',context);
async function run(){
const file={name:'51_DIAP_Softys_5w1h.xlsx',size:3,arrayBuffer:async()=>new Uint8Array([1,2,3]).buffer};
await context.pick(file);assert.equal(opens,1);assert(state.manual);assert.equal(calls[0].path,'/api/studio/import');
assert(calls[0].payload.reuse_existing);assert.equal(calls[0].payload.kind,'research');assert.equal(calls[0].payload.content,'AQID');
assert(messages.at(-1).includes('3 hojas'));assert.equal(nodes.upload.value,'');
fail=true;await context.pick(file);assert.equal(opens,1);assert(!state.busy);assert.equal(messages.at(-1),'Libro no válido');
await context.pick({...file,size:11*1024*1024});assert(messages.at(-1).includes('10 MiB'));assert.equal(calls.length,2);
}
run().catch(e=>{console.error(e);process.exitCode=1;});
'''
        result = subprocess.run([node, '-e', script], cwd=explorer.ROOT, capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)


if __name__ == '__main__':
    unittest.main()
