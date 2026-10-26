"""Draft persistence must not weaken validation or mutate any Excel source."""
from contextlib import closing
import json
import math
from pathlib import Path
import tempfile
import threading
import unittest
from unittest import mock
import urllib.error
import urllib.parse
import urllib.request

from openpyxl import Workbook, load_workbook

import pipeline_explorer as explorer
from studio_drafts import StudioDrafts
from studio_inputs import InputStudio, file_fingerprint


class StudioDraftTests(unittest.TestCase):
    def setUp(self):
        output = explorer.ROOT / "outputs"
        output.mkdir(exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(prefix="studio_draft_test_", dir=output)
        self.root = Path(self.temp.name) / "coverage"
        self.root.mkdir()
        self.external = Path(self.temp.name) / "research_original"
        self.external.mkdir()
        self.inputs = InputStudio(self.root, self.external)
        self.drafts = StudioDrafts(self.inputs)
        self.source = self.make_book(self.root, "52_CARB_Marca.xlsx", "P0_Marca")
        self.payload = {"file": self.source.name, "kind": "coverage",
                        "fingerprint": file_fingerprint(self.source), "sheet": "P0_Marca",
                        "row": 1, "column": 1, "number_locale": "auto", "edits": [
                            {"sheet": "P0_Marca", "row": 3, "column": 1,
                             "values": [["esto no es una fecha", "Sí", "No", "1,234", None, 0],
                                        ["2025-02-01", "=SUM(A1:A2)", "texto", "", 1.25, -4]]}]}

    def tearDown(self):
        self.temp.cleanup()

    @staticmethod
    def make_book(directory, filename, sheet):
        workbook = Workbook()
        workbook.active.title = sheet
        workbook.active.append(["Fecha", "Valor"])
        workbook.active.append(["2025-01-01", 100])
        target = directory / filename
        workbook.save(target)
        workbook.close()
        return target

    def test_invalid_capture_is_preserved_exactly_without_touching_excel(self):
        original = self.source.read_bytes()
        saved = self.drafts.save(self.payload)
        self.assertFalse(saved["source_changed"])
        self.assertEqual(saved["draft"]["edits"], self.payload["edits"])
        self.assertEqual(self.source.read_bytes(), original)
        restored = StudioDrafts(InputStudio(self.root, self.external)).get(self.source.name)
        self.assertEqual(restored, saved)
        self.assertEqual(self.source.read_bytes(), original)
        self.assertEqual(list(self.drafts.directory.glob("*.tmp")), [])

    def test_drafts_survive_new_workflow_and_do_not_enable_invalid_excel_save(self):
        state = explorer.ExplorerState(self.root, research_root=self.external)
        saved = state.workflow.save_draft(self.payload)
        recreated = explorer.ExplorerState(self.root, research_root=self.external)
        self.assertEqual(recreated.workflow.input_draft({"file": self.source.name}), saved)
        original = self.source.read_bytes()
        with self.assertRaises(ValueError):
            recreated.workflow.save_input(self.payload)
        self.assertEqual(self.source.read_bytes(), original)
        self.assertEqual(recreated.workflow.input_draft({"file": self.source.name}), saved)

    def test_source_changed_keeps_old_draft_instead_of_overwriting_source(self):
        old = self.drafts.save(self.payload)["draft"]
        with closing(load_workbook(self.source)) as workbook:
            workbook.active["B2"] = 200
            workbook.save(self.source)
        changed = self.source.read_bytes()
        self.assertTrue(self.drafts.get(self.source.name)["source_changed"])
        self.assertEqual(self.drafts.get(self.source.name)["draft"], old)
        saved = self.drafts.save(self.payload)
        self.assertTrue(saved["source_changed"])
        self.assertEqual(saved["draft"]["fingerprint"], old["fingerprint"])
        self.assertEqual(self.source.read_bytes(), changed)

    def test_external_aliases_share_one_draft_without_creating_excel_copy(self):
        source = self.make_book(self.external, self.source.name, "5_Marca_2")
        original = source.read_bytes()
        payload = self.payload | {"kind": "research", "file": "external:" + source.name,
                                  "fingerprint": file_fingerprint(source), "sheet": "5_Marca_2"}
        saved = self.drafts.save(payload)
        self.assertEqual(saved["draft"]["file"], "external/" + source.name)
        self.assertEqual(self.drafts.get("external/" + source.name, "research"), saved)
        self.assertEqual(len(list(self.drafts.directory.glob("*.json"))), 1)
        self.assertEqual(source.read_bytes(), original)
        self.assertFalse(self.inputs.research_local.exists())
        self.drafts.save(self.payload)
        self.assertEqual(len(list(self.drafts.directory.glob("*.json"))), 2)
        self.drafts.discard("external:" + source.name, "research")
        self.assertIsNone(self.drafts.get("external/" + source.name, "research")["draft"])
        self.assertIsNotNone(self.drafts.get(self.source.name)["draft"])
        self.assertEqual(source.read_bytes(), original)

    def test_multiple_sheets_paging_and_discard_are_independent_of_current_excel_schema(self):
        # A removed/renamed sheet is still useful in a stale draft; recovery must
        # not run Excel's current sheet/date/number validators against it.
        payload = self.payload | {"sheet": "Hoja anterior", "row": 12000, "column": 200,
                                  "number_locale": "decimal-comma", "edits": self.payload["edits"] + [
                                      {"sheet": "Otra hoja", "row": 20000, "column": 256, "values": [["Sí"]]}]}
        saved = self.drafts.save(payload)
        self.assertEqual(saved["draft"]["row"], 12000)
        self.assertEqual(saved["draft"]["column"], 200)
        self.assertEqual(saved["draft"]["number_locale"], "decimal-comma")
        self.assertEqual(saved["draft"]["edits"], payload["edits"])
        source_bytes = self.source.read_bytes()
        self.assertEqual(self.drafts.discard(self.source.name), {"draft": None, "source_changed": False})
        self.assertEqual(self.source.read_bytes(), source_bytes)
        self.assertIsNone(self.drafts.get(self.source.name)["draft"])
        self.drafts.discard(self.source.name)

    def test_invalid_shape_bounds_types_and_limits_leave_previous_draft_unchanged(self):
        self.drafts.save(self.payload)
        target = next(self.drafts.directory.glob("*.json"))
        original = target.read_bytes()
        source_bytes = self.source.read_bytes()
        block = {"sheet": "P0_Marca", "row": 3, "column": 1, "values": [[1]]}
        invalid = [{"fingerprint": "fake"}, {"fingerprint": "A" * 64}, {"row": True}, {"row": 0},
                   {"column": 257}, {"number_locale": []}, {"sheet": "Bad/Sheet"}, {"edits": None}]
        invalid += [{"edits": [block | changes]} for changes in (
            {"row": 20000, "values": [[1], [2]]}, {"column": 256, "values": [[1, 2]]},
            {"values": [[1, 2], [3]]}, {"values": []}, {"values": [[True]]},
            {"values": [[{"value": 1}]]}, {"values": [[math.inf]]}, {"values": [[math.nan]]},
            {"values": [["x" * 32768]]}, {"values": [[0] * 256] * 196},
            {"values": [["x" * 32767] * 130]},
        )]
        invalid.append({"edits": [block | {"sheet": f"Hoja {i}"} for i in range(100)]})
        for changes in invalid:
            with self.subTest(changes=str(changes)[:120]), self.assertRaises(ValueError):
                self.drafts.save(self.payload | changes)
            self.assertEqual(target.read_bytes(), original)
            self.assertEqual(self.source.read_bytes(), source_bytes)

    def test_atomic_replace_failure_cleans_temporary_and_keeps_previous_draft(self):
        self.drafts.save(self.payload)
        target = next(self.drafts.directory.glob("*.json"))
        original = target.read_bytes()
        with mock.patch("studio_drafts.os.replace", side_effect=OSError("Simulated disk failure")):
            with self.assertRaises(OSError):
                self.drafts.save(self.payload | {"row": 10})
        self.assertEqual(target.read_bytes(), original)
        self.assertEqual(list(self.drafts.directory.glob("*.tmp")), [])

    def test_paths_and_linked_draft_directory_are_rejected_before_writes(self):
        for file, kind in (("../fuera.xlsx", "coverage"), (str(self.source), "coverage"),
                           ("local/../fuera.xlsx", "research"), ("external/C:/fuera.xlsx", "research"),
                           ("no_existe.xlsx", "coverage"), (self.source.name, "other")):
            with self.subTest(file=file), self.assertRaises(ValueError):
                self.drafts.save(self.payload | {"file": file, "kind": kind})
            with self.assertRaises(ValueError):
                self.drafts.get(file, kind)
            with self.assertRaises(ValueError):
                self.drafts.discard(file, kind)
        self.assertFalse(self.drafts.directory.exists())
        native_is_symlink = Path.is_symlink
        def linked(path):
            return path == self.root / "inputs" or native_is_symlink(path)
        with mock.patch.object(Path, "is_symlink", linked):
            for operation in (lambda: self.drafts.save(self.payload),
                              lambda: self.drafts.get(self.source.name),
                              lambda: self.drafts.discard(self.source.name)):
                with self.assertRaisesRegex(ValueError, "enlaces"):
                    operation()
        self.assertFalse(self.drafts.directory.exists())

    def test_http_routes_keep_local_origin_and_session_guards(self):
        state = explorer.ExplorerState(self.root, research_root=self.external)
        server = explorer.ExplorerHTTPServer(("127.0.0.1", 0), explorer.make_handler(state))
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        base = f"http://127.0.0.1:{server.server_port}"
        def request(path, payload=None, headers=None, method=None):
            data = json.dumps(payload, ensure_ascii=False).encode("utf-8") if payload is not None else None
            req = urllib.request.Request(base + path, data=data, method=method,
                headers={"X-Explorer-Token": state.token, "Content-Type": "application/json"} | (headers or {}))
            with urllib.request.urlopen(req, timeout=10) as response:
                return json.load(response)
        query = "?" + urllib.parse.urlencode({"file": self.source.name, "kind": "coverage"})
        try:
            self.assertIsNone(request("/api/studio/draft" + query)["draft"])
            saved = request("/api/studio/draft", self.payload)
            self.assertEqual(request("/api/studio/draft" + query), saved)
            for path, payload in (("/api/studio/draft" + query, None),
                                  ("/api/studio/draft", self.payload),
                                  ("/api/studio/discard-draft", {"file": self.source.name})):
                for headers in ({"X-Explorer-Token": "wrong"}, {"Origin": "https://example.invalid"},
                                {"Host": "example.invalid"}):
                    with self.subTest(path=path, headers=headers):
                        # Authentication is checked before the request body. An
                        # empty POST avoids TCP reset races from unread bodies.
                        with self.assertRaises(urllib.error.HTTPError) as error:
                            request(path, None, headers, "POST" if payload is not None else "GET")
                        self.assertEqual(error.exception.code, 403)
                        error.exception.read()
                        error.exception.close()
            self.assertEqual(request("/api/studio/draft" + query), saved)
            request("/api/studio/discard-draft", {"file": self.source.name, "kind": "coverage"})
            self.assertIsNone(request("/api/studio/draft" + query)["draft"])
        finally:
            server.shutdown()
            server.server_close()
            thread.join(timeout=10)


if __name__ == "__main__":
    unittest.main()
