"""Contract checks for typed, atomic editing of real Excel input workbooks."""
import base64
from contextlib import closing
from datetime import datetime
from pathlib import Path
import tempfile
import threading
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900, to_excel

from studio_inputs import InputConflictError, InputStudio, file_fingerprint


class InputStudioTests(unittest.TestCase):
    def setUp(self):
        output = Path(__file__).resolve().parents[1] / "outputs"
        output.mkdir(exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(prefix="input_studio_test_", dir=output)
        self.root = Path(self.temp.name) / "coverage"
        self.root.mkdir()
        self.external = Path(self.temp.name) / "original_research"
        self.external.mkdir()
        self.studio = InputStudio(self.root, self.external)

    def tearDown(self):
        self.temp.cleanup()

    def create(self, **updates):
        return self.studio.create({"country_code": "52", "category_code": "CARB",
                                   "manufacturer": "Fabricante", "brands": ["Marca"],
                                   "unit": "VOL1_P", **updates})

    def save(self, book, values, row=3, column=1, **options):
        return self.studio.save(book["file"], fingerprint=book["fingerprint"],
            kind=book["kind"], edits=[{"sheet": book["sheet"], "row": row,
                                       "column": column, "values": values}], **options)

    def external_book(self):
        path = self.external / "52_CARB_Marca.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "5_Marca_2"
        sheet.append(["Fecha", "Moderno", "Tradicional"])
        sheet.append([datetime(2025, 1, 1), 12.5, 20])
        sheet["B2"].number_format = "0.00"
        sheet["B2"].fill = PatternFill("solid", fgColor="EAF5F2")
        sheet["D2"] = "=SUM(B2:C2)"
        workbook.save(path)
        return path

    def test_creation_preserves_positional_schema_and_names_without_overwriting(self):
        book = self.create(brands=["Marca", "Otra"])
        original = (self.root / book["file"]).read_bytes()
        self.assertEqual(book["file"], "52_CARB_Fabricante.xlsx")
        self.assertEqual([s["name"] for s in book["sheets"]], ["P0_Marca", "P0_Otra"])
        self.assertEqual(book["data_row"], 3)
        self.assertFalse(book["validation"]["ready"])
        with closing(load_workbook(self.root / book["file"])) as workbook:
            self.assertEqual(workbook["P0_Marca"]["B1"].value, "Weighted VOL1_P")
            self.assertEqual(workbook["P0_Marca"]["H2"].value, "SELL-IN: Marca")
            self.assertIsNone(workbook["P0_Marca"]["B3"].value)
        second = self.create()
        self.assertNotEqual(second["file"], book["file"])
        self.assertEqual((self.root / book["file"]).read_bytes(), original)

    def test_paste_dates_locale_numbers_blank_and_real_zero_have_distinct_types(self):
        book = self.create()
        result = self.save(book, [["ene 25", "1.234,5", "20%", "2,5", "0", None, "100", "2,345.6"],
                                  ["Feb-25 ", "100", "20,5", "2", "1", "1", "10", "200"]])
        with closing(load_workbook(self.root / book["file"])) as workbook:
            sheet = workbook[book["sheet"]]
            self.assertEqual(sheet["A3"].value, datetime(2025, 1, 1))
            self.assertEqual(sheet["B3"].value, 1234.5)
            self.assertEqual(sheet["C3"].value, 20)
            self.assertEqual(sheet["D3"].value, 2.5)
            self.assertEqual(sheet["E3"].value, 0)
            self.assertIsNone(sheet["F3"].value)
            self.assertEqual(sheet["H3"].value, 2345.6)
            self.assertEqual(sheet["A4"].value, datetime(2025, 2, 1))
        self.assertEqual(result["validation"]["error_count"], 0)
        self.assertTrue(any("KPI sin dato" in warning for warning in result["validation"]["warnings"]))

    def test_rectangular_save_is_atomic_and_conflicts_do_not_overwrite(self):
        book = self.create()
        path = self.root / book["file"]
        before = path.read_bytes()
        with self.assertRaisesRegex(ValueError, "B4"):
            self.save(book, [["2025-01-01", "10"], ["2025-02-01", "texto inválido"]])
        self.assertEqual(path.read_bytes(), before)
        saved = self.save(book, [["2025-01-01", "10"]])
        expected = path.read_bytes()
        with self.assertRaises(InputConflictError):
            self.save(book, [["2025-01-01", "999"]])
        self.assertEqual(path.read_bytes(), expected)
        self.assertNotEqual(saved["fingerprint"], book["fingerprint"])

    def test_existing_formulas_formats_and_source_are_preserved_on_external_copy(self):
        path = self.external_book()
        original = path.read_bytes()
        book = self.studio.workbook(f"external/{path.name}", kind="research")
        self.assertTrue(book["copy_on_save"])
        self.assertTrue(book["readonly"][1][3])
        with self.assertRaisesRegex(ValueError, "fórmula"):
            self.save(book, [["0"]], row=2, column=4)
        self.assertFalse(self.studio.research_local.exists())
        literal = self.save(book, [["=SUM(A1:A2)"]], row=2, column=2)
        with closing(load_workbook(self.studio.research_local / literal['name'])) as workbook:
            self.assertEqual(workbook[book['sheet']]['B2'].data_type, 's')
            self.assertEqual(workbook[book['sheet']]['B2'].value, '=SUM(A1:A2)')
        saved = self.save(book, [["30,5"]], row=2, column=2)
        self.assertTrue(saved["file"].startswith("local/"))
        self.assertEqual(path.read_bytes(), original)
        with closing(load_workbook(self.studio.research_local / saved["name"])) as workbook:
            sheet = workbook[book["sheet"]]
            self.assertEqual(sheet["B2"].value, '30,5')
            self.assertEqual(sheet["B2"].number_format, "0.00")
            self.assertEqual(sheet["B2"].fill.fgColor.rgb, "00EAF5F2")
            self.assertEqual(sheet["D2"].value, "=SUM(B2:C2)")
        self.assertEqual(self.studio.workbook(f"local:{saved['name']}", kind="research")["file"], saved["file"])

    def test_research_percentages_and_numeric_text_are_copied_verbatim(self):
        path = self.external_book()
        book = self.studio.workbook(f"external/{path.name}", kind="research")
        saved = self.save(book, [["20%", "0,20"]], row=2, column=2)
        with closing(load_workbook(self.studio.research_local / saved["name"])) as workbook:
            self.assertEqual(workbook[book["sheet"]]["B2"].value, '20%')
            self.assertEqual(workbook[book["sheet"]]["C2"].value, '0,20')

    def test_import_is_exact_copy_and_rejects_unsafe_paths_and_invalid_bytes(self):
        original = self.external_book().read_bytes()
        imported = self.studio.import_workbook({"kind": "research", "filename": "Captura.xlsx",
            "content": base64.b64encode(original).decode("ascii")})
        self.assertEqual((self.studio.research_local / imported["name"]).read_bytes(), original)
        for name in ("../fuera.xlsx", "C:\\fuera.xlsx", "~$Temporal.xlsx", "fuera.xlsm"):
            with self.subTest(name=name), self.assertRaises(ValueError):
                self.studio.import_workbook({"filename": name, "content": original})
        with self.assertRaises(ValueError):
            self.studio.import_workbook({"filename": "invalid.xlsx", "content": b"not a workbook"})
        with self.assertRaises(ValueError):
            self.studio.workbook("../Captura.xlsx")

    def test_add_sheet_and_duplicate_keep_source_values_and_schema(self):
        book = self.create()
        result = self.studio.add_sheet(book["file"], fingerprint=book["fingerprint"], name="Otra")
        self.assertEqual(result["sheet"], "P0_Otra")
        self.assertEqual(result["values"][0][1], "Weighted VOL1_P")
        filled = self.save(result, [["2025-01-01", "20"]])
        duplicate = self.studio.add_sheet(filled["file"], fingerprint=filled["fingerprint"],
                                          name="P0_Copia", source_sheet="P0_Otra")
        self.assertEqual(duplicate["values"][2][1], 20)
        with self.assertRaisesRegex(ValueError, "Ya existe"):
            self.studio.add_sheet(duplicate["file"], fingerprint=duplicate["fingerprint"], name="P0_copia")

    def test_limits_and_overlapping_blocks_reject_without_disk_mutation(self):
        book = self.create()
        path = self.root / book["file"]
        before = path.read_bytes()
        with self.assertRaises(ValueError):
            self.studio.workbook(book["file"], row=20000, limit=2)
        with self.assertRaisesRegex(ValueError, "rectangular|misma cantidad"):
            self.save(book, [["10", "20"], ["30"]], row=3, column=2)
        with self.assertRaisesRegex(ValueError, "superponen"):
            self.studio.save(book["file"], fingerprint=book["fingerprint"], edits=[
                {"sheet": book["sheet"], "row": 3, "column": 2, "values": [["10"]]},
                {"sheet": book["sheet"], "row": 3, "column": 2, "values": [["20"]]},
            ])
        self.assertEqual(path.read_bytes(), before)

    def test_two_concurrent_saves_with_same_version_allow_only_one(self):
        book = self.create()
        barrier, results = threading.Barrier(2), []
        def run(number):
            barrier.wait()
            try:
                self.save(book, [["2025-01-01", str(number)]])
                results.append("saved")
            except InputConflictError:
                results.append("conflict")
        threads = [threading.Thread(target=run, args=(number,)) for number in (10, 20)]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join(timeout=20)
        self.assertCountEqual(results, ["saved", "conflict"])

    def test_number_locale_uses_region_in_auto_and_preserves_manual_override(self):
        for value, locale, expected in (("1,234", "decimal-dot", 1234),
                                         ("1.234", "decimal-comma", 1234),
                                         ("1.234", "decimal-dot", 1.234),
                                         ("1,234", "decimal-comma", 1.234),
                                         ("0.054", "auto", 0.054),
                                         ("1,234,567", "auto", 1234567),
                                         ("10,921", "auto", 10921),
                                         ("10.921", "auto", 10.921)):
            with self.subTest(value=value, locale=locale):
                book = self.create()
                saved = self.save(book, [[value]], column=2, number_locale=locale)
                self.assertEqual(saved["values"][2][1], expected)
        book = self.create()
        saved = self.save(book, [['1.234']], column=2, number_locale_hint='decimal-comma')
        self.assertEqual(saved['values'][2][1],1234)
        # True typed numbers never need string-format inference.
        saved = self.save(saved, [[{"type": "number", "value": 1.234}]], column=2, number_locale_hint='decimal-comma')
        self.assertEqual(saved["values"][2][1], 1.234)

    def test_research_babysec_paste_preserves_text_without_touching_original(self):
        path = self.external_book()
        with closing(load_workbook(path)) as workbook:
            workbook.active.title='5_Babysec_2'
            workbook.active['G2']=10921
            workbook.active['G2'].number_format='#,##0'
            workbook.save(path)
        original = path.read_bytes()
        book = self.studio.workbook(f'external/{path.name}',kind='research')
        self.assertEqual(book['values'][1][6],10921)
        saved = self.save(book,[['10,921']],row=2,column=7)
        self.assertEqual(saved['values'][1][6],'10,921')
        self.assertEqual(path.read_bytes(),original)
        for values in ([['10,921','20,5']], [['10,921','20.5']], [['10.921','1.234,50']]):
            saved = self.save(saved,values,row=2,column=7)
            self.assertEqual(saved['values'][1][6:8],values[0])

    def test_research_allows_text_over_existing_date_and_numbers(self):
        path=self.external_book()
        book=self.studio.workbook(f'external/{path.name}',kind='research')
        saved=self.save(book,[['Total Softys','No disponible','  10,921  ']],row=2)
        self.assertEqual(saved['values'][1][:3],['Total Softys','No disponible','  10,921  '])
        same=self.save(saved,[['Otra etiqueta']],row=2)
        self.assertEqual(same['file'],saved['file'])
        self.assertEqual(len(list(self.studio.research_local.glob('*.xlsx'))),1)

    def test_month_display_is_consistent_without_changing_source_values_or_bytes(self):
        book = self.create()
        path = self.root / book["file"]
        with closing(load_workbook(path)) as workbook:
            sheet = workbook[book["sheet"]]
            sheet["A1"] = "Jan-22"  # Metadata does not become a calendar cell.
            for row, value in enumerate(["Jan-24  ", datetime(2024, 2, 1), "mar 24",
                                          to_excel(datetime(2024, 4, 1)), "mes inválido"], 3):
                sheet.cell(row, 1, value)
                sheet.cell(row, 2, 45323)
                sheet.cell(row, 8, 10)
            sheet["A4"].number_format = "mmm-yy"
            workbook.save(path)
        before, fingerprint = path.read_bytes(), file_fingerprint(path)
        shown = self.studio.workbook(book["file"])
        self.assertEqual(shown["date_columns"], [{"column": 1, "data_row": 3}])
        self.assertEqual([row[0] for row in shown["display_values"]][2:6],
                         ["ene 2024", "feb 2024", "mar 2024", "abr 2024"])
        self.assertEqual(shown["display_values"][0][0], "Jan-22")
        self.assertEqual(shown["display_types"][0][0], "text")
        self.assertEqual(shown["display_values"][6][0], "mes inválido")
        self.assertEqual(shown["display_types"][6][0], "invalid-date")
        self.assertEqual(shown["values"][2][0], "Jan-24  ")
        self.assertEqual(shown["types"][2][0], "text")
        self.assertEqual(shown["values"][3][0], "2024-02-01")
        self.assertEqual(shown["formats"][3][0], "mmm-yy")
        self.assertEqual(shown["values"][5][0], to_excel(datetime(2024, 4, 1)))
        self.assertEqual(shown["types"][5][0], "number")
        self.assertEqual(shown["values"][2][1], 45323)
        self.assertEqual(shown["display_values"][2][1], "45323")
        self.assertEqual(shown["display_types"][2][1], "number")
        self.assertFalse(shown["validation"]["ready"])
        self.assertEqual(shown["fingerprint"], fingerprint)
        self.assertEqual(file_fingerprint(path), fingerprint)
        self.assertEqual(path.read_bytes(), before)

    def test_month_display_uses_each_workbook_excel_epoch(self):
        for epoch in [CALENDAR_WINDOWS_1900, CALENDAR_MAC_1904]:
            with self.subTest(epoch=epoch):
                book = self.create()
                path = self.root / book["file"]
                workbook = Workbook()
                workbook.epoch = epoch
                sheet = workbook.active
                sheet.title = "P0_Marca"
                sheet.append(["Linea PowerView", "Weighted VOL1_P"])
                sheet.append(["Base", "Weighted VOL1_P"])
                serial = to_excel(datetime(2024, 2, 1), epoch)
                sheet.append([serial, serial])
                workbook.save(path)
                workbook.close()
                shown = self.studio.workbook(book["file"])
                self.assertEqual(shown["display_values"][2][0], "feb 2024")
                self.assertEqual(shown["values"][2][0], serial)
                self.assertEqual(shown["display_types"][2][1], "number")
                confidence = workbook.create_sheet("8_Categoria")
                confidence.append(["table"])
                confidence.append(["Marca", "Weighted PENET", "BUYERS"])
                confidence.append([serial, 10, 100])
                research_path = self.external / f"52_CARB_Epoch{epoch.year}.xlsx"
                workbook.save(research_path)
                research = self.studio.workbook(f"external/{research_path.name}", kind="research", sheet="8_Categoria")
                self.assertEqual(research["date_columns"], [{"column": 1, "data_row": 3}])
                self.assertEqual(research["display_values"][2][0], "feb 2024")
                self.assertEqual(research["values"][2][0], serial)

    def test_research_calendars_are_detected_from_block_headers_not_numeric_metrics(self):
        path = self.external / "52_CARB_Marca 5w1h.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "5_Marca_2"
        sheet.append(["Measures = Weighted VOL1_P"])
        sheet.append(["Compras", "Ventas_1", "Moderno", None, None, None, "Ventas_1", "Moderno"])
        sheet.append(["Jan-24  ", 45323, 10, None, None, None, datetime(2024, 1, 1), 45323])
        sheet.append([datetime(2024, 2, 1), 45352, 20, None, None, None, "feb 24", 45352])
        sheet["G3"].number_format = "mmm-yy"
        workbook.save(path)
        workbook.close()
        before = path.read_bytes()
        shown = self.studio.workbook(f"external/{path.name}", kind="research", row=3, column=2)
        self.assertEqual(shown["date_columns"], [{"column": 1, "data_row": 3}, {"column": 7, "data_row": 3}])
        self.assertEqual(shown["display_values"][0][5], "ene 2024")
        self.assertEqual(shown["display_values"][1][5], "feb 2024")
        self.assertEqual(shown["values"][0][0], 45323)
        self.assertEqual(shown["display_types"][0][0], "number")
        self.assertEqual(shown["values"][0][5], "2024-01-01")
        self.assertEqual(shown["display_types"][0][6], "number")
        self.assertEqual(path.read_bytes(), before)

    def test_confidence_calendar_skips_brand_header_and_non_calendar_segments(self):
        path = self.external / "52_CARB_Marca.xlsx"
        workbook = Workbook()
        confidence = workbook.active
        confidence.title = "8_Categoria"
        confidence.append(["table", "Weighted PENET"])
        confidence.append(["Marca", 10])
        confidence.append(["Jan-24  ", 45323])
        confidence.append([datetime(2024, 2, 1), 45352])
        confidence.append(["=A4", 45380])
        confidence.append(["Otra marca"])
        confidence.append(["mar 24", 10, 100])
        confidence.append(["abr 24"])  # Calendar stays readable before observations are filled.
        for name in ["2_Marca_K", "7_Categoria_Canal", "README"]:
            sheet = workbook.create_sheet(name)
            sheet.append(["Jan-24", "MAT Sep-24", 45323])
        workbook.save(path)
        workbook.close()
        shown = self.studio.workbook(f"external/{path.name}", kind="research")
        self.assertEqual(shown["date_columns"], [{"column": 1, "data_row": 3}])
        self.assertEqual(shown["display_values"][1][0], "Marca")
        self.assertEqual(shown["display_types"][1][0], "text")
        self.assertEqual(shown["display_values"][2][0], "ene 2024")
        self.assertEqual(shown["display_types"][4][0], "formula")
        self.assertEqual(shown["display_values"][5][0], "Otra marca")
        self.assertEqual(shown["display_types"][5][0], "text")
        self.assertEqual(shown["display_values"][6][0], "mar 2024")
        self.assertEqual(shown["display_values"][7][0], "abr 2024")
        self.assertEqual(shown["display_types"][7][0], "date")
        for name in ["2_Marca_K", "7_Categoria_Canal", "README"]:
            with self.subTest(sheet=name):
                shown = self.studio.workbook(f"external/{path.name}", kind="research", sheet=name)
                self.assertEqual(shown["date_columns"], [])
                self.assertEqual(shown["display_values"][0], ["Jan-24", "MAT Sep-24", "45323"])

    def test_context_metadata_is_shared_by_file_list_and_workbook_without_invented_labels(self):
        book = self.create(manufacturer="Fabricante_con_guion")
        metadata = book["metadata"]
        self.assertTrue(metadata["recognized"])
        self.assertEqual(metadata["country"], "México")
        self.assertEqual(metadata["country_code"], "52")
        self.assertEqual(metadata["category"], "Bebidas Gaseosas")
        self.assertEqual(metadata["category_code"], "CARB")
        self.assertEqual(metadata["basket"], "Bebidas")
        self.assertEqual(metadata["manufacturer"], "Fabricante_con_guion")
        listed = next(entry for entry in self.studio.files()["coverage"] if entry["file"] == book["file"])
        self.assertEqual(listed["metadata"], metadata)
        self.assertEqual(self.studio._metadata("52_CARB_Marca_5w1h.xlsx", "research")["manufacturer"], "Marca")
        self.assertEqual(self.studio._metadata("52_CARB_Marca 5w1h.xlsx", "research")["manufacturer"], "Marca")
        for name in ["Datos.xlsx", "98_NOPE_Desconocido.xlsx"]:
            with self.subTest(name=name):
                unknown = self.studio._metadata(name, "coverage")
                self.assertFalse(unknown["recognized"])
                self.assertIsNone(unknown["country"])
                self.assertIsNone(unknown["category"])


if __name__ == "__main__":
    unittest.main()
