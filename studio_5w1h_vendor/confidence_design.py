"""Presentation-only confidence table shared by the browser and report worker.

The cutoffs below are fixed visual bands, not statistical acceptance criteria.
The source calculations and their unrounded values remain in report_engine.
"""
from __future__ import annotations

import math
import re
from decimal import Decimal, ROUND_HALF_UP


ERROR_BANDS = (
    (0.01, "123C30", "58DEAB", "E3F4EA"),
    (0.05, "173C45", "D4E55D", "F0F3D7"),
    (0.10, "593D0D", "F9CE5C", "FFF2CF"),
    (0.20, "6D2D14", "FFAD58", "FCE8D8"),
    (0.35, "6A2029", "FF7A80", "FBE4E6"),
    (0.51, "9D0B16", "EF1B2D", "FFD3D6"),
    (math.inf, "FFFFFF", "FF465D", "151419"),
)
INK, POSITIVE, NEGATIVE = "183143", "168345", "C0392B"
REFERENCE_ROW_HEIGHTS = (.24, .19, .20, .21, .19, .23, .19)


def finite(value):
    return isinstance(value, (int, float)) and math.isfinite(value)


def error_style(value):
    """Accept the engine's ratio, never its rounded display percentage."""
    if not finite(value) or value < 0:
        return dict(ink="526878", accent="CBD5DF", background="F1F4F7", extreme=False)
    _, ink, accent, background = next(b for b in ERROR_BANDS if value <= b[0])
    return dict(ink=ink, accent=accent, background=background, extreme=value > .51)


def percent(value, *, ratio=False, signed=False):
    if not finite(value):
        return "—"
    value = value * 100 if ratio else value
    if signed:
        return f"{value:+.1f}%".replace("-", "−")
    return f"{value:.1f}%"


def buyers_label(value):
    """Show whole buyers while preserving the engine's unrounded average."""
    if not finite(value):
        return "—"
    return f"{Decimal(str(value)).quantize(Decimal('1'), rounding=ROUND_HALF_UP):,}"


def sign_color(value):
    return NEGATIVE if finite(value) and value < 0 else POSITIVE if finite(value) and value > 0 else INK


def penetration_label(row):
    match = re.match(r"(?:Marca\s+)?(.+?)\s*\((.+)\)\s*$", row.nivel_pen)
    if match:
        return match.group(1), match.group(2)
    if finite(row.pen12m) and 0 <= row.pen12m < 1:
        return "Fuera de rango", "Menor de 1%"
    return row.nivel_pen, ""


def position_label(row):
    if not all(finite(v) for v in (row.var_fabricante, row.lim_inf, row.lim_sup)):
        return "—"
    if row.var_fabricante < row.lim_inf:
        return "Por debajo"
    if row.var_fabricante > row.lim_sup:
        return "Por encima"
    return "Dentro"


def interval_axis(rows):
    values = [0.0] + [v for r in rows for v in (r.var_fabricante, r.mat_sep25, r.lim_inf, r.lim_sup) if finite(v)]
    low, high = min(values), max(values)
    span = high - low or 1
    target = span / 5
    power = 10 ** math.floor(math.log10(target))
    step = next(n * power for n in (1, 2, 5, 10) if n * power >= target)
    low, high = math.floor(low / step) * step, math.ceil(high / step) * step
    if low == high:
        high = low + step
    return dict(min=low, max=high, ticks=[low + i * step for i in range(round((high - low) / step) + 1)])


def confidence_view(rows):
    return dict(axis=interval_axis(rows), rows=[dict(
        buyers=r.compradores_prom, brand=r.marca, manufacturer=r.var_fabricante, worldpanel=r.mat_sep25,
        lower=r.lim_inf, upper=r.lim_sup, position=position_label(r),
        penetration=list(penetration_label(r)), error=r.error_muestral,
        error_style=error_style(r.error_muestral),
    ) for r in rows])


def page_capacity(max_height_in):
    return max(1, min(8, math.floor((max_height_in - 1.3) / .64)))


def slide_layout(slide_height_in, table_top_in):
    """Reserve readable reference rows, the comment box and the template's logo."""
    reference_top = slide_height_in - .18 - sum(REFERENCE_ROW_HEIGHTS)
    comment_top = slide_height_in - 3.5
    table_height = comment_top - table_top_in
    if table_height < 1.94:
        raise ValueError("La diapositiva no tiene altura suficiente para la tabla y su referencia.")
    return dict(reference_top=reference_top, comment_top=comment_top,
                table_height=table_height, capacity=page_capacity(table_height))


def add_table(slide, rows, brand_colors, left, top, width_emu, max_height_in=None):
    """Native table, data marks and error signals. No raster fallback or legend ranges."""
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt

    if not rows:
        return None
    if max_height_in and len(rows) > page_capacity(max_height_in):
        raise ValueError("La tabla de intervalos necesita distribuirse en varias diapositivas.")
    width, x0, y0 = int(width_emu), int(left), int(top)
    scale = width / int(Inches(12.5))
    h0, h1, body_h = [int(Inches(v)) for v in (.28, .46, .64)]
    shape = slide.shapes.add_table(len(rows) + 2, 8, x0, y0, width, h0 + h1 + body_h * len(rows))
    shape.name = "Intervalos de confianza · tabla editable"
    table = shape.table
    ratios = [90, 150, 95, 95, 310, 110, 225, 125]
    widths = [round(width * v / 1200) for v in ratios]
    widths[-1] = width - sum(widths[:-1])
    for column, value in zip(table.columns, widths):
        column.width = value
    table.rows[0].height, table.rows[1].height = h0, h1
    for row in list(table.rows)[2:]:
        row.height = body_h

    def set_cell(cell, text, bg, color=INK, bold=False, size=11, align=PP_ALIGN.CENTER):
        cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string(bg)
        cell.margin_top = cell.margin_bottom = 0
        cell.margin_left = cell.margin_right = Pt(4)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame; tf.clear(); tf.word_wrap = True
        p = tf.paragraphs[0]; p.alignment = align
        p.space_before = p.space_after = Pt(0); p.line_spacing = 1.0
        run = p.add_run(); run.text = text
        run.font.name = "Arial"; run.font.size = Pt(size * min(1, scale))
        run.font.bold = bold; run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
        props = cell._tc.get_or_add_tcPr()
        for edge in ("lnL", "lnR", "lnT", "lnB"):
            for previous in list(props.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}" + edge)):
                props.remove(previous)
            line = OxmlElement("a:" + edge); line.set("w", "7620")
            fill = OxmlElement("a:solidFill"); rgb = OxmlElement("a:srgbClr"); rgb.set("val", "FFFFFF")
            fill.append(rgb); line.append(fill); props.append(line)

    for r in (0, 1):
        for c in range(8):
            set_cell(table.cell(r, c), "", "214E76" if r == 0 else "EAF1F6", "FFFFFF" if r == 0 else INK, True, 10.5)
    table.cell(0, 0).merge(table.cell(1, 0))
    table.cell(0, 1).merge(table.cell(1, 1))
    table.cell(0, 2).merge(table.cell(0, 5))
    table.cell(0, 6).merge(table.cell(0, 7))
    for c, text in ((0, "Compradores\nprom. 12 meses"), (1, "Marca"), (2, "Comparación MAT (%)"), (6, "Detalle de muestra")):
        set_cell(table.cell(0, c), text, "214E76", "CAD7E2" if c == 0 else "FFFFFF", c != 0, 9 if c == 0 else 12, PP_ALIGN.LEFT if c == 1 else PP_ALIGN.CENTER)
    for c, text in enumerate(("", "", "Fabricante", "Worldpanel", "Intervalo WP", "Fabricante\nvs. intervalo", "Nivel de\npenetración", "Error muestral")):
        if c > 1:
            set_cell(table.cell(1, c), text, "EAF1F6", bold=True, size=10.5)

    def rect(name, x, y, w, h, color):
        s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(x), int(y), max(1, int(w)), max(1, int(h)))
        s.name = name; s.fill.solid(); s.fill.fore_color.rgb = RGBColor.from_string(color); s.line.fill.background()
        return s

    def line(name, x, y, w, h, color, dashed=False):
        from pptx.enum.shapes import MSO_CONNECTOR
        s = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(x), int(y), int(x + w), int(y + h))
        s.name = name; s.line.color.rgb = RGBColor.from_string(color); s.line.width = Pt(.7)
        if dashed:
            s.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        return s

    def textbox(name, text, x, y, w, h, color=INK, size=9, bold=False):
        s = slide.shapes.add_textbox(int(x), int(y), max(1, int(w)), max(1, int(h))); s.name = name
        tf = s.text_frame; tf.clear(); tf.word_wrap = False
        tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; p.space_before = p.space_after = Pt(0)
        r = p.add_run(); r.text = text; r.font.name = "Arial"; r.font.size = Pt(size * min(1, scale))
        r.font.bold = bold; r.font.color.rgb = RGBColor.from_string(color)
        return s

    axis = interval_axis(rows)
    col_x = [x0 + sum(widths[:c]) for c in range(8)]
    plot_x, plot_w = col_x[4] + widths[4] * .09, widths[4] * .82
    def project(v):
        return plot_x + (v - axis["min"]) / (axis["max"] - axis["min"]) * plot_w
    for i, row in enumerate(rows):
        r = i + 2; y = y0 + h0 + h1 + i * body_h
        band = "FFFFFF" if i % 2 == 0 else "F3F7FA"
        for c in range(8):
            set_cell(table.cell(r, c), "", band)
        set_cell(table.cell(r, 0), buyers_label(row.compradores_prom), band, "6B7E8A", False, 10)
        set_cell(table.cell(r, 1), row.marca, band, brand_colors[i], True, 12, PP_ALIGN.LEFT)
        for c, value, bg in ((2, row.var_fabricante, "FCF6DF"), (3, row.mat_sep25, "EAF1F6")):
            set_cell(table.cell(r, c), percent(value, signed=True), bg, sign_color(value), True, 12)
        position = position_label(row)
        set_cell(table.cell(r, 5), position, band, "236B58" if position == "Dentro" else "79571C", size=10)
        level, detail = penetration_label(row)
        set_cell(table.cell(r, 6), level + ("\n" + detail if detail else ""), band, "000000", size=11)
        tf = table.cell(r, 6).text_frame
        if detail:
            # Keep the penetration range bold, independently of the error bands.
            tf.clear()
            for n, text in enumerate((level, detail)):
                p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
                p.alignment = PP_ALIGN.CENTER; p.space_before = p.space_after = Pt(0)
                run = p.add_run(); run.text = text; run.font.name = "Arial"; run.font.size = Pt(11 * min(1, scale))
                run.font.bold = bool(n); run.font.color.rgb = RGBColor(0, 0, 0)
        style = error_style(row.error_muestral)
        set_cell(table.cell(r, 7), percent(row.error_muestral, ratio=True), style["background"], style["ink"], True, 13)
        pad = int(Inches(.06)); error_x = col_x[7]; error_w = widths[7]
        rect(f"{row.marca}: error señal vertical", error_x + pad, y + pad, Inches(.025), body_h - 2 * pad, style["accent"] if style["extreme"] else style["ink"])
        rect(f"{row.marca}: error señal horizontal", error_x + 2 * pad, y + body_h - pad - Inches(.025), error_w - 3 * pad, Inches(.025), style["accent"])
        if style["extreme"]:
            warning = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, error_x + 2 * pad, y + int(Inches(.14)), int(Inches(.10)), int(Inches(.10)))
            warning.name = row.marca + ": error extremo"; warning.fill.background()
            warning.line.color.rgb = RGBColor.from_string(style["accent"]); warning.line.width = Pt(.8)
            textbox(row.marca + ": alerta de error", "!", error_x + 2 * pad, y + int(Inches(.155)), Inches(.10), Inches(.08), style["accent"], 5, True)
            table.cell(r, 7).margin_left = int(Inches(.18))
        center = y + int(body_h * .32)
        line(row.marca + ": eje", plot_x, center, plot_w, 0, "CBD5DF")
        line(row.marca + ": cero", project(0), y + body_h * .09, 0, body_h * .54, "7C8F9B", True)
        if finite(row.lim_inf) and finite(row.lim_sup):
            a, b = project(row.lim_inf), project(row.lim_sup)
            rect(row.marca + ": intervalo WP", a, center - Inches(.035), b - a, Inches(.07), "B8D0DF")
            for endpoint, name in ((a, "inferior"), (b, "superior")):
                line(row.marca + ": límite " + name, endpoint, center - Inches(.055), 0, Inches(.11), "B8D0DF")
        for value, geometry, color, label in ((row.mat_sep25, MSO_SHAPE.OVAL, "293F52", "Worldpanel"), (row.var_fabricante, MSO_SHAPE.DIAMOND, "BF9900", "Fabricante")):
            if finite(value):
                d = int(Inches(.065 if label == "Worldpanel" else .085))
                mark = slide.shapes.add_shape(geometry, int(project(value) - d / 2), int(center - d / 2), d, d)
                mark.name = row.marca + ": " + label; mark.fill.solid(); mark.fill.fore_color.rgb = RGBColor.from_string(color); mark.line.fill.background()
        label = textbox(row.marca + ": límites WP", "", col_x[4], y + body_h * .65, widths[4], body_h * .31)
        p = label.text_frame.paragraphs[0]
        for value, color in ((percent(row.lim_inf, signed=True), sign_color(row.lim_inf)), (" a ", "526878"), (percent(row.lim_sup, signed=True), sign_color(row.lim_sup))):
            run = p.add_run(); run.text = value; run.font.name = "Arial"; run.font.size = Pt(9 * min(1, scale)); run.font.color.rgb = RGBColor.from_string(color)
    bottom = y0 + h0 + h1 + len(rows) * body_h
    for tick in axis["ticks"]:
        textbox("Escala intervalo " + str(tick), f"{tick:g}%", project(tick) - Inches(.23), bottom + Inches(.04), Inches(.46), Inches(.18), "526878", 8)
    for i, (label, color) in enumerate((("Intervalo WP", "B8D0DF"), ("Worldpanel", "293F52"), ("Fabricante", "BF9900"))):
        lx = x0 + int(Inches(i * 1.65))
        if label == "Intervalo WP":
            rect("Leyenda " + label, lx, bottom + Inches(.35), Inches(.14), Inches(.05), color)
        else:
            marker = slide.shapes.add_shape(MSO_SHAPE.OVAL if label == "Worldpanel" else MSO_SHAPE.DIAMOND,
                                           int(lx), int(bottom + Inches(.32)), int(Inches(.07)), int(Inches(.07)))
            marker.name = "Leyenda " + label; marker.fill.solid()
            marker.fill.fore_color.rgb = RGBColor.from_string(color); marker.line.fill.background()
        textbox("Leyenda texto " + label, label, lx + Inches(.18), bottom + Inches(.28), Inches(1.35), Inches(.22), "526878", 9)
    return shape


# This is the user's reference, not a rule used by the confidence calculations.
# Its original 100-case boundary is retained verbatim.
BUYER_REFERENCE = (
    ("—", "Menos de 70 casos", "Información NO disponible", "FCE1E4", "B21D32"),
    ("Nivel 1", "Entre 70 y 100 casos", "Variables relativas (Share de la cadena, Penetración, Ranking)", "FAE9DE", "97502A"),
    ("Nivel 2", "Entre 100 y 200 casos", "Variables medias (Compra Media, Compra por ocasión, Gasto Medio, Frecuencia)", "FFF5D6", "8A6715"),
    ("Nivel 3", "Entre 201 y 300 casos", "Variables absolutas (Volumen, Gasto Total)", "F1F2D4", "717322"),
    ("Nivel 4", "Entre 301 y 400 casos", "Perfiles por Demográficos / Análisis Especiales (RTR’s) / Variaciones de volumen entre períodos", "E2F0E9", "28735D"),
    ("Nivel 5", "Más de 400 casos", "Demográficos por Canal/Cadena", "D2EBD9", "276B41"),
)


def add_buyer_reference(slide, left, top, width_emu):
    """Editable lookup footer under the comment box; no sampling-error legend."""
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt

    reference_width = round(int(width_emu) * .72)
    reference_left = int(left) + (int(width_emu) - reference_width) // 2
    shape = slide.shapes.add_table(7, 3, reference_left, int(top), reference_width, Inches(sum(REFERENCE_ROW_HEIGHTS)))
    shape.name = "Referencia de compradores · tamaños muestrales · editable"
    table = shape.table
    widths = [int(reference_width * .075), int(reference_width * .19)]
    widths.append(reference_width - sum(widths))
    for c, width in zip(table.columns, widths):
        c.width = width
    for row, height in zip(table.rows, REFERENCE_ROW_HEIGHTS):
        row.height = Inches(height)

    def cell(r, c, text, fill, ink=INK, bold=False, size=9, center=False):
        target = table.cell(r, c)
        target.fill.solid(); target.fill.fore_color.rgb = RGBColor.from_string(fill)
        target.margin_top = target.margin_bottom = 0
        target.margin_left = target.margin_right = Pt(5)
        target.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = target.text_frame; tf.clear(); tf.word_wrap = True
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER if center else PP_ALIGN.LEFT
        p.space_before = p.space_after = Pt(0); p.line_spacing = 1.0
        run = p.add_run(); run.text = text
        run.font.name = "Arial"; run.font.size = Pt(size); run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(ink)
        props = target._tc.get_or_add_tcPr()
        for edge in ("lnL", "lnR", "lnT", "lnB"):
            for old in list(props.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}" + edge)):
                props.remove(old)
            line = OxmlElement("a:" + edge); line.set("w", "7620")
            fill_node = OxmlElement("a:solidFill"); rgb = OxmlElement("a:srgbClr"); rgb.set("val", "FFFFFF")
            fill_node.append(rgb); line.append(fill_node); props.append(line)

    table.cell(0, 1).merge(table.cell(0, 2))
    cell(0, 0, "Nivel", "214E76", "FFFFFF", True, center=True)
    cell(0, 1, "Tamaños muestrales para análisis · panel de consumidores Worldpanel by Numerator", "214E76", "FFFFFF", True, 9.5)
    for r, (level, cases, usage, background, accent) in enumerate(BUYER_REFERENCE, 1):
        cell(r, 0, level, accent, "FFFFFF", True, center=True)
        cell(r, 1, cases, background, accent, True)
        cell(r, 2, usage, background)
    return shape
