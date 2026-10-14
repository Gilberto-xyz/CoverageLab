"""Compatibilidad de auxiliares y contrato de los gráficos adicionales del PPT."""
import json
from pathlib import Path
import tempfile
import types
import unittest
from unittest import mock

import pipeline_explorer as explorer
from delivery_files import auxiliary, organize_auxiliaries


class DeliveryExtrasTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.core = explorer.engine()

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=explorer.ROOT / "outputs")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.directory = self.root / "Output" / "Entrega"
        self.directory.mkdir(parents=True)

    def test_old_delivery_is_migrated_and_recognized_again(self):
        output = self.directory / "Template.xlsx"
        output.write_bytes(b"existing-deliverable")
        job = {"id": "old", "status": "complete", "directory": str(self.directory),
               "started_at": "2026-10-13T12:00:00-06:00", "file": "52_CARB_Test.xlsx",
               "outputs": [{"label": "template", "path": str(output), "name": output.name}]}
        for name, value in (("job.json", job), ("selection.json", {"selections": {"A": 2}}),
                            ("result.json", {"directory": str(self.directory)})):
            (self.directory / name).write_text(json.dumps(value), encoding="utf-8")
        (self.directory / "generation.log").write_text("Registro conservado", encoding="utf-8")
        for _ in range(2):
            state = explorer.ExplorerState(self.root)
            self.assertEqual(state.job("old")["log"], "Registro conservado")
            self.assertEqual(state.output_path(state.job("old"), "selection"), self.directory / "_sistema" / "selection.json")
        self.assertEqual(output.read_bytes(), b"existing-deliverable")
        self.assertEqual({path.name for path in self.directory.iterdir()}, {"_sistema", "Template.xlsx"})

    def test_old_read_only_auxiliary_still_reads_in_place(self):
        old = self.directory / "selection.json"
        old.write_text("{}", encoding="utf-8")
        with mock.patch.object(Path, "replace", side_effect=PermissionError):
            self.assertEqual(organize_auxiliaries(self.directory, [old.name]), {})
        self.assertEqual(auxiliary(self.directory, old.name), old)

    def test_old_research_delivery_retargets_auxiliary_links_and_work_copy(self):
        outputs = []
        for label, name in (("Excel", "Informe.xlsx"), ("PPT", "Informe.pptx"),
                            ("Provenance", "investigacion_5w1h.json"), ("Registro", "generacion_5w1h.log")):
            path = self.directory / name
            path.write_bytes(b"preserved")
            outputs.append({"label": label, "name": name, "path": str(path)})
        copied = self.directory / "_input"
        copied.mkdir()
        (copied / "Entrada.xlsx").write_bytes(b"original copy")
        (copied / "request.json").write_text(json.dumps({"source": str(copied / "Entrada.xlsx")}), encoding="utf-8")
        (self.directory / "_result.json").write_text("{}", encoding="utf-8")
        job = {"id":"research-old", "kind":"research", "file":"local/52_CARB_Test.xlsx",
               "fingerprint":"f"*64, "status":"complete", "directory":str(self.directory),
               "started_at":"2026-10-13T12:00:00-06:00", "outputs":outputs}
        (self.directory / "research-job.json").write_text(json.dumps(job), encoding="utf-8")
        for _ in range(2):
            state = explorer.ExplorerState(self.root)
            recovered = state.workflow.research_jobs[job["id"]]
            self.assertTrue(all(Path(out["path"]).is_file() for out in recovered["outputs"]))
        self.assertEqual({path.name for path in self.directory.iterdir()}, {"Informe.xlsx", "Informe.pptx", "_sistema"})
        request = json.loads((self.directory / "_sistema" / "_input" / "request.json").read_text(encoding="utf-8"))
        self.assertEqual(Path(request["source"]).read_bytes(), b"original copy")

    def test_internal_path_escape_is_rejected(self):
        with self.assertRaises(ValueError):
            auxiliary(self.directory, "../outside.json", write=True)
        outside = self.root / "outside"
        outside.mkdir()
        try:
            (self.directory / "_sistema").symlink_to(outside, target_is_directory=True)
        except OSError:
            self.skipTest("Esta sesión no permite crear enlaces simbólicos")
        with self.assertRaises(ValueError):
            auxiliary(self.directory, "job.json", write=True)

    def test_extras_reject_unknown_or_excluded_sheet_and_duplicate_keys(self):
        for value in (None, [], {"Excluded": ["frequency"]}, {"A": ["unknown"]},
                      {"A": ["frequency", "frequency"]}, {"A": "frequency"}):
            with self.subTest(value=value), self.assertRaises(ValueError):
                self.core.validate_ppt_extra_charts(value, ["A"])
        self.assertEqual(self.core.validate_ppt_extra_charts({}, ["A"]), {})

    def test_every_kpi_is_editable_and_default_adds_no_slides(self):
        core = self.core
        ppt = core.Presentation(explorer.ROOT / "Modelo_PPT.pptx")
        builder = types.SimpleNamespace(ppt=ppt, lang_index=2)
        frame = core.pd.DataFrame({core.COL_DATA:core.pd.date_range("2024-01-01", periods=14, freq="MS")})
        for column, _ in core.PPT_EXTRA_KPIS.values():
            frame[column] = list(range(1, 15))
        frame.loc[0, core.COL_FREQ] = float("nan")
        count = len(ppt.slides)
        self.assertEqual(core.add_extra_kpi_slides(builder, frame, [], "Marca", "kg"), 0)
        self.assertEqual(len(ppt.slides), count)
        self.assertEqual(core.add_extra_kpi_slides(builder, frame, list(core.PPT_EXTRA_KPIS), "Marca", "kg"), 5)
        ppt.save(self.root / "proof.pptx")
        reloaded = core.Presentation(self.root / "proof.pptx")
        for slide in list(reloaded.slides)[count:]:
            chart = next(shape.chart for shape in slide.shapes if shape.has_chart)
            self.assertEqual(chart.series[0].values[-1], 14)
            texts = " ".join(shape.text for shape in slide.shapes if shape.has_text_frame)
            self.assertIn("Promedio 12 meses: 8" if slide.name.startswith("KPI:buyers_series") else "Promedio 12 meses: 8.50", texts)
            self.assertIn("Promedio total", texts)
            self.assertTrue(all(shape.left >= 0 and shape.top >= 0 and shape.left+shape.width <= ppt.slide_width and shape.top+shape.height <= ppt.slide_height for shape in slide.shapes))

    def test_annual_comparison_preserves_pipeline_yoy_coverage_and_editable_data(self):
        from io import BytesIO
        from openpyxl import load_workbook
        core = self.core
        frame = core.pd.DataFrame({core.COL_DATA:core.pd.date_range("2023-01-01", periods=36, freq="MS"),
            core.COL_SELL_OUT:list(range(100,136)), core.COL_SELL_IN:list(range(200,236))})
        coverage = core.compute_coverage_dataframe(frame, "Peru", "Relativa", True)
        p = 3
        selected = coverage[f"P{p}"]
        rows = core.annual_mat_rows(frame, p, selected)
        self.assertIsNone(rows[0]["sell_in"])
        self.assertIsNone(rows[1]["sell_in_yoy"])
        self.assertEqual(rows[-1]["sell_in"], sum(frame[core.COL_SELL_IN].iloc[-15:-3]))
        self.assertEqual(rows[-1]["coverage"], selected.iloc[-1])
        self.assertAlmostEqual(rows[-1]["sell_in_yoy"], (rows[-1]["sell_in"]/rows[-2]["sell_in"]-1)*100)
        self.assertEqual(core.validate_ppt_extra_charts({"A":["annual_mat"]}, ["A"]), {"A":["annual_mat"]})
        ppt = core.Presentation(explorer.ROOT / "Modelo_PPT.pptx")
        builder = types.SimpleNamespace(ppt=ppt, lang_index=2)
        self.assertEqual(core.add_extra_kpi_slides(builder, frame, ["annual_mat"], "Marca", "kg", p, selected), 1)
        ppt.save(self.root / "annual.pptx")
        chart = next(shape.chart for shape in core.Presentation(self.root / "annual.pptx").slides[-1].shapes if shape.has_chart)
        self.assertEqual(len(chart.plots), 2)
        self.assertTrue(chart.plots[0]._element.tag.endswith('}barChart'))
        self.assertTrue(chart.plots[1]._element.tag.endswith('}lineChart'))
        self.assertEqual(chart.series[0].values, tuple(row["sell_in"] for row in rows[1:]))
        self.assertEqual(chart.series[1].values, tuple(row["worldpanel"] for row in rows[1:]))
        self.assertEqual(chart.series[2].values, tuple(row["coverage"] for row in rows[1:]))
        axes = chart._chartSpace.chart.plotArea.xpath('c:valAx')
        self.assertEqual(len(axes), 2)
        self.assertEqual(axes[-1].xpath('c:axPos')[0].get('val'), 'r')
        self.assertEqual(axes[-1].xpath('c:numFmt')[0].get('formatCode'), '0"%"')
        table = next(shape.table for shape in core.Presentation(self.root / "annual.pptx").slides[-1].shapes if shape.has_table)
        self.assertEqual(table.cell(1,2).text, f"{rows[-1]['sell_in_yoy']:+.1f}%")
        self.assertEqual(table.cell(1,1).text, '—')
        embedded = load_workbook(BytesIO(chart.part.chart_workbook.xlsx_part.blob), data_only=True)
        self.assertEqual(embedded.active.cell(3, 4).value, rows[-1]["coverage"])
        self.assertEqual(core.add_extra_kpi_slides(builder, frame, ["annual_mat"], "Marca", "kg", p,
            core.pd.Series([float('nan')]*len(frame))), 0)


if __name__ == "__main__":
    unittest.main()
