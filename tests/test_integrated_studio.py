"""Real workbook and HTTP checks for the integrated capture → analysis workflow."""
from __future__ import annotations

import base64
import contextlib
from datetime import datetime
import io
import json
from pathlib import Path
import shutil
import subprocess
import tempfile
import threading
import time
import unittest
from unittest import mock
import urllib.error
import urllib.parse
import urllib.request
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

import pipeline_explorer as explorer


class IntegratedStudioTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.core = explorer.engine()

    def setUp(self):
        (explorer.ROOT / "outputs").mkdir(exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(prefix="integrated_test_", dir=explorer.ROOT / "outputs")
        self.root = Path(self.temp.name)
        self.external = self.root / "original_5w1h"
        self.external.mkdir()
        self.state = explorer.ExplorerState(self.root, research_root=self.external)

    def tearDown(self):
        self.temp.cleanup()

    @staticmethod
    def coverage_spec(**updates):
        spec = {"kind": "coverage", "country_code": "52", "category_code": "CARB",
                "manufacturer": "Captura", "brands": ["Cola", "Limón"]}
        return spec | updates

    def create_coverage(self, **updates):
        return self.state.workflow.create_input(self.coverage_spec(**updates))

    @staticmethod
    def monthly_rows(count=36):
        rows = []
        for index in range(count):
            year, month = 2023 + index // 12, index % 12 + 1
            # Excel paste is textual; locale decimal commas must retain numeric meaning.
            rows.append([f"{year}-{month:02d}-01", str(1000 + index * 10), "20,5", "2,5",
                         "1,25", "2", str(500 + index), str(1600 + index * 13)])
        return rows

    def fill_coverage(self, workbook):
        name = workbook["sheets"][0]["name"]
        return self.state.workflow.save_input({"file": workbook["file"], "kind": "coverage",
            "fingerprint": workbook["fingerprint"], "edits": [
                {"sheet": name, "row": 1, "column": 2, "values": [["Weighted VOL1_P"]]},
                {"sheet": name, "row": 3, "column": 1, "values": self.monthly_rows()},
            ]})

    def analyze_coverage(self, name):
        with contextlib.redirect_stdout(io.StringIO()):
            return self.state.analyze({"file": name, "animate_trend_pipeline": False})

    def make_research(self, destination=None, filename="52_CARB_Investigación.xlsx", name="5_Cola_2"):
        destination = destination or self.root / "inputs" / "5w1h"
        destination.mkdir(parents=True, exist_ok=True)
        path = destination / filename
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = name
        # Real PowerView data often has metadata above the table and a separated sales block.
        sheet.append(["CCSD_BR = Cola\\Total CCSD_BR"])
        sheet.append(["Measures = Weighted VOL1_P"])
        sheet.append(["Compras", "Moderno", "Tradicional", None, "Ventas_1", "Moderno", "Tradicional"])
        for index in range(36):
            date = datetime(2023 + index // 12, index % 12 + 1, 1)
            sheet.append([date, 100 + index * 2, 200 + index * 3, None,
                          date, 300 + index * 5, 500 + index * 7])
        sheet["B4"].fill = PatternFill("solid", fgColor="FFDDEEFF")
        sheet["B4"].number_format = "0.00"
        workbook.save(path)
        workbook.close()
        return path

    def test_create_capture_has_blank_brand_sheets_and_never_overwrites(self):
        first = self.create_coverage()
        self.assertTrue(first["file"].startswith("52_CARB_Captura"))
        self.assertEqual(len(first["sheets"]), 2)
        self.assertFalse(first["validation"]["ready"])
        path = self.root / first["file"]
        original = path.read_bytes()
        with contextlib.closing(load_workbook(path)) as workbook:
            self.assertEqual(workbook.sheetnames, ["P0_Cola", "P0_Limón"])
            self.assertEqual(workbook["P0_Cola"]["A1"].value, "Linea PowerView")
            self.assertEqual(workbook["P0_Cola"]["H2"].value, "SELL-IN: Cola")
            self.assertIsNone(workbook["P0_Cola"]["B3"].value)
        second = self.create_coverage()
        self.assertNotEqual(first["file"], second["file"])
        self.assertEqual(path.read_bytes(), original)

    def test_pasted_capture_survives_reload_and_matches_real_pipeline_calculations(self):
        created = self.create_coverage(brands=["Cola"])
        saved = self.fill_coverage(created)
        reloaded = self.state.workflow.workbook({"file": saved["file"], "kind": "coverage", "row": 3})
        self.assertEqual(reloaded["fingerprint"], saved["fingerprint"])
        self.assertNotEqual(saved["fingerprint"], created["fingerprint"])
        with contextlib.closing(load_workbook(self.root / saved["file"], data_only=True)) as workbook:
            self.assertEqual(workbook["P0_Cola"]["C3"].value, 20.5)
            self.assertEqual(workbook["P0_Cola"]["D3"].value, 2.5)
            self.assertEqual(workbook["P0_Cola"]["E3"].value, 1.25)
            self.assertEqual(workbook["P0_Cola"]["A3"].value, datetime(2023, 1, 1))
        analysis = self.analyze_coverage(saved["file"])
        sheet = analysis["sheets"][0]
        self.assertIsNone(sheet["error"])
        self.assertEqual(len(sheet["dates"]), 36)
        self.assertEqual(sheet["purchase_average"], [2.5] * 36)
        self.assertEqual(sheet["penetration"], [20.5] * 36)
        expected = round(sum(1000 + i * 10 for i in range(24, 36)) /
                         sum(1600 + i * 13 for i in range(21, 33)) * 100, 1)
        self.assertEqual(sheet["pipelines"][3]["coverage"], expected)
        self.assertEqual(sheet["pipelines"][3]["sell_in"][:3], [None] * 3)
        self.assertEqual(self.state.jobs, {})
        self.assertFalse((self.root / "Output").exists())
        json.dumps(analysis, allow_nan=False)

    def test_edit_invalidates_review_and_stale_save_is_atomic(self):
        saved = self.fill_coverage(self.create_coverage(brands=["Cola"]))
        analysis = self.analyze_coverage(saved["file"])
        path = self.root / saved["file"]
        modified = self.state.workflow.save_input({"file": saved["file"], "kind": "coverage",
            "fingerprint": saved["fingerprint"], "edits": [
                {"sheet": "P0_Cola", "row": 3, "column": 2, "values": [["1.234,5"]]}]})
        with self.assertRaises(ValueError):
            self.state.reviewed_selection({"analysis_id": analysis["id"], "selections": {"P0_Cola": 3}})
        expected_bytes = path.read_bytes()
        with self.assertRaises(ValueError):
            self.state.workflow.save_input({"file": saved["file"], "kind": "coverage",
                "fingerprint": saved["fingerprint"], "edits": [
                    {"sheet": "P0_Cola", "row": 3, "column": 2, "values": [["999"]]}]})
        self.assertEqual(path.read_bytes(), expected_bytes)
        self.assertNotEqual(modified["fingerprint"], saved["fingerprint"])
        refreshed = self.analyze_coverage(saved["file"])
        self.assertEqual(refreshed["sheets"][0]["sell_out"][0], 1234.5)
        self.assertNotEqual(refreshed["id"], analysis["id"])

    def test_formula_and_merged_cells_are_protected_without_partial_edits(self):
        created = self.create_coverage(brands=["Cola"])
        path = self.root / created["file"]
        with contextlib.closing(load_workbook(path)) as workbook:
            sheet = workbook["P0_Cola"]
            sheet["J3"] = "=SUM(B3:B5)"
            sheet["J3"].number_format = "0.00"
            sheet["J3"].fill = PatternFill("solid", fgColor="FFAABBCC")
            sheet.merge_cells("J5:K5")
            workbook.save(path)
        current = self.state.workflow.workbook({"file": created["file"], "kind": "coverage",
                                                "row": 3, "column": 10, "columns": 2, "limit": 3})
        self.assertEqual(current["values"][0][0], "=SUM(B3:B5)")
        self.assertTrue(current["readonly"][0][0])
        self.assertTrue(current["readonly"][2][1])
        before = path.read_bytes()
        for row, column in ((3, 10), (5, 11)):
            with self.subTest(row=row, column=column):
                with self.assertRaises(ValueError):
                    self.state.workflow.save_input({"file": created["file"], "kind": "coverage",
                        "fingerprint": current["fingerprint"], "edits": [
                            {"sheet": "P0_Cola", "row": 3, "column": 2, "values": [["123"]]},
                            {"sheet": "P0_Cola", "row": row, "column": column, "values": [["456"]]},
                        ]})
                self.assertEqual(path.read_bytes(), before)
        updated = self.state.workflow.save_input({"file": created["file"], "kind": "coverage",
            "fingerprint": current["fingerprint"], "edits": [
                {"sheet": "P0_Cola", "row": 3, "column": 2, "values": [["123"]]}]})
        self.assertNotEqual(updated["fingerprint"], current["fingerprint"])
        with contextlib.closing(load_workbook(path)) as workbook:
            self.assertEqual(workbook["P0_Cola"]["J3"].value, "=SUM(B3:B5)")
            self.assertEqual(workbook["P0_Cola"]["J3"].number_format, "0.00")
            self.assertEqual(workbook["P0_Cola"]["J3"].fill.fgColor.rgb, "FFAABBCC")

    def test_calendar_validation_agrees_with_pipeline_requirements(self):
        for problem in ("missing_month", "same_month_twice", "unordered"):
            with self.subTest(problem=problem):
                created = self.create_coverage(brands=["Cola"], manufacturer=problem)
                rows = self.monthly_rows()
                if problem == "missing_month":
                    for index in range(18, 36):
                        number = 2023 * 12 + index + 1
                        rows[index][0] = f"{number // 12}-{number % 12 + 1:02d}-01"
                elif problem == "same_month_twice":
                    rows[1][0] = "2023-01-15"
                else:
                    rows[0][0], rows[1][0] = rows[1][0], rows[0][0]
                saved = self.state.workflow.save_input({"file": created["file"], "kind": "coverage",
                    "fingerprint": created["fingerprint"], "edits": [
                        {"sheet": "P0_Cola", "row": 3, "column": 1, "values": rows}]})
                self.assertFalse(saved["validation"]["ready"])
                self.assertTrue(saved["validation"]["errors"])
                analysis = self.analyze_coverage(saved["file"])
                self.assertTrue(analysis["sheets"][0]["error"])

    def test_adding_brand_inherits_unit_and_duplicate_preserves_source_data(self):
        saved = self.fill_coverage(self.create_coverage(brands=["Cola"], unit="VOL1_P"))
        analysis = self.analyze_coverage(saved["file"])
        added = self.state.workflow.add_sheet({"file": saved["file"], "kind": "coverage",
            "fingerprint": saved["fingerprint"], "name": "Uva"})
        self.assertEqual(added["sheet"], "P0_Uva")
        with self.assertRaises(ValueError):
            self.state.reviewed_selection({"analysis_id": analysis["id"], "selections": {"P0_Cola": 0}})
        cloned = self.state.workflow.add_sheet({"file": saved["file"], "kind": "coverage",
            "fingerprint": added["fingerprint"], "name": "P3_Copia", "source_sheet": "P0_Cola"})
        with contextlib.closing(load_workbook(self.root / cloned["file"])) as workbook:
            self.assertEqual(workbook.sheetnames, ["P0_Cola", "P0_Uva", "P3_Copia"])
            self.assertEqual(workbook["P0_Uva"]["B1"].value, "Weighted VOL1_P")
            self.assertIsNone(workbook["P0_Uva"]["B3"].value)
            self.assertEqual(workbook["P3_Copia"]["B3"].value, workbook["P0_Cola"]["B3"].value)
        preview = self.analyze_coverage(saved["file"])
        copy = next(sheet for sheet in preview["sheets"] if sheet["name"] == "P3_Copia")
        self.assertIsNone(copy["error"])
        self.assertEqual(copy["default_pipeline"], 3)

    def test_invalid_or_overlapping_paste_does_not_write_any_cells(self):
        created = self.create_coverage(brands=["Cola"])
        path = self.root / created["file"]
        original = path.read_bytes()
        valid = {"sheet": "P0_Cola", "row": 3, "column": 2, "values": [["100"]]}
        for invalid in ({"sheet": "P0_Cola", "row": 4, "column": 2, "values": [["100", "2"], ["3"]]},
                        {"sheet": "P0_Cola", "row": 4, "column": 2, "values": [["=SUM(A1:A2)"]]},
                        {"sheet": "P0_Cola", "row": 4, "column": 2, "values": [["dato inválido"]]},
                        valid):
            with self.subTest(invalid=invalid):
                with self.assertRaises(ValueError):
                    self.state.workflow.save_input({"file": created["file"], "kind": "coverage",
                        "fingerprint": created["fingerprint"], "edits": [valid, invalid]})
                self.assertEqual(path.read_bytes(), original)

    def test_import_valid_workbook_and_reject_invalid_bytes_and_paths(self):
        source = self.make_research(destination=self.external)
        payload = {"kind": "research", "filename": source.name,
                   "content": base64.b64encode(source.read_bytes()).decode("ascii")}
        imported = self.state.workflow.import_input(payload)
        self.assertTrue(imported["file"].startswith("local/"))
        self.assertEqual(len(imported["sheets"]), 1)
        before = sorted(str(path.relative_to(self.root)) for path in self.root.rglob("*.xlsx"))
        for invalid in ({"filename": "../escape.xlsx"}, {"filename": "captura.xlsm"},
                        {"content": "%%%"}, {"content": base64.b64encode(b"not an Excel").decode("ascii")}):
            with self.subTest(invalid=invalid):
                with self.assertRaises(ValueError):
                    self.state.workflow.import_input(payload | invalid)
        self.assertEqual(sorted(str(path.relative_to(self.root)) for path in self.root.rglob("*.xlsx")), before)
        self.assertFalse((self.root.parent / "escape.xlsx").exists())

    def test_editing_external_research_creates_local_copy_and_preserves_original(self):
        path = self.make_research(destination=self.external)
        original = path.read_bytes()
        source_id = "external/" + path.name
        current = self.state.workflow.workbook({"file": source_id, "kind": "research", "sheet": "5_Cola_2"})
        copied = self.state.workflow.save_input({"file": source_id, "kind": "research",
            "fingerprint": current["fingerprint"], "edits": [
                {"sheet": "5_Cola_2", "row": 4, "column": 2, "values": [["987,5"]]}]})
        self.assertTrue(copied["file"].startswith("local/"))
        self.assertEqual(path.read_bytes(), original)
        self.assertEqual(copied["copied_from"], source_id)
        self.assertIn(copied["file"], {entry["id"] for entry in self.state.workflow.research.catalog()["files"]})
        with contextlib.closing(load_workbook(self.root / "inputs" / "5w1h" / copied["file"].split("/", 1)[1])) as workbook:
            self.assertEqual(workbook["5_Cola_2"]["B4"].value, '987,5')
            self.assertEqual(workbook["5_Cola_2"]["B4"].number_format, "0.00")
            self.assertEqual(workbook["5_Cola_2"]["B4"].fill.fgColor.rgb, "FFDDEEFF")
        analysis = self.state.workflow.research_analyze({"file": copied["file"], "sheet": "5_Cola_2"})
        self.assertTrue(analysis["ready"])
        monthly = next(chart for chart in analysis["charts"] if chart["id"] == "monthly")
        moderno = next(series for series in monthly["series"] if series["name"] == "Moderno · Compras")
        self.assertEqual(moderno["values"][0], 987.5)
        self.assertEqual(path.read_bytes(), original)

    def test_new_research_template_contains_no_demo_observations(self):
        with contextlib.redirect_stdout(io.StringIO()):
            created = self.state.workflow.create_input(self.coverage_spec(kind="research", brands=["Cola"],
                segments=["1", "2", "3-1", "4", "5-1", "5-2", "6", "6-1", "7", "8"],
                start_date="2024-01", months=36, unit="K"))
        self.assertTrue(created["file"].startswith("local/"))
        inspected = self.state.workflow.research.inspect(created["file"])
        self.assertTrue(inspected["requires_input"])
        self.assertGreaterEqual(len(inspected["sheets"]), 10)
        self.assertTrue(all(not sheet["has_data"] for sheet in inspected["sheets"]))
        for sheet in inspected["sheets"]:
            with self.subTest(segment=sheet["segment"]):
                preview = self.state.workflow.research_analyze({"file": created["file"], "sheet": sheet["name"]})
                self.assertFalse(preview["ready"])
                self.assertEqual(preview["charts"], [])
                self.assertEqual(preview["tables"], [])
                self.assertTrue(preview["warnings"])
        self.assertFalse((self.root / "Output").exists())

    def test_research_preview_matches_source_monthly_mat_and_full_total_shares(self):
        path = self.make_research()
        source_id = "local/" + path.name
        original = path.read_bytes()
        analysis = self.state.workflow.research_analyze({"file": source_id, "sheet": "5_Cola_2", "pipeline": 3})
        self.assertTrue(analysis["ready"])
        self.assertEqual(analysis["provenance"]["sha256"], explorer.fingerprint(path))
        charts = {chart["id"]: chart for chart in analysis["charts"]}
        monthly = {series["name"]: series for series in charts["monthly"]["series"]}
        self.assertEqual(monthly["Moderno · Compras"]["values"], [100 + i * 2 for i in range(36)])
        self.assertEqual(monthly["Moderno · Ventas"]["values"][:3], [None] * 3)
        self.assertEqual(monthly["Moderno · Ventas"]["values"][3:], [300 + i * 5 for i in range(33)])
        mat = {series["name"]: series for series in charts["mat"]["series"]}
        self.assertEqual(mat["Moderno · Compras"]["values"][-1], sum(100 + i * 2 for i in range(24, 36)))
        self.assertEqual(mat["Moderno · Ventas"]["values"][-1], sum(300 + i * 5 for i in range(21, 33)))
        shares = {series["name"]: series for series in charts["share-Compras"]["series"]}
        expected_share = sum(100 + i * 2 for i in range(24, 36)) / sum(300 + i * 5 for i in range(24, 36))
        self.assertAlmostEqual(shares["Moderno"]["values"][-1], expected_share, places=5)
        filtered = self.state.workflow.research_analyze({"file": source_id, "sheet": "5_Cola_2",
            "pipeline": 3, "start": "2025-01", "end": "2025-06", "elements": ["Moderno"]})
        filtered_charts = {chart["id"]: chart for chart in filtered["charts"]}
        self.assertEqual(filtered_charts["mat"]["dates"], [f"2025-{month:02d}" for month in range(1, 7)])
        filtered_mat = next(series for series in filtered_charts["mat"]["series"] if series["origin"] == "Compras")
        self.assertEqual(filtered_mat["values"][0], sum(100 + i * 2 for i in range(13, 25)))
        filtered_share = filtered_charts["share-Compras"]["series"][0]
        cut_share = sum(100 + i * 2 for i in range(18, 30)) / sum(300 + i * 5 for i in range(18, 30))
        self.assertAlmostEqual(filtered_share["values"][-1], cut_share, places=5)
        self.assertTrue(any("base" in warning for warning in filtered["warnings"]))
        self.assertEqual(path.read_bytes(), original)
        self.assertFalse((self.root / "Output").exists())
        json.dumps(filtered, allow_nan=False)

    def test_research_preview_rejects_invalid_identifiers_modes_and_filters(self):
        path = self.make_research()
        source_id = "local/" + path.name
        for invalid in ({"file": "local/../escape.xlsx"}, {"file": str(path)}, {"sheet": "SinDatos"},
                        {"mode": "4"}, {"pipeline": True}, {"pipeline": 7},
                        {"start": "2025-09", "end": "2024-01"}, {"elements": ["Inventado"]}):
            with self.subTest(invalid=invalid):
                with self.assertRaises(ValueError):
                    self.state.workflow.research_analyze({"file": source_id, "sheet": "5_Cola_2"} | invalid)
        self.assertEqual(self.state.workflow.research_jobs, {})
        self.assertFalse((self.root / "Output").exists())

    def test_research_edits_and_external_changes_invalidate_previous_export_review(self):
        path = self.make_research()
        source_id = "local/" + path.name
        review = self.state.workflow.research_analyze({"file": source_id, "sheet": "5_Cola_2"})
        current = self.state.workflow.workbook({"file": source_id, "kind": "research"})
        self.state.workflow.save_input({"file": source_id, "kind": "research",
            "fingerprint": current["fingerprint"], "edits": [
                {"sheet": "5_Cola_2", "row": 4, "column": 2, "values": [["200"]]}]})
        with self.assertRaises(ValueError):
            self.state.workflow.export_research({"analysis_id": review["id"]})
        review = self.state.workflow.research_analyze({"file": source_id, "sheet": "5_Cola_2"})
        with contextlib.closing(load_workbook(path)) as workbook:
            workbook["5_Cola_2"]["B4"] = 300
            workbook.save(path)
        with self.assertRaises(ValueError):
            self.state.workflow.export_research({"analysis_id": review["id"]})
        self.assertEqual(self.state.workflow.research_jobs, {})
        self.assertFalse((self.root / "Output").exists())

    def test_research_folder_uses_shared_foreground_helper_after_delivery_validation(self):
        path = self.make_research()
        source_id = "local/" + path.name
        directory = self.root / "Output" / "Investigación"
        directory.mkdir(parents=True)
        job = {"id": "folder-proof", "file": source_id, "status": "complete",
               "directory": str(directory), "outputs": [], "fingerprint": explorer.fingerprint(path)}
        self.state.workflow.research_jobs[job["id"]] = job
        with mock.patch("studio_workflow.open_folder", return_value={"foreground": True,
                "directory": str(directory), "message": "Carpeta en primer plano."}) as opener:
            result = self.state.workflow.open_research_output({"id": job["id"], "file": source_id, "output": "folder"})
            self.assertTrue(result["foreground"])
            opener.assert_called_once_with(directory.resolve())
            opener.reset_mock()
            for updates in ({"file": "external/other.xlsx"}, {"id": "not-registered"}):
                with self.subTest(updates=updates), self.assertRaises(ValueError):
                    self.state.workflow.open_research_output({"id": job["id"], "file": source_id,
                        "output": "folder"} | updates)
            for updates in ({"status": "running"}, {"directory": str(self.root.parent)},
                            {"directory": str(directory / "missing")}):
                self.state.workflow.research_jobs[job["id"]] = job | updates
                with self.subTest(job=updates), self.assertRaises(ValueError):
                    self.state.workflow.open_research_output({"id": job["id"], "file": source_id, "output": "folder"})
            opener.assert_not_called()

    def test_research_export_generates_real_excel_ppt_provenance_and_restorable_job(self):
        path = self.make_research()
        source_id = "local/" + path.name
        original = path.read_bytes()
        # A filtered exploration still exports the complete report explicitly reviewed by the UI.
        review = self.state.workflow.research_analyze({"file": source_id, "sheet": "5_Cola_2", "mode": "1",
                                                      "start": "2025-01", "elements": ["Moderno"]})
        job = self.state.workflow.export_research({"analysis_id": review["id"], "pipeline_overrides": {"Cola": 3}})
        deadline = time.monotonic() + 120
        while time.monotonic() < deadline:
            result = self.state.workflow.research_job(job["id"])
            if result["status"] != "running":
                break
            time.sleep(0.1)
        self.assertEqual(result["status"], "complete", result.get("error"))
        self.assertNotIn("start", result["options"])
        self.assertNotIn("elements", result["options"])
        self.assertEqual(set(result["available_outputs"]), {"Excel", "PPT", "Provenance", "Registro"})
        directory = Path(result["directory"])
        self.assertTrue(directory.is_relative_to(self.root / "Output"))
        self.assertNotEqual(directory, self.root / "Output")
        outputs = {entry["label"]: Path(entry["path"]) for entry in result["outputs"]}
        for label in ("Excel", "PPT"):
            with zipfile.ZipFile(outputs[label]) as archive:
                self.assertIsNone(archive.testzip())
                if label == "PPT":
                    slides = [name for name in archive.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml")]
                    self.assertGreater(len(slides), 1)
                    self.assertTrue(any(name.startswith("ppt/media/") for name in archive.namelist()))
                else:
                    self.assertTrue(any(name.startswith("xl/charts/chart") for name in archive.namelist()))
        with contextlib.closing(load_workbook(outputs["Excel"], data_only=False)) as workbook:
            formulas = [cell.value for sheet in workbook for row in sheet for cell in row if cell.data_type == "f"]
            self.assertTrue(formulas)
            self.assertIn("5_Cola_2", workbook.sheetnames)
        provenance = json.loads(outputs["Provenance"].read_text(encoding="utf-8"))
        self.assertEqual(provenance["source_sha256"], explorer.fingerprint(path))
        self.assertEqual(provenance["pipelines_applied"], {"5_Cola_2": 3})
        self.assertEqual(path.read_bytes(), original)
        restored = explorer.ExplorerState(self.root, research_root=self.external)
        recovered = restored.workflow.research_job(file_id=source_id)
        self.assertEqual(recovered["id"], job["id"])
        self.assertEqual(set(recovered["available_outputs"]), {"Excel", "PPT", "Provenance", "Registro"})
        with mock.patch("studio_workflow.os.startfile", create=True) as opener, mock.patch("studio_workflow.sys.platform", "win32"):
            self.state.workflow.open_research_output({"id": job["id"], "file": source_id, "output": "PPT"})
            opener.assert_called_once_with(str(outputs["PPT"]))
            with self.assertRaises(ValueError):
                self.state.workflow.open_research_output({"id": job["id"], "file": "external/other.xlsx", "output": "PPT"})

    def test_all_new_http_routes_require_same_local_session(self):
        server = explorer.ExplorerHTTPServer(("127.0.0.1", 0), explorer.make_handler(self.state))
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        base = f"http://127.0.0.1:{server.server_port}"
        paths = {
            "GET": ["/api/studio/catalog", "/api/studio/files", "/api/studio/workbook",
                    "/api/research/catalog", "/api/research/inspect", "/api/research/job"],
            "POST": ["/api/studio/create", "/api/studio/save", "/api/studio/add-sheet",
                     "/api/studio/import", "/api/research/analyze", "/api/research/export",
                     "/api/research/open-output"],
        }
        before = list(self.root.rglob("*"))
        try:
            for method, endpoints in paths.items():
                for endpoint in endpoints:
                    for headers in ({}, {"X-Explorer-Token": self.state.token, "Origin": "https://example.com"},
                                    {"X-Explorer-Token": self.state.token, "Host": "example.com"}):
                        with self.subTest(method=method, endpoint=endpoint, headers=headers):
                            request = urllib.request.Request(base + endpoint, method=method,
                                data=b"{}" if method == "POST" else None, headers=headers)
                            with self.assertRaises(urllib.error.HTTPError) as caught:
                                urllib.request.urlopen(request)
                            self.assertEqual(caught.exception.code, 403)
            self.assertEqual(list(self.root.rglob("*")), before)
            self.assertIsNone(self.state.workflow._inputs)
            self.assertIsNone(self.state.workflow._research)
        finally:
            server.shutdown()
            server.server_close()
            thread.join()

    def test_web_excel_paste_preserves_empty_cells_quotes_and_multiline_values(self):
        node = shutil.which("node")
        if not node:
            self.skipTest("Node no está disponible para comprobar el parser web.")
        script = r'''
const fs=require('node:fs'),vm=require('node:vm'),assert=require('node:assert/strict');
const source=fs.readFileSync('explorer_web/studio.js','utf8');
const start=source.indexOf('  function parseTSV(text) {');
const end=source.indexOf('  function beginEdit(',start);
assert(start>=0 && end>start,'No se encontró el parser de pegado');
const context={};vm.createContext(context);
vm.runInContext(source.slice(start,end)+';this.parseTSV=parseTSV;',context);
const plain=value=>JSON.parse(JSON.stringify(value));
assert.deepEqual(plain(context.parseTSV('A\t"línea 1\nlínea 2"\t"dice ""hola"""\r\n\t\t\r\n3\t4\t5\r\n')),
  [['A','línea 1\nlínea 2','dice "hola"'],['','',''],['3','4','5']]);
assert.deepEqual(plain(context.parseTSV('2024-01-01\t1.234,5\t20%\n2024-02-01\t\t0.054')),
  [['2024-01-01','1.234,5','20%'],['2024-02-01','','0.054']]);
assert.throws(()=>context.parseTSV('"valor sin cerrar\t123'),/comillas/);
'''
        process = subprocess.run([node, "-e", script], cwd=explorer.ROOT, capture_output=True,
                                 text=True, encoding="utf-8")
        self.assertEqual(process.returncode, 0, process.stdout + process.stderr)

    def test_research_edit_transition_keeps_explicit_file_and_sheet_during_async_load(self):
        node = shutil.which("node")
        if not node:
            self.skipTest("Node no está disponible para comprobar la transición web.")
        script = r'''
const fs=require('node:fs'),vm=require('node:vm'),assert=require('node:assert/strict');
const source=fs.readFileSync('explorer_web/studio.js','utf8');
const researchSource=fs.readFileSync('explorer_web/research.js','utf8');
function extract(startMarker,endMarker) {
  const start=source.indexOf(startMarker),end=source.indexOf(endMarker,start);
  assert(start>=0 && end>start,'No se encontraron las funciones de transición');
  return source.slice(start,end);
}
let functions=extract('  async function refreshFiles()', '  function compactEdits()')+
                extract('  async function checkpointDraft()', '  async function saved(')+
                extract('  async function activateInputs()', "  document.querySelectorAll('.module-nav [data-module]').forEach(button=>button.addEventListener");
if(process.argv.includes('--previous-activation'))functions=functions.replace('!editor.book && !editor.file','!editor.book');
const registration=researchSource.match(/nodes\.edit\.addEventListener\('click',\s*\(\)\s*=>\s*bridge\.openEditor[^;]+;/);
assert(registration,'No se encontró la acción real Llenar / editar datos');
const elements=new Map(),requests=[],messages=[];
function byId(id) {
  if(!elements.has(id))elements.set(id,{value:1,replaceChildren(){}});
  return elements.get(id);
}
const coverageFile='52_CARB_Cobertura.xlsx',researchFile='external/52_CARB_Investigación.xlsx';
const selectedSheet='5_Cola_2';
const editor={kind:'coverage',file:coverageFile,book:{file:coverageFile,sheet:'P0_Cola'},catalog:{},
  files:null,pending:new Map(),originals:new Map(),undo:[],selected:null,anchor:null,busy:false,request:0};
const context={editor,byId,URLSearchParams,Option:class{constructor(text,value){this.text=text;this.value=value;}},
  module:'research',pipelineHash:'#pipelines',location:{hash:'#research'},history:{replaceState(){}},
  document:{body:{dataset:{}},querySelectorAll(){return[];}},valueTarget:null,
  window:{CoverageApp:{context:()=>({file:{name:coverageFile}}),geometry(){},redraw(){}}},
  fillFiles(){},summary(){},commitValue(){},renderGrid(){},validation(){},busy(value){editor.busy=value;},
  message(text,error){if(text)messages.push({text,error});},
  api(path){
    if(path==='/api/studio/files')return Promise.resolve({coverage:[],research:[]});
    if(path.startsWith('/api/studio/draft?'))return Promise.resolve({draft:null,source_changed:false});
    assert(path.startsWith('/api/studio/workbook?'));
    const params=Object.fromEntries(new URL('http://localhost'+path).searchParams);
    return new Promise(resolve=>requests.push({params,resolve}));
  }};
vm.createContext(context);vm.runInContext(functions,context);
let click;
context.nodes={sheet:{value:selectedSheet},edit:{addEventListener(type,handler){assert.equal(type,'click');click=handler;}}};
context.state={file:researchFile};context.bridge={openEditor:(...args)=>context.openEditor(...args)};
vm.runInContext(registration[0],context);
const flush=async()=>{for(let count=0;count<20;count++)await Promise.resolve();};
const finish=request=>request.resolve({file:request.params.file,kind:request.params.kind,
  name:request.params.file.split('/').at(-1),sheet:request.params.sheet || 'P0_Cola',row:1,column:1,
  origin:request.params.kind==='research' ? 'external' : 'local',copy_on_save:request.params.kind==='research',
  sheets:[{name:request.params.sheet || 'P0_Cola',rows:38,columns:8}],fingerprint:'version',validation:{ready:true}});
async function run() {
  // The edit click starts a workbook request while activation also runs and book is null.
  const opening=click();await flush();
  assert.equal(requests.length,1,'Activar entradas lanzó una carga de cobertura que sustituyó la investigación');
  assert.equal(requests[0].params.kind,'research');
  assert.equal(requests[0].params.file,researchFile);
  assert.equal(requests[0].params.sheet,selectedSheet,'Se perdió el segmento seleccionado al abrir el editor');
  assert.equal(editor.kind,'research');assert.equal(editor.file,researchFile);
  assert.equal(editor.book,null,'La respuesta diferida debe conservar el estado de carga');
  finish(requests[0]);await opening;await flush();
  assert.equal(editor.book.file,researchFile);assert.equal(editor.book.sheet,selectedSheet);
  assert.equal(byId('input-kind').value,'research');assert.equal(byId('input-sheet').value,selectedSheet);
  assert.equal(messages.filter(entry=>entry.error).length,0);
  // A genuinely empty editor should still adopt the active coverage context on first activation.
  editor.file=null;editor.book=null;
  const activation=context.activateInputs();await flush();
  assert.equal(requests.length,2);assert.equal(requests[1].params.kind,'coverage');
  assert.equal(requests[1].params.file,coverageFile);
  finish(requests[1]);await activation;
  assert.equal(editor.book.file,coverageFile);assert.equal(editor.kind,'coverage');
}
run().catch(error=>{console.error(error);process.exitCode=1;});
'''
        process = subprocess.run([node, "-e", script], cwd=explorer.ROOT, capture_output=True,
                                 text=True, encoding="utf-8")
        self.assertEqual(process.returncode, 0, process.stdout + process.stderr)
        previous = subprocess.run([node, "-e", script, "--", "--previous-activation"],
                                  cwd=explorer.ROOT, capture_output=True, text=True, encoding="utf-8")
        self.assertNotEqual(previous.returncode, 0, "La prueba debe detectar la carrera del encabezado anterior.")
        self.assertIn("Activar entradas lanzó una carga de cobertura", previous.stderr)

    def test_http_capture_roundtrip_and_review_invalidation(self):
        server = explorer.ExplorerHTTPServer(("127.0.0.1", 0), explorer.make_handler(self.state))
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        base = f"http://127.0.0.1:{server.server_port}"

        def request(path, payload=None):
            data = None if payload is None else json.dumps(payload, ensure_ascii=False).encode("utf-8")
            req = urllib.request.Request(base + path, data=data, headers={
                "X-Explorer-Token": self.state.token, "Content-Type": "application/json"})
            with urllib.request.urlopen(req) as response:
                return json.load(response)

        try:
            created = request("/api/studio/create", self.coverage_spec(brands=["Cola"]))
            saved = request("/api/studio/save", {"file": created["file"], "kind": "coverage",
                "fingerprint": created["fingerprint"], "edits": [
                    {"sheet": "P0_Cola", "row": 1, "column": 2, "values": [["Weighted VOL1_P"]]},
                    {"sheet": "P0_Cola", "row": 3, "column": 1, "values": self.monthly_rows()},
                ]})
            analysis = request("/api/analyze", {"file": saved["file"], "animate_trend_pipeline": False})
            self.assertIsNone(analysis["sheets"][0]["error"])
            query = urllib.parse.urlencode({"file": saved["file"], "kind": "coverage",
                                            "sheet": "P0_Cola", "row": 3, "column": 2, "columns": 2, "limit": 2})
            view = request("/api/studio/workbook?" + query)
            self.assertEqual(view["values"], [[1000, 20.5], [1010, 20.5]])
            path = self.root / saved["file"]
            before = path.read_bytes()
            ambiguous = {"file": saved["file"], "kind": "coverage",
                "fingerprint": saved["fingerprint"], "edits": [
                    {"sheet": "P0_Cola", "row": 3, "column": 2, "values": [["1,234"]]}]}
            request("/api/studio/save", ambiguous)
            view = request("/api/studio/workbook?" + query)
            self.assertEqual(view["values"][0][0], 1234)
            with self.assertRaises(urllib.error.HTTPError) as caught:
                request("/api/save", {"analysis_id": analysis["id"], "selections": {"P0_Cola": 3}})
            self.assertEqual(caught.exception.code, 400)
            self.assertEqual(self.state.jobs, {})
        finally:
            server.shutdown()
            server.server_close()
            thread.join()


if __name__ == "__main__":
    unittest.main()
