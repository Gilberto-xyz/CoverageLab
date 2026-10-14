"""Integration checks for preview, source identity and manual export selection."""
import contextlib
import io
import json
import math
from pathlib import Path
import shutil
import subprocess
import tempfile
import threading
import time
import unittest
from unittest import mock
import urllib.error
import urllib.request

import pipeline_explorer as explorer


class PipelineExplorerTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.core = explorer.engine()

    def setUp(self):
        (explorer.ROOT / "outputs").mkdir(exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(prefix="explorer_test_", dir=explorer.ROOT / "outputs")
        self.root = Path(self.temp.name)
        self.source = self.root / "52_CARB_Explorador.xlsx"
        self.make_workbook()
        shutil.copy2(explorer.ROOT / "Modelo_PPT.pptx", self.root / "Modelo_PPT.pptx")
        self.state = explorer.ExplorerState(self.root)
        self.options = explorer.options_from_request({"animate_trend_pipeline": False})
        self.original_cwd = Path.cwd()

    def tearDown(self):
        import os
        os.chdir(self.original_cwd)
        self.temp.cleanup()

    def make_workbook(self, months=36):
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.remove(workbook.active)
        for name in ("P1_Cola", "P4_Limon"):
            sheet = workbook.create_sheet(name)
            sheet.append(["Table", "Weighted R_VOL1", "PENET", "VO1_BUY", "VO1_DAY", "FREQ", "BUYERS", "SELL-IN"])
            for index in range(months):
                date = self.core.pd.Timestamp(2023, 1, 1) + self.core.pd.DateOffset(months=index)
                volume = 1200 + index * 17 + 180 * math.sin(index * .7)
                sheet.append([date.to_pydatetime(), volume, 20 + index * .25, 6 + index * .1, 2 + index * .03, 3 + index * .02, 900 + index * 7, volume * 1.7 + math.cos(index) * 75])
        workbook.create_sheet("Sin datos").append(["vacío"])
        workbook.save(self.source)

    def analyze(self):
        with contextlib.redirect_stdout(io.StringIO()):
            return self.state.analyze({"file": self.source.name, "animate_trend_pipeline": False})

    def test_preview_shift_and_coverage_use_the_same_months(self):
        analysis = self.analyze()
        sheet = analysis["sheets"][0]
        self.assertEqual(sheet["default_pipeline"], 1)
        self.assertEqual(len(sheet["pipelines"]), 7)
        self.assertEqual(sheet["pipelines"][3]["sell_in"][:3], [None] * 3)
        self.assertEqual(sheet["pipelines"][3]["sell_in"][3:], sheet["sell_in"][:-3])
        expected = round(sum(sheet["sell_out"][-12:]) / sum(sheet["sell_in"][-15:-3]) * 100, 1)
        self.assertEqual(sheet["pipelines"][3]["coverage"], expected)
        self.assertIsNotNone(analysis["sheets"][-1]["error"])
        json.dumps(analysis, allow_nan=False)

    def test_short_history_disables_unavailable_pipelines(self):
        self.make_workbook(months=13)
        analysis = self.analyze()
        pipelines = analysis["sheets"][0]["pipelines"]
        self.assertTrue(pipelines[1]["available"])
        self.assertFalse(pipelines[2]["available"])
        self.assertIsNone(pipelines[0]["correlation"])
        with self.assertRaises(ValueError):
            self.state.reviewed_selection({"analysis_id": analysis["id"], "selections": {"P1_Cola": 6}})

    def test_monthly_penetration_retains_source_dates_and_ppt_colors(self):
        analysis = self.analyze()
        sheet = analysis["sheets"][0]
        self.assertEqual(sheet["penetration"], [20 + index * .25 for index in range(36)])
        self.assertEqual(len(sheet["penetration"]), len(sheet["dates"]))
        self.assertEqual(sheet["penetration"][-1], 28.75)
        colors = analysis["chart_colors"]
        self.assertEqual(colors["sellin"], self.core.COLOR_SELLIN_TREND_LINE)
        self.assertEqual(colors["worldpanel"], self.core.COLOR_SELLOUT_TREND_LINE)
        self.assertEqual(colors["coverage"], self.core.COLOR_COBERTURA_BAR)
        self.assertEqual(colors["penetration"], self.core.COLOR_PENETRACION_BAR)

    def test_file_catalog_decodes_names_and_retains_invalid_files(self):
        invalid = self.root / "Archivo sin códigos.xlsx"
        shutil.copy2(self.source, invalid)
        entries = {entry["name"]: entry for entry in self.state.file_catalog()}
        valid = entries[self.source.name]
        self.assertEqual(valid["country"], "Mexico")
        self.assertEqual(valid["country_code"], "52")
        self.assertEqual(valid["category_code"], "CARB")
        self.assertEqual(valid["manufacturer"], "Explorador")
        self.assertTrue(valid["category"])
        self.assertEqual(valid["size"], self.source.stat().st_size)
        self.assertIsNotNone(entries[invalid.name]["error"])
        multi = self.root / "55_MULT_Unilever.xlsx"
        shutil.copy2(self.source, multi)
        catalog = {entry["name"]: entry for entry in self.state.file_catalog()}
        self.assertTrue(catalog[multi.name]["per_sheet"])
        self.assertEqual(catalog[multi.name]["country"], "Brasil")

    def test_template_details_match_excel_windows_and_monthly_indicators(self):
        sheet = self.analyze()["sheets"][0]
        self.assertAlmostEqual(sheet["purchase_average"][-1], 9.5)
        self.assertAlmostEqual(sheet["purchase_occasion"][-1], 3.05)
        self.assertAlmostEqual(sheet["frequency"][-1], 3.7)
        self.assertEqual(sheet["buyers_series"][-1], 1145)
        averages = sheet["template"]["averages"]
        self.assertEqual(averages["Penet_MAT_Actual"], sum(sheet["penetration"][-12:]) / 12)
        self.assertEqual(averages["Buyers_MAT_Anterior"], sum(sheet["buyers_series"][-24:-12]) / 12)
        rows = sheet["template"]["variations"]
        sell_in = sheet["sell_in"]
        expected = (sum(sell_in[-9:-3]) / sum(sell_in[-15:-9]) - 1) * 100
        self.assertAlmostEqual(rows[1]["pipelines"][3], expected)
        same_period = (sum(sell_in[-9:-3]) / sum(sell_in[-21:-15]) - 1) * 100
        self.assertAlmostEqual(rows[3]["pipelines"][3], same_period)
        self.assertEqual(rows[3]["pipeline_periods"][3]["current"], [sheet["dates"][-9], sheet["dates"][-4]])
        self.assertIsNone(rows[5]["pipelines"][3])
        self.assertIsNotNone(sheet["template"]["correlations"][0]["pipelines"][3])
        self.assertIsNone(sheet["template"]["correlations"][2]["pipelines"][3])
        self.assertEqual(sheet["pipelines"][3]["yoy"][:26], [None] * 26)
        self.assertAlmostEqual(sheet["wp_yoy"][-1], (sum(sheet["sell_out"][-12:]) / sum(sheet["sell_out"][-24:-12]) - 1) * 100)

    def test_template_variations_keep_excel_zero_exclusions(self):
        with self.core.pd.ExcelFile(self.source) as workbook:
            df, _ = self.core.load_and_preprocess_sheet(workbook, "P1_Cola")
        df.loc[df.index[-2], self.core.COL_SELL_IN] = 0
        details = explorer.template_details(df, self.core)
        self.assertIsNone(details["variations"][0]["pipelines"][0])
        self.assertIsNotNone(details["variations"][0]["pipelines"][3])
        self.assertIsNotNone(details["variations"][0]["worldpanel"])
        json.dumps(details, allow_nan=False)

    def test_load_exposes_complete_dashboard_without_generating_files(self):
        before = {path.name for path in self.root.iterdir()}
        with mock.patch.object(self.core.CoverageStudioUltraApp, "process_file", side_effect=AssertionError("No debe exportar al cargar")), mock.patch.object(explorer.subprocess, "Popen", side_effect=AssertionError("No debe iniciar un generador")):
            analysis = self.analyze()
        sheet = analysis["sheets"][0]
        self.assertEqual(len(sheet["dashboard"]["kpis"]), 5)
        self.assertEqual(len(sheet["template"]["variations"]), 9)
        self.assertEqual(len(sheet["template"]["correlations"]), 3)
        self.assertEqual(self.state.jobs, {})
        self.assertEqual({path.name for path in self.root.iterdir()}, before)
        self.assertFalse((self.root / "outputs").exists())
        self.assertFalse((self.root / "Output").exists())
        json.dumps(analysis, allow_nan=False)

    def test_mat_annual_cuts_match_excel_sums_growth_and_pipeline(self):
        sheet = self.analyze()["sheets"][0]
        rows = sheet["dashboard"]["mat_periods"]
        self.assertEqual([row["date"] for row in rows], ["2023-12-01", "2024-12-01", "2025-12-01"])
        self.assertIsNone(rows[0]["sell_in"][3])
        self.assertIsNone(rows[1]["sell_in_yoy"][3])
        expected = sum(sheet["sell_in"][-15:-3])
        previous = sum(sheet["sell_in"][-27:-15])
        self.assertAlmostEqual(rows[-1]["sell_in"][3], expected)
        self.assertAlmostEqual(rows[-1]["sell_in_yoy"][3], (expected / previous - 1) * 100)
        self.assertAlmostEqual(rows[-1]["worldpanel"], sum(sheet["sell_out"][-12:]))
        self.assertEqual(rows[-1]["coverage"][3], sheet["pipelines"][3]["coverage"])

    def test_kpi_diagnostic_matches_excel_formulas_and_zero_average_signal(self):
        sheet = self.analyze()["sheets"][0]
        dashboard = sheet["dashboard"]
        kpi = next(value for value in dashboard["kpis"] if value["key"] == "purchase_average")
        recent = sheet["purchase_average"][-12:]
        average = sum(recent) / 12
        self.assertAlmostEqual(kpi["average"], average)
        self.assertAlmostEqual(kpi["total_average"], sum(sheet["purchase_average"]) / len(sheet["dates"]))
        self.assertEqual(kpi["recent_count"], 12)
        self.assertEqual(kpi["total_count"], len(sheet["dates"]))
        self.assertEqual(kpi["recent_start"], sheet["dates"][-12])
        self.assertEqual(kpi["history_start"], sheet["dates"][0])
        self.assertEqual(kpi["end"], sheet["dates"][-1])
        self.assertAlmostEqual(kpi["change"], recent[-1] / average - 1)
        self.assertAlmostEqual(kpi["range"], (max(recent) - min(recent)) / abs(average))
        self.assertEqual(kpi["color"], self.core.EXCEL_DASHBOARD_KPI_COLORS[self.core.COL_COMPRA_MEDIA])
        self.assertEqual(dashboard["change_threshold"], .25)
        self.assertEqual(dashboard["range_threshold"], .75)
        with self.core.pd.ExcelFile(self.source) as workbook:
            df, _ = self.core.load_and_preprocess_sheet(workbook, "P1_Cola")
        df[self.core.COL_BUYERS] = 0
        diagnostic = explorer.dashboard_details(df, self.core, sheet["pipelines"])
        buyers = next(value for value in diagnostic["kpis"] if value["key"] == "buyers_series")
        self.assertTrue(buyers["review"])
        self.assertEqual(buyers["change"], 0)
        self.assertEqual(buyers["range"], 0)

    def test_kpi_averages_count_zero_and_exclude_missing_values(self):
        sheet = self.analyze()["sheets"][0]
        with self.core.pd.ExcelFile(self.source) as workbook:
            df, _ = self.core.load_and_preprocess_sheet(workbook, "P1_Cola")
        df = df.tail(12).copy()
        df[self.core.COL_COMPRA_MEDIA] = [None, 0] + [0.02] * 10
        kpi = explorer.dashboard_details(df, self.core, sheet["pipelines"])["kpis"][0]
        self.assertEqual(kpi["recent_count"], 11)
        self.assertEqual(kpi["total_count"], 11)
        self.assertAlmostEqual(kpi["average"], 0.2 / 11)
        self.assertAlmostEqual(kpi["total_average"], 0.2 / 11)
        df[self.core.COL_COMPRA_MEDIA] = float('nan')
        kpi = explorer.dashboard_details(df, self.core, sheet["pipelines"])["kpis"][0]
        self.assertIsNone(kpi["average"])
        self.assertIsNone(kpi["total_average"])
        self.assertEqual(kpi["total_count"], 0)
        self.assertTrue(kpi["review"])

    def test_web_charts_use_shared_period_without_compressing_gaps(self):
        node = shutil.which("node")
        if not node:
            self.skipTest("Node no está instalado; las ventanas se validan en el navegador.")
        script = r'''
const fs=require('node:fs'),vm=require('node:vm'),assert=require('node:assert/strict');
const context={window:{}};vm.createContext(context);
vm.runInContext(fs.readFileSync('explorer_web/dashboard.js','utf8'),context);
const api=context.window.CoverageDashboard;
const dates=Array.from({length:32},(_,i)=>new Date(Date.UTC(2024,i,1)).toISOString().slice(0,10));
const series=(first,value=1)=>dates.map((_,i)=>i<first ? null : value);
const coverage={dates,series:[{values:series(12)},{values:series(0,5.3)}]};
assert.equal(api.chartData(coverage).dates[0],'2025-01-01');
const yoy={dates,series:[{values:series(23)},{values:series(23,-19)}]};
assert.equal(api.chartData(yoy).dates[0],'2025-12-01');
const shifted={dates,series:[{values:series(26)},{values:series(23)}]};
assert.equal(api.chartData(shifted).dates[0],'2026-03-01');
assert.equal(api.chartData(coverage,'12').dates[0],'2025-09-01');
const gap={dates:dates.slice(0,6),series:[{values:[null,0,1,null,2,null],yoy:[0,1,2,3,4,5]},{values:[3,4,5,6,7,8]}]};
const view=api.chartData(gap);
assert.equal(view.dates.length,4);assert.equal(view.dates[0],dates[1]);
assert.equal(view.series[0].values[0],0,'Cero es un dato válido');
assert.equal(view.series[1].values[2],null,'Los huecos deben cortar ambas curvas');
assert.equal(view.series[0].yoy[3],4,'Las etiquetas anuales conservan su fecha');
assert.equal(gap.series[1].values[3],6,'No modificar el análisis original');
assert.equal(api.chartData({dates,series:[{values:series(32)}]}).dates.length,0);
assert.equal(api.chartData({dates,series:[]}).dates.length,0);
assert.equal(api.kpiDigits({key:'purchase_average'},.054375),4);
'''
        subprocess.run([node, "-e", script], cwd=explorer.ROOT, check=True, capture_output=True, text=True)

    def test_preview_table_palette_matches_excel_constants(self):
        colors = self.analyze()["table_colors"]
        self.assertEqual(colors["correlation"], ["#" + value for value in self.core.EXCEL_CORRELATION_COLORS])
        self.assertEqual(colors["positive"]["fill"], "#" + self.core.EXCEL_VARIATION_POSITIVE_FILL)
        self.assertEqual(colors["negative"]["ink"], "#" + self.core.EXCEL_VARIATION_NEGATIVE_FONT)
        self.assertEqual(colors["coverage_current"], "#" + self.core.EXCEL_COVERAGE_CURRENT_FILL)
        self.assertEqual(colors["coverage_previous"], "#" + self.core.EXCEL_COVERAGE_PREVIOUS_FILL)

    def test_web_small_values_remain_readable_and_heatmap_preserves_excel_colors(self):
        node = shutil.which("node")
        if not node:
            self.skipTest("Node no está instalado; estas utilidades se validan en el navegador.")
        script = r'''
const fs=require('node:fs'),vm=require('node:vm'),assert=require('node:assert/strict');
const context={window:{}};vm.createContext(context);
vm.runInContext(fs.readFileSync('explorer_web/dashboard.js','utf8'),context);
const api=context.window.CoverageDashboard;
for(const values of [[.04,.05,.07],[.01,.01,.02],[.01,.01,.01]]) {
  const scale=api.axisScale(values);
  assert.ok(scale.max<.1 && scale.max>Math.max(...values),'El eje debe preservar el rango pequeño de Softys');
  const labels=Array.from({length:5},(_,i)=>api.axisTick(scale.min+i/4*(scale.max-scale.min),scale.max-scale.min));
  assert.equal(new Set(labels).size,5,'Las marcas pequeñas deben distinguirse');
}
const negatives=api.axisScale([-.02,0,.04]);assert.ok(negatives.min<-.02 && negatives.max>.04);
const zero=api.axisScale([0,0]);assert.ok(zero.max>zero.min && Number.isFinite(zero.max));
const palette=['#F8696B','#FFEB84','#63BE7B'];
assert.equal(api.correlationStyle(0,[0,.5,1,null],palette).fill,palette[0]);
assert.equal(api.correlationStyle(.5,[0,.5,1,null],palette).fill,palette[1]);
assert.equal(api.correlationStyle(1,[0,.5,1,null],palette).fill,palette[2]);
assert.equal(api.correlationStyle(.7,[.7,.7],palette).fill,palette[1]);
assert.equal(api.correlationStyle(null,[0,.5,1],palette),null);
'''
        subprocess.run([node, "-e", script], cwd=explorer.ROOT, check=True, capture_output=True, text=True)

    def test_manual_p0_has_diagnostics_without_changing_auto_candidates(self):
        analysis = self.analyze()
        with self.core.pd.ExcelFile(self.source) as workbook:
            df, _ = self.core.load_and_preprocess_sheet(workbook, "P1_Cola")
        variations = self.core.compute_variations_dataframe(df)
        automatic = self.core.build_optimal_pipeline_candidates(df, variations)
        reviewed = self.core.build_optimal_pipeline_candidates(df, variations, include_reference=True)
        self.assertEqual([candidate.pipeline for candidate in automatic], list(range(1, 7)))
        diagnostics = self.core.build_pipeline_selection_diagnostics(0, reviewed, "Elegido manualmente")
        self.assertAlmostEqual(diagnostics['Correlación seleccionada'], analysis['sheets'][0]['pipelines'][0]['correlation'], places=3)

    def test_invalid_month_sequences_are_blocked(self):
        core = self.core
        dates = core.pd.date_range("2023-01-01", periods=14, freq="MS")
        for values in (list(dates[:3]) + list(dates[2:]), list(reversed(dates)), list(dates[:5]) + list(dates[6:])):
            self.assertIsNotNone(explorer.monthly_problem(core.pd.DataFrame({core.COL_DATA: values}), core))
        self.assertIsNone(explorer.monthly_problem(core.pd.DataFrame({core.COL_DATA: dates}), core))

    def test_selections_reject_invalid_values_and_changed_source(self):
        analysis = self.analyze()
        for selection in ({}, {"Otra hoja": 1}, {"P1_Cola": True}, {"P1_Cola": 7}, {"Sin datos": 0}):
            with self.assertRaises(ValueError):
                self.state.reviewed_selection({"analysis_id": analysis["id"], "selections": selection})
        with self.source.open("ab") as source:
            source.write(b"changed")
        with self.assertRaisesRegex(ValueError, "cambió"):
            self.state.reviewed_selection({"analysis_id": analysis["id"], "selections": {"P1_Cola": 3}})

    def test_saved_choices_survive_a_new_session(self):
        analysis = self.analyze()
        self.state.save_selection({"analysis_id": analysis["id"], "selections": {"P1_Cola": 3}})
        self.state = explorer.ExplorerState(self.root)
        restored = self.analyze()
        self.assertEqual(restored["saved_selection"]["selections"], {"P1_Cola": 3})
        with self.assertRaises(ValueError):
            self.state.source("../52_CARB_Explorador.xlsx")

    def test_equal_bytes_with_different_names_keep_their_own_metadata(self):
        first = self.analyze()
        alternate = self.root / "51_CARB_OtroFabricante.xlsx"
        shutil.copy2(self.source, alternate)
        self.state.save_selection({"analysis_id": first["id"], "selections": {"P1_Cola": 3}})
        with contextlib.redirect_stdout(io.StringIO()):
            second = self.state.analyze({"file": alternate.name, "animate_trend_pipeline": False})
        self.assertEqual(first["fingerprint"], second["fingerprint"])
        self.assertNotEqual(first["selection_key"], second["selection_key"])
        self.assertEqual(second["country"], "Peru")
        self.assertEqual(second["manufacturer"], "OtroFabricante")
        self.assertIsNone(second["saved_selection"])

    def test_http_api_requires_local_session_and_origin(self):
        server = explorer.ExplorerHTTPServer(("127.0.0.1", 0), explorer.make_handler(self.state))
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        base = f"http://127.0.0.1:{server.server_port}"
        try:
            with self.assertRaises(OSError):
                explorer.ExplorerHTTPServer(("127.0.0.1", server.server_port), explorer.make_handler(self.state))
            with urllib.request.urlopen(base + "/") as response:
                self.assertIn(self.state.token, response.read().decode())
            for headers in ({}, {"X-Explorer-Token": self.state.token, "Origin": "https://example.com"},
                            {"X-Explorer-Token": self.state.token, "Host": "example.com"}):
                request = urllib.request.Request(base + "/api/files", headers=headers)
                with self.assertRaises(urllib.error.HTTPError) as error:
                    urllib.request.urlopen(request)
                self.assertEqual(error.exception.code, 403)
            request = urllib.request.Request(base + "/api/files", headers={"X-Explorer-Token": self.state.token})
            with urllib.request.urlopen(request) as response:
                self.assertEqual(json.load(response)["files"], [self.source.name])
            for endpoint in ("/api/open-results", "/api/open-output"):
                for headers in ({}, {"X-Explorer-Token": self.state.token, "Origin": "https://example.com"}):
                    request = urllib.request.Request(base + endpoint, data=b'{"scope":"all"}', headers=headers)
                    with self.assertRaises(urllib.error.HTTPError) as error:
                        urllib.request.urlopen(request)
                    self.assertEqual(error.exception.code, 403)
            self.assertFalse((self.root / "Output").exists())
            with mock.patch.object(explorer.os, "startfile", create=True) as opener, mock.patch.object(explorer.sys, "platform", "win32"):
                request = urllib.request.Request(base + "/api/open-results", data=b'{"scope":"all"}', headers={"X-Explorer-Token": self.state.token})
                with urllib.request.urlopen(request) as response:
                    self.assertEqual(Path(json.load(response)["directory"]), self.root / "Output")
                opener.assert_called_once_with(str(self.root / "Output"))
        finally:
            server.shutdown()
            server.server_close()
            thread.join()

    def test_open_results_only_opens_known_local_folders(self):
        with mock.patch.object(explorer.os, "startfile", create=True) as opener, mock.patch.object(explorer.sys, "platform", "win32"):
            for payload in ({"scope": "other", "path": str(self.root.parent)}, {"scope": "job", "id": "../outside"}):
                with self.assertRaises(ValueError):
                    self.state.open_results(payload)
            self.state.jobs["escaped"] = {"directory": str(self.root.parent)}
            with self.assertRaises(ValueError):
                self.state.open_results({"scope": "job", "id": "escaped"})
            opener.assert_not_called()
            directory = self.root / "Output" / "Entrega"
            directory.mkdir(parents=True)
            self.state.jobs["local"] = {"directory": str(directory)}
            self.state.open_results({"scope": "job", "id": "local"})
            opener.assert_called_once_with(str(directory))

    def test_restore_jobs_rejects_output_paths_outside_the_delivery(self):
        directory = self.root / "Output" / "Entrega"
        directory.mkdir(parents=True)
        job = {"id": "bad", "directory": str(directory), "status": "complete", "started_at": "2026-10-13T12:00:00-06:00",
               "outputs": [{"path": str(self.source), "label": "template"}]}
        (directory / "job.json").write_text(json.dumps(job), encoding="utf-8")
        self.assertEqual(explorer.ExplorerState(self.root).jobs, {})

    def test_restored_deliveries_follow_the_selected_source_not_global_latest(self):
        other = self.root / "51_BLEA_OtraMarca.xlsx"
        shutil.copy2(self.source, other)
        for source, job_id, hour in ((self.source, "first", 12), (other, "other", 13)):
            directory = self.root / "Output" / job_id
            directory.mkdir(parents=True)
            output = directory / "template.xlsx"
            output.write_bytes(b"test")
            (directory / "selection.json").write_text(json.dumps({"file": source.name, "options": {"coverage_type": "Absoluta"}}), encoding="utf-8")
            job = {"id": job_id, "directory": str(directory), "status": "complete",
                   "started_at": f"2026-10-13T{hour}:00:00-06:00", "outputs": [{"path": str(output), "label": "template", "name": output.name}]}
            (directory / "job.json").write_text(json.dumps(job), encoding="utf-8")
        restored = explorer.ExplorerState(self.root)
        self.assertEqual(restored.latest_job()["id"], "other")
        self.assertEqual(restored.latest_job(self.source.name)["id"], "first")
        self.assertEqual(restored.latest_job(other.name)["file"], other.name)
        self.assertEqual(restored.latest_job(self.source.name)["options"]["coverage_type"], "Absoluta")
        first_run = self.root / "52_CARB_SinEntrega.xlsx"
        shutil.copy2(self.source, first_run)
        self.assertIsNone(restored.latest_job(first_run.name))

    def test_open_outputs_only_opens_completed_matching_registered_files(self):
        directory = self.root / "Output" / "Entrega"
        directory.mkdir(parents=True)
        paths = {"template": directory / "template.xlsx", "presentation": directory / "slides.pptx",
                 "bank": directory / "bank.xlsx", "pipeline_report": directory / "pipelines.xlsx",
                 "selection": directory / "selection.json", "log": directory / "generation.log"}
        for file in paths.values():
            file.write_text("test", encoding="utf-8")
        self.state.jobs["local"] = {"id": "local", "status": "complete", "file": self.source.name,
                                    "directory": str(directory), "outputs": [{"label": label, "path": str(file)} for label, file in paths.items() if label not in {"selection", "log"}]}
        payload = {"id": "local", "file": self.source.name, "output": "template"}
        with mock.patch.object(explorer.os, "startfile", create=True) as opener, mock.patch.object(explorer.sys, "platform", "win32"):
            self.assertEqual(set(self.state.job("local")["available_outputs"]), set(paths))
            for label, file in paths.items():
                self.state.open_output(dict(payload, output=label))
                self.assertEqual(opener.call_args, mock.call(str(file)))
            opener.reset_mock()
            other = self.root / "51_BLEA_OtraMarca.xlsx"
            shutil.copy2(self.source, other)
            for invalid in (dict(payload, output="../file"), dict(payload, output=[]), dict(payload, file=other.name)):
                with self.assertRaises(ValueError):
                    self.state.open_output(invalid)
            self.state.jobs["local"]["status"] = "running"
            with self.assertRaises(ValueError):
                self.state.open_output(payload)
            self.assertEqual(self.state.job("local")["available_outputs"], [])
            self.state.jobs["local"]["status"] = "complete"
            output = self.state.jobs["local"]["outputs"][0]
            output["path"] = str(self.source)
            with self.assertRaises(ValueError):
                self.state.open_output(payload)
            executable = directory / "unsafe.exe"
            executable.write_bytes(b"test")
            output["path"] = str(executable)
            with self.assertRaises(ValueError):
                self.state.open_output(payload)
            output["path"] = str(paths["template"])
            paths["template"].unlink()
            self.assertNotIn("template", self.state.job("local")["available_outputs"])
            with self.assertRaises(ValueError):
                self.state.open_output(payload)
            opener.assert_not_called()

    def test_text_outputs_fall_back_to_notepad_only_without_a_file_association(self):
        directory = self.root / "Output" / "Entrega"
        directory.mkdir(parents=True)
        file = directory / "generation.log"
        file.write_text("Registro", encoding="utf-8")
        self.state.jobs["local"] = {"id": "local", "status": "complete", "file": self.source.name,
                                    "directory": str(directory), "outputs": []}
        missing_association = OSError("Sin aplicación asociada")
        missing_association.winerror = 1155
        with mock.patch.object(explorer.os, "startfile", create=True, side_effect=missing_association), mock.patch.object(explorer.sys, "platform", "win32"), mock.patch.object(explorer.subprocess, "Popen") as opener:
            self.state.open_output({"id": "local", "file": self.source.name, "output": "log"})
            opener.assert_called_once_with(["notepad.exe", str(file)])

    def test_real_export_respects_override_and_excludes_other_sheets(self):
        from openpyxl import load_workbook
        self.source.rename(self.root / "55_MULT_Unilever.xlsx")
        self.source = self.root / "55_MULT_Unilever.xlsx"
        source = load_workbook(self.source)
        source["P1_Cola"].title = "P1_BISC_T.UL Galletas"
        source["P4_Limon"].title = "P4_MarcaDemo"
        source.save(self.source)
        source.close()
        analysis = self.analyze()
        selected = analysis["sheets"][1]
        self.assertEqual(selected["metadata"]["categoria_codigo"], "BISC")
        job = self.state.export({"analysis_id": analysis["id"], "selections": {"P4_MarcaDemo": 3}})
        deadline = time.monotonic() + 90
        while self.state.job(job["id"])["status"] == "running" and time.monotonic() < deadline:
            time.sleep(.1)
        job = self.state.job(job["id"])
        self.assertEqual(job["status"], "complete", job.get("error") or job.get("log"))
        outputs = {out["label"]: out["path"] for out in job["outputs"]}
        directory = Path(job["directory"])
        self.assertEqual(directory.parent, self.root / "Output")
        self.assertTrue(directory.name.startswith(job["title"] + "__"))
        stage = self.root / "Output" / ".runs" / job["id"] / job["title"]
        expected_name = self.core.build_bounded_output_filename(str(stage), "Template_" + job["title"] + ".xlsx")
        self.assertEqual(Path(outputs["template"]).name, expected_name)
        self.assertEqual(job["options"], analysis["options"])
        for name in ("selection.json", "result.json", "job.json", "generation.log"):
            self.assertTrue((directory / name).is_file())
        config = json.loads((directory / "selection.json").read_text(encoding="utf-8"))
        self.assertEqual(config["options"], analysis["options"])
        self.assertEqual(config["selections"], {"P4_MarcaDemo": 3})
        self.assertEqual(explorer.ExplorerState(self.root).jobs[job["id"]]["status"], "complete")
        self.assertFalse((self.root / "Output" / ".runs" / job["id"]).exists())
        self.assertTrue(all(Path(path).is_file() for path in outputs.values()))
        bank = load_workbook(outputs["bank"], data_only=True)
        ws = bank.active
        headers = {cell.value: cell.column for cell in ws[1]}
        self.assertEqual(ws.cell(2, headers["Pipeline"]).value, 3)
        self.assertEqual(ws.cell(2, headers["Fabricante/Marca"]).value, "MarcaDemo")
        self.assertEqual(ws.cell(2, headers["Categoria"]).value, "Galletas")
        self.assertEqual(ws.cell(2, headers["Cobertura Año Mov Actual"]).value, selected["pipelines"][3]["coverage"])
        bank.close()
        template = load_workbook(outputs["template"])
        self.assertIn("P4_MarcaDemo", template.sheetnames)
        self.assertNotIn("P1_BISC_T.UL Galletas", template.sheetnames)
        self.assertNotIn("Sin datos", template.sheetnames)
        ws = template["P4_MarcaDemo"]
        headers = {cell.value: cell.column_letter for cell in ws[1] if cell.value is not None}
        coverage_chart = ws._charts[0]
        self.assertIn(f"${headers['P3']}$", coverage_chart.series[0].val.numRef.f)
        template.close()
        ppt = self.core.Presentation(outputs["presentation"])
        texts = " ".join(shape.text for slide in ppt.slides for shape in slide.shapes if shape.has_text_frame)
        self.assertIn("MarcaDemo | P3", texts)
        self.assertNotIn("T.UL Galletas |", texts)
        report = load_workbook(outputs["pipeline_report"], data_only=True)
        ws = report.active
        headers = {cell.value: cell.column for cell in ws[1]}
        self.assertEqual(ws.cell(2, headers["Pipeline"]).value, 3)
        self.assertIn("explorador", " ".join(str(cell.value) for cell in ws[2]))
        report.close()


if __name__ == "__main__":
    unittest.main()
