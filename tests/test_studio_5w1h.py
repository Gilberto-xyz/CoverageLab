"""5W1H calculation parity, blank inputs, source isolation and real report tests."""
import ast
from datetime import datetime
import hashlib
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook
from pptx import Presentation

from studio_5w1h import ENGINE_DIR, ResearchStudio, create_research_template


class ResearchStudioTests(unittest.TestCase):
    def setUp(self):
        output_root = Path(__file__).resolve().parents[1] / "outputs"
        output_root.mkdir(exist_ok=True)
        self.temporary = tempfile.TemporaryDirectory(prefix="research_test_", dir=output_root)
        self.root = Path(self.temporary.name)
        self.external = self.root / "external"
        self.external.mkdir()
        self.studio = ResearchStudio(self.root, self.external)
        self.source = self.root / "inputs" / "5w1h" / "52_CARB_Test_5w1h.xlsx"
        self.file_id = "local/52_CARB_Test_5w1h.xlsx"

    def tearDown(self):
        self.temporary.cleanup()

    def create(self, segments=None, brands=None):
        create_research_template({"country_code": "52", "category_code": "CARB", "manufacturer": "Test",
                                  "brands": brands or ["Prueba"], "segments": segments or ["5-2"],
                                  "start_date": "2023-01", "months": 36, "unit": "K"}, self.source)

    def fill_channels(self):
        self.create()
        wb = load_workbook(self.source)
        ws = wb["5_Prueba_2"]
        for row in range(2, 38):
            # 24 first months = 100, latest12 = 120/80/100.
            current = row >= 26
            for col, number in ((2, 120 if current else 100), (3, 80 if current else 100), (4, 100)):
                ws.cell(row, col, number)
                ws.cell(row, col + 5, number * 2)
        ws.cell(1, 6, "Ventas_0")
        wb.save(self.source)
        wb.close()

    def test_original_critical_functions_preserved(self):
        provenance = json.loads((ENGINE_DIR / "provenance.json").read_text(encoding="utf-8"))
        original = provenance["original_calculation_ast_sha256"]
        module = ast.parse((ENGINE_DIR / "report_engine.py").read_text(encoding="utf-8"))
        actual = {n.name: hashlib.sha256(ast.dump(n, include_attributes=False).encode()).hexdigest()
                  for n in module.body if isinstance(n, ast.FunctionDef) and n.name in original}
        self.assertEqual(actual, original)

    def test_tree_uses_percentage_or_points_mode_and_units(self):
        self.create(["2"])
        wb = load_workbook(self.source)
        ws = wb["2_Prueba_K"]
        for row in range(3, ws.max_row + 1):
            label = str(ws.cell(row, 1).value)
            if "Vert" in label:
                continue
            ws.cell(row, 2, 100)
            ws.cell(row, 3, 120)
            if label == "Weighted PENET":
                ws.cell(row, 2, 10)
                ws.cell(row, 3, 20)
        wb.save(self.source)
        result = self.studio.analyze(self.file_id, ws.title)
        row = next(row for row in result["tables"][0]["rows"] if row[0] == "% Penetracion")
        self.assertEqual(row[1:5], [10, 20, 100, "%"])
        self.assertEqual(result["provenance"]["unit"], "Kilos")
        value_row = next(row for row in result["tables"][0]["rows"] if row[0] == "Valor $ 000s")
        self.assertEqual(value_row[1:3], [0, 0])  # Original report rounds 100/120 currency to 000s.
        ws.title = "2_Prueba_K_pp"
        wb.save(self.source)
        wb.close()
        result = self.studio.analyze(self.file_id, "2_Prueba_K_pp")
        row = next(row for row in result["tables"][0]["rows"] if row[0] == "% Penetracion")
        self.assertEqual(row[1:5], [10, 20, 10, " pp"])

    def test_partial_tree_is_not_exportable(self):
        self.create(["2"])
        wb = load_workbook(self.source)
        ws = wb["2_Prueba_K"]
        ws.cell(3, 2, 100)
        ws.cell(3, 3, 120)
        wb.save(self.source)
        wb.close()
        result = self.studio.analyze(self.file_id, "2_Prueba_K")
        self.assertFalse(result["ready"])
        self.assertTrue(any("ambos períodos" in warning for warning in result["warnings"]))
        with self.assertRaises(ValueError):
            self.studio.export(self.file_id)

    def test_distribution_is_percent_points_and_excludes_total_row(self):
        self.create(["7"])
        wb = load_workbook(self.source)
        ws = next(ws for ws in wb if ws.title.startswith("7_"))
        for row, numbers in ((4, (100, 100)), (5, (70, 60)), (6, (20, 30)), (7, (10, 10))):
            for col, value in enumerate(numbers, 2):
                ws.cell(row, col, value)
        name = ws.title
        wb.save(self.source)
        wb.close()
        result = self.studio.analyze(self.file_id, name)
        self.assertTrue(result["ready"])
        chart = result["charts"][0]
        self.assertEqual(len(chart["dates"]), 3)
        self.assertEqual(chart["series"][0]["values"], [70, 20, 10])
        self.assertEqual(chart["series"][0]["unit"], "%")

    def test_confidence_intervals_are_percentage_changes(self):
        self.create(["8"])
        wb = load_workbook(self.source)
        ws = next(ws for ws in wb if ws.title.startswith("8_"))
        for row in range(3, 39):
            ws.cell(row, 2, 20)
            ws.cell(row, 3, 1000)
            ws.cell(row, 4, 100 if row < 27 else 120)
        for row in range(1, ws.max_row + 1):
            metric = str(ws.cell(row, 6).value).strip()
            if row > 26:
                continue
            pair = {"Weighted R_VOL1": (1000, 1200), "Weighted PENET": (20, 25),
                    "Weighted FREQ": (4, 5), "Weighted HHOLDS": (10000, 10000)}.get(metric)
            if pair:
                ws.cell(row, 7, pair[0]); ws.cell(row, 8, pair[1])
        name = ws.title
        wb.save(self.source); wb.close()
        result = self.studio.analyze(self.file_id, name)
        self.assertTrue(result["ready"], result["warnings"])
        model = result["charts"][0]
        self.assertEqual(model["series"][1]["unit"], "%")
        self.assertAlmostEqual(model["series"][1]["values"][0], 20)
        self.assertLess(model["series"][0]["values"][0], 20)
        self.assertGreater(model["series"][2]["values"][0], 20)

    def test_all_templates_blank_and_configured(self):
        self.create(["1", "2", "3-1", "3-2", "3-3", "3-4", "3-5", "4", "5-1", "5-2", "5-3", "5-4", "6", "6-1", "7", "8"], ["Prueba", "Segunda"])
        info = self.studio.inspect(self.file_id)
        self.assertTrue(info["requires_input"])
        self.assertTrue(all(not s["has_data"] for s in info["sheets"]))
        for sheet in info["sheets"]:
            self.assertFalse(self.studio.analyze(self.file_id, sheet["name"])["ready"])
        wb = load_workbook(self.source)
        self.assertEqual(wb["5_Prueba_2"].cell(37, 1).value, datetime(2025, 12, 1))
        self.assertEqual(wb["2_Prueba_K"].cell(2, 3).value, "MAT Dec-25")
        wb.close()
        with self.assertRaises(ValueError):
            self.studio.export(self.file_id)

    def test_channels_participation_variation_and_aporte(self):
        self.fill_channels()
        result = self.studio.analyze(self.file_id, "5_Prueba_2")
        self.assertTrue(result["ready"])
        shares = next(c for c in result["charts"] if c["id"] == "share-Compras")
        self.assertAlmostEqual(sum(s["values"][-1] for s in shares["series"]), 1.0)
        self.assertAlmostEqual(shares["series"][0]["values"][-1], 0.4)
        table = next(t for t in result["tables"] if t["id"] == "contribution-Compras")
        self.assertEqual(table["rows"][2][1:4], ["20.0%", "-20.0%", "0.0%"])
        self.assertEqual(table["rows"][3][-1], "0.0%")
        self.assertEqual(result["stats"][0]["value"], 0)
        self.assertEqual(result["provenance"]["sha256"], hashlib.sha256(self.source.read_bytes()).hexdigest())

    def test_filters_keep_whole_base_and_pipeline_alignment(self):
        self.fill_channels()
        result = self.studio.analyze(self.file_id, "5_Prueba_2", start="2025-01", end="2025-12", elements=["Canal Moderno"], pipeline=2)
        monthly = next(c for c in result["charts"] if c["id"] == "monthly")
        self.assertEqual(len(monthly["dates"]), 12)
        purchases, sales = monthly["series"]
        self.assertEqual(purchases["values"][0], 120)
        self.assertEqual(sales["values"][0], 200)
        self.assertEqual(sales["values"][2], 240)
        shares = next(c for c in result["charts"] if c["id"] == "share-Compras")
        self.assertAlmostEqual(shares["series"][0]["values"][-1], 0.4)
        self.assertTrue(any("total de la hoja" in warning for warning in result["warnings"]))
        self.assertEqual(self.studio.analyze(self.file_id, "5_Prueba_2", mode="3")["charts"][0]["series"][0]["origin"], "Compras")
        separate = self.studio.analyze(self.file_id, "5_Prueba_2", mode="2")
        self.assertIn("monthly-Ventas", [c["id"] for c in separate["charts"]])

    def test_no_elements_is_not_all_elements(self):
        self.fill_channels()
        result = self.studio.analyze(self.file_id, "5_Prueba_2", elements=[])
        self.assertFalse(result["ready"])
        self.assertEqual(result["charts"], [])
        self.assertEqual(result["provenance"]["filters"]["elements"], [])
        self.assertTrue(all(len(table["columns"]) == 2 for table in result["tables"]))
        self.assertTrue(any("No hay elementos seleccionados" in warning for warning in result["warnings"]))
        self.source.unlink()
        self.create(["7"])
        wb = load_workbook(self.source)
        ws = next(ws for ws in wb if ws.title.startswith("7_"))
        ws.cell(5, 2, 70)
        name = ws.title
        wb.save(self.source); wb.close()
        result = self.studio.analyze(self.file_id, name, elements=[])
        self.assertFalse(result["ready"])
        self.assertEqual(result["charts"], [])
        self.assertEqual(len(result["tables"][0]["columns"]), 1)

    def test_invalid_scope_and_short_history(self):
        self.fill_channels()
        for kwargs in ({"mode": "4"}, {"pipeline": 7}, {"elements": ["Inventado"]}, {"start": "2026-01", "end": "2025-01"}, {"start": "no-date"}):
            with self.assertRaises(ValueError):
                self.studio.analyze(self.file_id, "5_Prueba_2", **kwargs)
        with self.assertRaises(ValueError):
            self.studio.analyze(self.file_id, "README")
        for bad in ("local/../52_CARB_Test_5w1h.xlsx", "C:/privado.xlsx", "external/../../privado.xlsx"):
            with self.assertRaises(ValueError):
                self.studio.source(bad)
        short = self.studio.analyze(self.file_id, "5_Prueba_2", end="2023-12")
        self.assertTrue(short["ready"])
        self.assertFalse(short["tables"])
        self.assertTrue(any("24" in warning for warning in short["warnings"]))

    def test_real_report_outputs_and_external_source_unchanged(self):
        self.fill_channels()
        external = self.external / self.source.name
        external.write_bytes(self.source.read_bytes())
        initial = external.read_bytes()
        result = self.studio.export("external/" + external.name,
                                    {"mode": "1", "sheets": ["5_Prueba_2"], "pipeline_overrides": {"Prueba": 2},
                                     "fingerprint": hashlib.sha256(initial).hexdigest()}, self.root / "Output")
        self.assertEqual(external.read_bytes(), initial)
        self.assertEqual(set(result["outputs"]), {"Excel", "PPT", "Provenance", "Registro"})
        self.assertEqual(result["pipelines_applied"], {"5_Prueba_2": 2})
        self.assertGreater(len(Presentation(result["outputs"]["PPT"]).slides), 1)
        workbook = load_workbook(result["outputs"]["Excel"])
        self.assertTrue(any(ws._charts for ws in workbook))
        self.assertTrue(any(cell.data_type == "f" for ws in workbook for row in ws for cell in row))
        workbook.close()
        manifest = json.loads(Path(result["outputs"]["Provenance"]).read_text(encoding="utf-8"))
        self.assertEqual(manifest["source_sha256"], hashlib.sha256(initial).hexdigest())
        delivery = Path(result["directory"])
        self.assertEqual(delivery.parent, self.root / "Output")
        self.assertTrue(delivery.name.startswith(result["title"] + "__"))
        self.assertIn("-12-25_5W1H__", delivery.name)
        self.assertTrue(all(Path(path).is_file() and Path(path).parent == delivery for path in result["outputs"].values()))
        worker_result = json.loads((delivery / "_result.json").read_text(encoding="utf-8"))
        self.assertEqual(worker_result["Excel"], result["outputs"]["Excel"])
        self.assertEqual(worker_result["PPT"], result["outputs"]["PPT"])
        worker_request = json.loads((delivery / "_input" / "request.json").read_text(encoding="utf-8"))
        self.assertEqual(Path(worker_request["destination"]), delivery)
        self.assertTrue(Path(worker_request["source"]).is_file())
        with self.assertRaises(ValueError):
            self.studio.export(self.file_id, {"fingerprint": "obsolete"})
        with self.assertRaises(ValueError):
            self.studio.export(self.file_id, {"elements": ["Canal Moderno"]})

    def test_explicit_delivery_directory_is_not_renamed(self):
        self.fill_channels()
        initial = self.source.read_bytes()
        direct = self.root / "carpeta_explicita"
        result = self.studio.export(self.file_id, {"mode": "3", "sheets": ["5_Prueba_2"]}, direct)
        self.assertEqual(Path(result["directory"]), direct)
        self.assertTrue(all(Path(path).is_file() and Path(path).parent == direct for path in result["outputs"].values()))
        self.assertEqual(self.source.read_bytes(), initial)


if __name__ == "__main__":
    unittest.main()
