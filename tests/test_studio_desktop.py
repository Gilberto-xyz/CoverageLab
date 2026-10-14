"""Conexión de entradas con Excel: solo libros de prueba y lanzador simulado."""
from __future__ import annotations

import contextlib
import json
from pathlib import Path
import sys
import tempfile
import threading
import types
import unittest
from unittest import mock
import urllib.error
import urllib.request

from openpyxl import Workbook, load_workbook

import pipeline_explorer as explorer
import studio_desktop as desktop
from studio_inputs import InputConflictError, file_fingerprint


class DesktopStudioTests(unittest.TestCase):
    def setUp(self):
        (explorer.ROOT / "outputs").mkdir(exist_ok=True)
        self.temporary = tempfile.TemporaryDirectory(prefix="desktop_test_", dir=explorer.ROOT / "outputs")
        self.root = Path(self.temporary.name)
        self.external = self.root / "external"
        self.external.mkdir()
        self.state = explorer.ExplorerState(self.root, research_root=self.external)
        self.executable = self.root / "Office con espacios" / "EXCEL.EXE"
        self.executable.parent.mkdir()
        self.executable.write_bytes(b"mock-only; never executed")
        self.launcher = mock.Mock()
        self.connection = desktop.DesktopStudio(self.state.workflow.inputs,
                                                finder=lambda: self.executable, launcher=self.launcher)
        self.state.workflow._desktop = self.connection

    def tearDown(self):
        self.temporary.cleanup()

    def create(self):
        return self.state.workflow.create_input({"kind": "coverage", "country_code": "52",
            "category_code": "CARB", "manufacturer": "Excel de prueba", "brands": ["Cola", "Limón"]})

    @staticmethod
    def request(book, **updates):
        return {"file": book["file"], "kind": book["kind"], "fingerprint": book["fingerprint"],
                "sheet": book["sheets"][0]["name"], "row": 1, "limit": 60,
                "column": 1, "columns": 32} | updates

    def external_book(self):
        path = self.external / "52_CARB_Excel externo.xlsx"
        with contextlib.closing(Workbook()) as book:
            sheet = book.active
            sheet.title = "5_Cola_2"
            sheet.append(["Compras", "Moderno", "Tradicional"])
            sheet.append(["ene 2025", 100, 200])
            sheet["D2"] = "=B2+C2"
            sheet["D2"].number_format = "#,##0.00"
            book.save(path)
        return self.state.workflow.workbook({"file": "external/" + path.name, "kind": "research"}), path

    def test_local_open_preserves_file_and_requested_sheet(self):
        book = self.create()
        path = self.root / book["file"]
        before = path.read_bytes()
        result = self.state.workflow.open_excel(self.request(book, sheet="P0_Limón", row=2, column=2))
        self.assertTrue(result["desktop_opened"])
        self.assertEqual((result["sheet"], result["row"], result["column"]), ("P0_Limón", 2, 2))
        self.assertNotIn("copied_from", result)
        self.assertEqual(path.read_bytes(), before)
        self.launcher.assert_called_once_with(self.executable, path)

    def test_external_open_copies_exact_bytes_and_original_stays_intact(self):
        book, original = self.external_book()
        before = original.read_bytes()
        result = self.state.workflow.open_excel(self.request(book, row=2))
        copied, origin = self.connection.inputs._path(result["file"], "research")
        self.assertEqual(origin, "local")
        self.assertEqual(result["copied_from"], book["file"])
        self.assertTrue(result["desktop_opened"])
        self.assertEqual(result["sheet"], "5_Cola_2")
        self.assertEqual(copied.read_bytes(), before)
        self.assertEqual(original.read_bytes(), before)
        self.launcher.assert_called_once_with(self.executable, copied)
        second = self.state.workflow.open_excel(self.request(book))
        self.assertNotEqual(second["file"], result["file"])
        self.assertEqual(copied.read_bytes(), before)

    def test_no_excel_does_not_copy_or_launch(self):
        book, original = self.external_book()
        self.connection.finder = lambda: None
        before = original.read_bytes()
        self.assertFalse(self.connection.availability()["available"])
        with self.assertRaisesRegex(ValueError, "Excel de escritorio no está disponible"):
            self.connection.open(self.request(book))
        self.assertFalse(self.connection.inputs.research_local.exists())
        self.assertEqual(original.read_bytes(), before)
        self.launcher.assert_not_called()

    def test_stale_or_missing_fingerprint_never_launches(self):
        book = self.create()
        with contextlib.closing(load_workbook(self.root / book["file"])) as workbook:
            workbook["P0_Cola"]["B3"] = 901
            workbook.save(self.root / book["file"])
        with self.assertRaises(InputConflictError):
            self.connection.open(self.request(book))
        for value in (None, "", "not-a-fingerprint", "f" * 63):
            with self.subTest(fingerprint=value), self.assertRaises(ValueError):
                self.connection.open(self.request(book, fingerprint=value))
        self.launcher.assert_not_called()

    def test_external_race_after_import_removes_own_copy_and_never_launches(self):
        book, original = self.external_book()
        import_real = self.connection.inputs.import_workbook

        def changing_import(payload):
            result = import_real(payload)
            with contextlib.closing(load_workbook(original)) as workbook:
                workbook.active["B2"] = 999
                workbook.save(original)
            return result

        with mock.patch.object(self.connection.inputs, "import_workbook", side_effect=changing_import):
            with self.assertRaises(InputConflictError):
                self.connection.open(self.request(book))
        self.assertEqual(list(self.connection.inputs.research_local.glob("*.xlsx")), [])
        self.launcher.assert_not_called()

    def test_invalid_file_paths_and_sheet_never_copy_or_launch(self):
        book, original = self.external_book()
        for value in ("external/../" + original.name, str(original), "local/missing.xlsx", "external/~$book.xlsx"):
            with self.subTest(file=value), self.assertRaises(ValueError):
                self.connection.open(self.request(book, file=value))
        with self.assertRaisesRegex(ValueError, "hoja seleccionada"):
            self.connection.open(self.request(book, sheet="Hoja inexistente"))
        self.assertFalse(self.connection.inputs.research_local.exists())
        self.launcher.assert_not_called()

    def test_reload_reads_saved_values_and_invalidates_only_affected_coverage(self):
        book = self.create()
        self.state.analyses = {"matching": {"file": book["file"]}, "other": {"file": "otro.xlsx"}}
        path = self.root / book["file"]
        with contextlib.closing(load_workbook(path)) as workbook:
            workbook["P0_Limón"]["B3"] = 1234.5
            workbook.save(path)
        result = self.state.workflow.reload_excel(self.request(book, sheet="P0_Limón", row=3, column=2))
        self.assertEqual(result["values"][0][0], 1234.5)
        self.assertNotEqual(result["fingerprint"], book["fingerprint"])
        self.assertEqual(result["fingerprint"], file_fingerprint(path))
        self.assertEqual(list(self.state.analyses), ["other"])
        self.launcher.assert_not_called()

    def test_reload_research_invalidates_review_without_writing_or_launching(self):
        book, original = self.external_book()
        self.state.workflow.research_analyses["old"] = {"file": book["file"]}
        before = original.read_bytes()
        result = self.state.workflow.reload_excel(self.request(book))
        self.assertEqual(self.state.workflow.research_analyses, {})
        self.assertEqual(original.read_bytes(), before)
        self.assertEqual(result["file"], book["file"])
        self.launcher.assert_not_called()

    def test_registry_detection_accepts_only_existing_excel_executable(self):
        registry = types.SimpleNamespace(HKEY_CURRENT_USER=1, HKEY_LOCAL_MACHINE=2,
            KEY_WOW64_64KEY=256, KEY_WOW64_32KEY=512, KEY_READ=1, REG_SZ=1, REG_EXPAND_SZ=2,
            OpenKey=mock.Mock(side_effect=lambda *args: contextlib.nullcontext(object())),
            QueryValueEx=mock.Mock(return_value=('"' + str(self.executable) + '"', 1)))
        with mock.patch.dict(sys.modules, {"winreg": registry}), mock.patch.object(desktop.sys, "platform", "win32"):
            self.assertEqual(desktop.find_excel_executable(), self.executable.resolve())
            for invalid in (str(self.executable) + " /safe", str(self.root / "missing" / "excel.exe"),
                            str(self.root / "application.exe"), str(self.executable.parent), "excel.exe"):
                registry.QueryValueEx.return_value = (invalid, 1)
                self.assertIsNone(desktop.find_excel_executable())

    def test_unsupported_platform_has_no_registry_or_launch_action(self):
        with mock.patch.object(desktop.sys, "platform", "linux"):
            self.assertIsNone(desktop.find_excel_executable())
            connection = desktop.DesktopStudio(self.connection.inputs)
            self.assertFalse(connection.availability()["available"])
            self.assertIn("Windows", connection.availability()["message"])
        self.launcher.assert_not_called()

    def test_native_launcher_uses_exact_argument_list_without_shell(self):
        book = self.create()
        path = self.root / book["file"]
        with mock.patch.object(desktop.subprocess, "Popen") as launch:
            desktop.launch_excel(self.executable, path)
        args, kwargs = launch.call_args
        self.assertEqual(args, ([str(self.executable.resolve()), str(path)],))
        self.assertIs(kwargs["shell"], False)
        self.assertEqual(kwargs["stdin"], desktop.subprocess.DEVNULL)

    def test_http_desktop_routes_require_same_local_session(self):
        server = explorer.ExplorerHTTPServer(("127.0.0.1", 0), explorer.make_handler(self.state))
        threading.Thread(target=server.serve_forever, daemon=True).start()
        base = f"http://127.0.0.1:{server.server_port}"
        book = self.create()

        def request(path, payload=None, headers=None):
            data = json.dumps(payload).encode() if payload is not None else None
            options = {"X-Explorer-Token": self.state.token, "Origin": base,
                       "Content-Type": "application/json"} | (headers or {})
            with urllib.request.urlopen(urllib.request.Request(base + path, data=data, headers=options)) as response:
                return json.loads(response.read())

        try:
            result = request("/api/studio/desktop")
            self.assertEqual(set(result), {"available", "message"})
            self.assertTrue(result["available"])
            opened = request("/api/studio/open-excel", self.request(book))
            self.assertTrue(opened["desktop_opened"])
            self.assertEqual(request("/api/studio/reload-excel", self.request(book))["file"], book["file"])
            for path, payload in (("/api/studio/desktop", None),
                                  ("/api/studio/open-excel", self.request(book)),
                                  ("/api/studio/reload-excel", self.request(book))):
                for headers in ({"X-Explorer-Token": "invalid"}, {"Origin": "https://other.invalid"},
                                {"Host": "other.invalid"}):
                    with self.subTest(path=path, headers=headers), self.assertRaises(urllib.error.HTTPError) as error:
                        request(path, payload, headers)
                    self.assertEqual(error.exception.code, 403)
                    error.exception.close()
            self.launcher.assert_called_once()
        finally:
            server.shutdown()
            server.server_close()


if __name__ == "__main__":
    unittest.main()
