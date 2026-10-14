"""Contract checks for typed, atomic editing of real Excel input workbooks."""
import base64
from contextlib import closing
from datetime import datetime
from pathlib import Path
import tempfile
import threading
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from studio_inputs import InputConflictError, InputStudio


class InputStudioTests(unittest.TestCase):
    def setUp(self):
        output = Path(__file__).resolve().parents[1] / "outputs"
        output.mkdir(exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(prefix="input_studio_test_", dir=output)
        self.root = Path(self.temp.name) / "coverage"
        self.root.mkdir()
        self.external = Path(self.temp.name) / "original_research"
        self.external.mkdir()
        self.studio = InputStudio(self.root, self.external)

    def tearDown(self):
        self.temp.cleanup()

    def create(self, **updates):
        return self.studio.create({"country_code": "52", "category_code": "CARB",
                                   "manufacturer": "Fabricante", "brands": ["Marca"],
                                   "unit": "VOL1_P", **updates})

    def save(self, book, values, row=3, column=1, **options):
        return self.studio.save(book["file"], fingerprint=book["fingerprint"],
            kind=book["kind"], edits=[{"sheet": book["sheet"], "row": row,
                                       "column": column, "values": values}], **options)

    def external_book(self):
        path = self.external / "52_CARB_Marca.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "5_Marca_2"
        sheet.append(["Fecha", "Moderno", "Tradicional"])
        sheet.append([datetime(2025, 1, 1), 12.5, 20])
        sheet["B2"].number_format = "0.00"
        sheet["B2"].fill = PatternFill("solid", fgColor="EAF5F2")
        sheet["D2"] = "=SUM(B2:C2)"
        workbook.save(path)
        return path

    def test_creation_preserves_positional_schema_and_names_without_overwriting(self):
        book = self.create(brands=["Marca", "Otra"])
        original = (self.root / book["file"]).read_bytes()
        self.assertEqual(book["file"], "52_CARB_Fabricante.xlsx")
        self.assertEqual([s["name"] for s in book["sheets"]], ["P0_Marca", "P0_Otra"])
        self.assertEqual(book["data_row"], 3)
        self.assertFalse(book["validation"]["ready"])
        with closing(load_workbook(self.root / book["file"])) as workbook:
            self.assertEqual(workbook["P0_Marca"]["B1"].value, "Weighted VOL1_P")
            self.assertEqual(workbook["P0_Marca"]["H2"].value, "SELL-IN: Marca")
            self.assertIsNone(workbook["P0_Marca"]["B3"].value)
        second = self.create()
        self.assertNotEqual(second["file"], book["file"])
        self.assertEqual((self.root / book["file"]).read_bytes(), original)

    def test_paste_dates_locale_numbers_blank_and_real_zero_have_distinct_types(self):
        book = self.create()
        result = self.save(book, [["ene 25", "1.234,5", "20%", "2,5", "0", None, "100", "2,345.6"],
                                  ["Feb-25 ", "100", "20,5", "2", "1", "1", "10", "200"]])
        with closing(load_workbook(self.root / book["file"])) as workbook:
            sheet = workbook[book["sheet"]]
            self.assertEqual(sheet["A3"].value, datetime(2025, 1, 1))
            self.assertEqual(sheet["B3"].value, 1234.5)
            self.assertEqual(sheet["C3"].value, 20)
            self.assertEqual(sheet["D3"].value, 2.5)
            self.assertEqual(sheet["E3"].value, 0)
            self.assertIsNone(sheet["F3"].value)
            self.assertEqual(sheet["H3"].value, 2345.6)
            self.assertEqual(sheet["A4"].value, datetime(2025, 2, 1))
        self.assertEqual(result["validation"]["error_count"], 0)
        self.assertTrue(any("KPI sin dato" in warning for warning in result["validation"]["warnings"]))

    def test_rectangular_save_is_atomic_and_conflicts_do_not_overwrite(self):
        book = self.create()
        path = self.root / book["file"]
        before = path.read_bytes()
        with self.assertRaisesRegex(ValueError, "B4"):
            self.save(book, [["2025-01-01", "10"], ["2025-02-01", "texto inválido"]])
        self.assertEqual(path.read_bytes(), before)
        saved = self.save(book, [["2025-01-01", "10"]])
        expected = path.read_bytes()
        with self.assertRaises(InputConflictError):
            self.save(book, [["2025-01-01", "999"]])
        self.assertEqual(path.read_bytes(), expected)
        self.assertNotEqual(saved["fingerprint"], book["fingerprint"])

    def test_existing_formulas_formats_and_source_are_preserved_on_external_copy(self):
        path = self.external_book()
        original = path.read_bytes()
        book = self.studio.workbook(f"external/{path.name}", kind="research")
        self.assertTrue(book["copy_on_save"])
        self.assertTrue(book["readonly"][1][3])
        with self.assertRaisesRegex(ValueError, "fórmula"):
            self.save(book, [["0"]], row=2, column=4)
        self.assertFalse(self.studio.research_local.exists())
        with self.assertRaisesRegex(ValueError, "fórmulas"):
            self.save(book, [["=SUM(A1:A2)"]], row=2, column=2)
        saved = self.save(book, [["30,5"]], row=2, column=2)
        self.assertTrue(saved["file"].startswith("local/"))
        self.assertEqual(path.read_bytes(), original)
        with closing(load_workbook(self.studio.research_local / saved["name"])) as workbook:
            sheet = workbook[book["sheet"]]
            self.assertEqual(sheet["B2"].value, 30.5)
            self.assertEqual(sheet["B2"].number_format, "0.00")
            self.assertEqual(sheet["B2"].fill.fgColor.rgb, "00EAF5F2")
            self.assertEqual(sheet["D2"].value, "=SUM(B2:C2)")
        self.assertEqual(self.studio.workbook(f"local:{saved['name']}", kind="research")["file"], saved["file"])

    def test_research_percentages_are_points_and_plain_numbers_are_not_rescaled(self):
        path = self.external_book()
        book = self.studio.workbook(f"external/{path.name}", kind="research")
        saved = self.save(book, [["20%", "0,20"]], row=2, column=2)
        with closing(load_workbook(self.studio.research_local / saved["name"])) as workbook:
            self.assertEqual(workbook[book["sheet"]]["B2"].value, 20)
            self.assertEqual(workbook[book["sheet"]]["B2"].number_format, '0.0"%"')
            self.assertEqual(workbook[book["sheet"]]["C2"].value, 0.2)

    def test_import_is_exact_copy_and_rejects_unsafe_paths_and_invalid_bytes(self):
        original = self.external_book().read_bytes()
        imported = self.studio.import_workbook({"kind": "research", "filename": "Captura.xlsx",
            "content": base64.b64encode(original).decode("ascii")})
        self.assertEqual((self.studio.research_local / imported["name"]).read_bytes(), original)
        for name in ("../fuera.xlsx", "C:\\fuera.xlsx", "~$Temporal.xlsx", "fuera.xlsm"):
            with self.subTest(name=name), self.assertRaises(ValueError):
                self.studio.import_workbook({"filename": name, "content": original})
        with self.assertRaises(ValueError):
            self.studio.import_workbook({"filename": "invalid.xlsx", "content": b"not a workbook"})
        with self.assertRaises(ValueError):
            self.studio.workbook("../Captura.xlsx")

    def test_add_sheet_and_duplicate_keep_source_values_and_schema(self):
        book = self.create()
        result = self.studio.add_sheet(book["file"], fingerprint=book["fingerprint"], name="Otra")
        self.assertEqual(result["sheet"], "P0_Otra")
        self.assertEqual(result["values"][0][1], "Weighted VOL1_P")
        filled = self.save(result, [["2025-01-01", "20"]])
        duplicate = self.studio.add_sheet(filled["file"], fingerprint=filled["fingerprint"],
                                          name="P0_Copia", source_sheet="P0_Otra")
        self.assertEqual(duplicate["values"][2][1], 20)
        with self.assertRaisesRegex(ValueError, "Ya existe"):
            self.studio.add_sheet(duplicate["file"], fingerprint=duplicate["fingerprint"], name="P0_copia")

    def test_limits_and_overlapping_blocks_reject_without_disk_mutation(self):
        book = self.create()
        path = self.root / book["file"]
        before = path.read_bytes()
        with self.assertRaises(ValueError):
            self.studio.workbook(book["file"], row=20000, limit=2)
        with self.assertRaisesRegex(ValueError, "rectangular|misma cantidad"):
            self.save(book, [["10", "20"], ["30"]], row=3, column=2)
        with self.assertRaisesRegex(ValueError, "superponen"):
            self.studio.save(book["file"], fingerprint=book["fingerprint"], edits=[
                {"sheet": book["sheet"], "row": 3, "column": 2, "values": [["10"]]},
                {"sheet": book["sheet"], "row": 3, "column": 2, "values": [["20"]]},
            ])
        self.assertEqual(path.read_bytes(), before)

    def test_two_concurrent_saves_with_same_version_allow_only_one(self):
        book = self.create()
        barrier, results = threading.Barrier(2), []
        def run(number):
            barrier.wait()
            try:
                self.save(book, [["2025-01-01", str(number)]])
                results.append("saved")
            except InputConflictError:
                results.append("conflict")
        threads = [threading.Thread(target=run, args=(number,)) for number in (10, 20)]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join(timeout=20)
        self.assertCountEqual(results, ["saved", "conflict"])

    def test_number_locale_requires_explicit_choice_for_ambiguous_excel_separators(self):
        for value, locale, expected in (("1,234", "decimal-dot", 1234),
                                         ("1.234", "decimal-comma", 1234),
                                         ("1.234", "decimal-dot", 1.234),
                                         ("1,234", "decimal-comma", 1.234),
                                         ("0.054", "auto", 0.054),
                                         ("1,234,567", "auto", 1234567)):
            with self.subTest(value=value, locale=locale):
                book = self.create()
                saved = self.save(book, [[value]], column=2, number_locale=locale)
                self.assertEqual(saved["values"][2][1], expected)
        book = self.create()
        before = (self.root / book["file"]).read_bytes()
        for value in ("1,234", "1.234"):
            with self.subTest(value=value), self.assertRaisesRegex(ValueError, "ambiguo"):
                self.save(book, [[value]], column=2)
        self.assertEqual((self.root / book["file"]).read_bytes(), before)
        # True typed numbers never need string-format inference.
        saved = self.save(book, [[{"type": "number", "value": 1.234}]], column=2)
        self.assertEqual(saved["values"][2][1], 1.234)


if __name__ == "__main__":
    unittest.main()
